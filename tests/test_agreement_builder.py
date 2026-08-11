from io import BytesIO

from docx import Document
from jinja2.sandbox import SandboxedEnvironment

from services.documents.agreement_builder_service import (
    default_agreement_builder_document,
    normalize_agreement_builder_document,
    render_agreement_builder_template,
    validate_agreement_builder_document,
)
from services.documents.agreement_template_context_service import (
    AgreementVariableCatalog,
    agreement_sample_context,
    build_agreement_render_context,
)
from services.documents.docx_template_parser import parse_docx_template
from services.form_config_service import FormConfigService


def test_catalog_is_unique_complete_and_matches_real_context():
    fields = [{"name": "stanowisko", "label": "Stanowisko", "type": "text"}]
    catalog = AgreementVariableCatalog.variables(fields)
    context = agreement_sample_context(fields, form_definition={"title": "Projekt", "slug": "projekt"})

    assert len(catalog) == len({item["name"] for item in catalog})
    assert all({"name", "label", "category", "type", "example", "description", "placeholder"} <= set(item) for item in catalog)
    assert not ({item["name"] for item in catalog} - set(context))
    assert context["stanowisko"] == "Specjalista ds. projektów"
    assert next(item for item in catalog if item["name"] == "stanowisko")["category"] == "Pola formularza"

    real_context = build_agreement_render_context(
        {"data_json": {}, "agreement_number": "UM/1", "agreement_sequence": 1},
        form_definition={"title": "Projekt", "slug": "projekt", "fields": fields},
        training={"id": "one", "name": "Excel", "price": "1200"},
        all_trainings=[{"id": "one", "name": "Excel", "price": "1200"}],
        fields=fields,
    )
    assert not ({item["name"] for item in catalog} - set(real_context))
    assert real_context["stanowisko"] == ""


def test_context_exposes_address_aliases_statuses_and_training_aggregates():
    context = build_agreement_render_context(
        {
            "submission_id": "ABC-1",
            "imiona": "Anna Maria",
            "nazwisko": "Nowak",
            "ulica": "Polna",
            "nr_budynku": "4",
            "nr_lokalu": "2",
            "kod_pocztowy": "65-001",
            "miejscowosc": "Zielona Góra",
            "wojewodztwo": "lubuskie",
            "process_status": "OFFICER_ACCEPTED",
            "data_json": {"nazwa_pracodawcy": "Przykład Sp. z o.o."},
            "selected_trainings": [{"id": "one", "name": "Excel", "price": "1200", "currency": "PLN"}],
        },
        form_definition={"title": "Projekt", "project": {"number": "P/1", "program": "EFS+"}},
        training={"id": "one", "name": "Excel", "price": "1200", "currency": "PLN"},
        all_trainings=[
            {"id": "one", "name": "Excel", "price": "1200", "currency": "PLN", "is_locked": True},
            {"id": "two", "name": "Kadry", "price": "800", "currency": "PLN"},
        ],
    )

    assert context["participant_full_name"] == "Anna Maria Nowak"
    assert context["participant_address_inline"] == "ul. Polna 4/2, 65-001 Zielona Góra, woj. lubuskie"
    assert context["submission_status_label"] == "Wniosek zaakceptowany"
    assert context["project_number"] == "P/1"
    assert context["training_name"] == "Excel"
    assert context["selected_trainings_count"] == 1
    assert context["all_selected_trainings_total_formatted"] == "2 000,00 zł"
    assert context["locked_trainings_total_formatted"] == "1 200,00 zł"
    assert "Excel" in str(context["selected_trainings_table"])
    assert context["nazwa_pracodawcy"] == "Przykład Sp. z o.o."


def test_context_derives_training_details_from_snapshot_dates_when_top_level_values_are_empty():
    training = {
        "id": "dated",
        "name": "Szkolenie terminowe",
        "location": "",
        "dates": [
            {"start_date": "2026-09-10", "start_time": "09:00", "end_time": "15:00", "location": "Sala 2", "description": "Dzień pierwszy"},
            {"start_date": "2026-09-11", "start_time": "10:00", "end_time": "14:00", "location": "Sala 3"},
        ],
    }

    context = build_agreement_render_context({}, training=training, all_trainings=[training])

    assert context["training_start_date"] == "2026-09-10"
    assert context["training_end_date"] == "2026-09-11"
    assert context["training_location"] == "Sala 2"
    assert context["training_description"] == "Dzień pierwszy"


def test_builder_renders_formal_blocks_lists_tables_and_components():
    document = {
        "version": 1,
        "blocks": [
            {"type": "heading", "level": 1, "alignment": "center", "content": "Umowa {{ agreement_number }}"},
            {"type": "paragraph", "alignment": "justify", "content": "Uczestnik: <strong>{{ participant_name }}</strong>"},
            {"type": "agreement_section", "number": "§ 3.", "title": "Prawa i obowiązki"},
            {"type": "ordered_list", "items": [{"content": "Punkt", "level": 0}, {"content": "Podpunkt", "level": 1}]},
            {"type": "table", "header": True, "rows": [["A", "B"], ["1", "2"]]},
            {"type": "training_table", "scope": "all_selected_trainings", "columns": ["index", "name", "price", "date"], "show_total": True},
            {"type": "participant_data", "fields": ["participant_name", "pesel", "participant_address", "telefon"]},
            {"type": "signatures", "left_label": "Beneficjent", "right_label": "Uczestnik projektu"},
            {"type": "project_info", "fields": ["project_name", "project_number"], "show_logo": True},
            {"type": "page_break"},
        ],
    }

    errors = validate_agreement_builder_document(document)
    template = render_agreement_builder_template(document)
    rendered = SandboxedEnvironment(autoescape=True).from_string(template).render(**agreement_sample_context())

    assert errors == []
    assert '<section class="document-section">' in template
    assert "document-list__item--level-1" in template
    assert "document-training-table" in template
    assert "document-participant-data" in template
    assert "document-signatures" in template
    assert "document-project-info" in template
    assert "document-page-break" in template
    assert "UM/PRZYKLAD/2026" in rendered
    assert "Przykładowe szkolenie" in rendered


def test_builder_normalizes_legacy_alignment_and_renders_bounded_format_classes():
    legacy = {
        "version": 1,
        "blocks": [
            {
                "type": "paragraph",
                "alignment": "center",
                "style": "position:fixed;left:0",
                "content": "Wyróżniony akapit",
                "format": {"bold": True, "italic": True, "underline": True},
            },
            {"type": "heading", "level": 1, "alignment": "right", "content": "Nagłówek"},
        ],
    }

    normalized = normalize_agreement_builder_document(legacy)
    template = render_agreement_builder_template(normalized)

    assert normalized["blocks"][0]["format"] == {
        "bold": True,
        "italic": True,
        "underline": True,
        "alignment": "center",
    }
    assert "alignment" not in normalized["blocks"][0]
    assert "style" not in normalized["blocks"][0]
    assert 'document-align-center document-bold document-italic document-underline' in template
    assert 'document-align-right document-bold' in template
    assert "position:fixed" not in template


def test_builder_validator_reports_unknown_variable_jinja_and_empty_blocks():
    unknown = {"version": 1, "blocks": [{"type": "paragraph", "content": "{{ participant_company_xyz }}"}]}
    broken = {"version": 1, "blocks": [{"type": "paragraph", "content": "{% if pesel %}"}]}
    empty = {"version": 1, "blocks": []}

    unknown_errors = validate_agreement_builder_document(unknown)
    broken_errors = validate_agreement_builder_document(broken)

    assert unknown_errors[0].variable == "participant_company_xyz"
    assert unknown_errors[0].path == "blocks.0"
    assert "Niepoprawna składnia Jinja" in broken_errors[0].message
    assert validate_agreement_builder_document(empty)[0].message == "Dokument nie zawiera treści."


def test_builder_preserves_legacy_submission_get_expression_while_sanitizing_html():
    document = {
        "version": 1,
        "blocks": [
            {"type": "paragraph", "content": '<script>alert(1)</script>{{ submission.get("ulica", "") }}'},
        ],
    }

    normalized = normalize_agreement_builder_document(document)
    template = render_agreement_builder_template(normalized)
    rendered = SandboxedEnvironment(autoescape=True).from_string(template).render(submission={"ulica": "Polna"})

    assert "<script>" not in template
    assert '{{ submission.get("ulica", "") }}' in template
    assert "Polna" in rendered


def test_docx_import_builds_editable_sections_multilevel_lists_variables_and_signatures():
    document = Document()
    document.add_heading("Umowa {{ agreement_number }}", level=1)
    document.add_paragraph("§ 1.")
    document.add_paragraph("Definicje")
    first = document.add_paragraph("Uczestnik {{ participant_name }}", style="List Number")
    second = document.add_paragraph("PESEL {{ pesel }}", style="List Number")
    second.paragraph_format.left_indent = first.paragraph_format.left_indent
    document.add_paragraph("{% if nr_lokalu %}")
    document.add_paragraph("Adres: {{ participant_address_inline }}")
    document.add_paragraph("{% endif %}")
    signatures = document.add_table(rows=2, cols=2)
    signatures.cell(0, 0).text = "Beneficjent"
    signatures.cell(0, 1).text = "Uczestnik projektu"
    signatures.cell(1, 0).text = "................"
    signatures.cell(1, 1).text = "................"
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_docx_template(buffer.getvalue())
    builder = parsed.builder_document

    assert builder and builder["version"] == 1
    assert any(block["type"] == "agreement_section" and "§ 1" in block["number"] for block in builder["blocks"])
    assert any(block["type"] == "ordered_list" for block in builder["blocks"])
    assert any(block["type"] == "table" and "Beneficjent" in str(block["rows"]) for block in builder["blocks"])
    assert {"agreement_number", "participant_name", "pesel", "nr_lokalu", "participant_address_inline"} <= set(parsed.variables)


def test_docx_import_splits_multiline_main_heading_into_heading_and_context_paragraphs():
    document = Document()
    heading = document.add_heading(level=1)
    heading.add_run("Umowa uczestnictwa nr {{ agreement_number }}")
    heading.add_run().add_break()
    heading.add_run("w projekcie Rozwój kompetencji")
    heading.add_run().add_break()
    heading.add_run("współfinansowanym z EFS+")
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_docx_template(buffer.getvalue())
    blocks = parsed.builder_document["blocks"]

    assert blocks[0]["type"] == "heading"
    assert blocks[0]["content"] == "Umowa uczestnictwa nr {{ agreement_number }}"
    assert blocks[0]["format"]["alignment"] == "center"
    assert [block["type"] for block in blocks[1:3]] == ["paragraph", "paragraph"]
    assert all(block["format"]["alignment"] == "center" for block in blocks[1:3])
    assert parsed.html.count("<h1") == 1
    assert parsed.html.count("document-heading-context") == 2


def test_default_builder_is_valid_and_uses_per_training_variables():
    document = default_agreement_builder_document()
    template = render_agreement_builder_template(document)

    assert validate_agreement_builder_document(document) == []
    assert "agreement_number" in template
    assert "selected_trainings" in template


def test_form_config_prefers_builder_for_new_forms_and_preserves_legacy_html():
    service = FormConfigService()

    fresh = service.normalize_form_config({"title": "Nowy", "workflow": {"requires_contract": True}})
    legacy = service.normalize_form_config({"title": "Stary", "workflow": {"requires_contract": True, "contract_template_html": "<p>Legacy</p>"}})

    assert fresh["workflow"]["contract_template_source"] == "builder"
    assert fresh["workflow"]["contract_builder_document"]["version"] == 1
    assert next(item for item in fresh["documents"] if item["id"] == "agreement")["template_source"] == "builder"
    assert legacy["workflow"]["contract_template_source"] == "html"


def test_complex_docx_structural_replica_imports_sections_one_to_eight_and_legacy_variables():
    document = Document()
    document.add_heading("Umowa nr {{ agreement_number }}", level=1)
    document.add_paragraph("Uczestnik {{ participant_name }}, PESEL {{ pesel }}")
    document.add_paragraph("Adres {{ participant_address_inline }}, telefon {{ telefon }}")
    document.add_paragraph("Data {{ generated_date }}, zgłoszenie {{ submission_id }}")
    document.add_paragraph("Szkolenie {{ training_name }}, kwota {{ all_selected_trainings_total_formatted }}")
    for section_number in range(1, 9):
        document.add_paragraph(f"§ {section_number}.")
        document.add_paragraph(f"Tytuł paragrafu {section_number}")
        document.add_paragraph(f"Punkt paragrafu {section_number}", style="List Number")
        nested = document.add_paragraph(f"Podpunkt paragrafu {section_number}", style="List Number 2")
        num_pr = nested._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 1
    signatures = document.add_table(rows=2, cols=2)
    signatures.cell(0, 0).text = "Beneficjent"
    signatures.cell(0, 1).text = "Uczestnik projektu"
    signatures.cell(1, 0).text = "................"
    signatures.cell(1, 1).text = "................"
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_docx_template(buffer.getvalue())
    builder = parsed.builder_document
    sections = [block for block in builder["blocks"] if block["type"] == "agreement_section"]
    list_items = [item for block in builder["blocks"] if block["type"] == "ordered_list" for item in block["items"]]

    assert len(sections) == 8
    assert [section["number"] for section in sections] == [f"§ {number}." for number in range(1, 9)]
    assert any(item["level"] == 1 for item in list_items)
    assert any(block["type"] == "table" and "Uczestnik projektu" in str(block["rows"]) for block in builder["blocks"])
    assert {"agreement_number", "participant_name", "pesel", "participant_address_inline", "telefon", "generated_date", "submission_id", "training_name", "all_selected_trainings_total_formatted"} <= set(parsed.variables)
