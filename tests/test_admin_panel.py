import io
import json
import logging
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

pytest.importorskip("sqlalchemy")

from flask import url_for
from sqlalchemy import select, text
from werkzeug.security import generate_password_hash

from conftest import InMemoryStorage
from config import Config
from database import create_session_factory
from models import (
    AccessRole,
    AccessRolePermission,
    ContactPage,
    EmailLog,
    Form,
    FormField,
    FormPermission,
    FormUserRole,
    FormRegulation,
    FormSubmission,
    Logo,
    MailFooter,
    MailTemplate,
    MailTemplateAsset,
    PlatformMailTemplate,
    Permission,
    ServiceDocument,
    SubmissionDecision,
    SubmissionAssignmentHistory,
    SubmissionFile,
    SubmissionInternalNote,
    SubmissionTraining,
    SubmissionWorkflowEvent,
    SystemMailSettings,
    User,
    UserGlobalRole,
)
from services.admin_form_service import build_definition_from_html, sync_form_fields
from form_loader import normalize_form_definition, validate_form_definition


class AdminTestConfig(Config):
    TESTING = True
    WTF_CSRF_ENABLED = False
    SERVER_NAME = "localhost"
    NEXTCLOUD_BASE_URL = "https://nextcloud.test"
    NEXTCLOUD_USERNAME = "tester"
    NEXTCLOUD_APP_PASSWORD = "secret"
    NEXTCLOUD_FORMS_DIR = "Formularze"
    NEXTCLOUD_OUTPUT_DIR = "output"
    SMTP_HOST = "smtp.test"
    SMTP_USER = "user"
    SMTP_PASSWORD = "secret"
    MAIL_FROM = "noreply@example.com"
    AUTO_CREATE_DB_SCHEMA = True


@pytest.fixture()
def admin_app(tmp_path, form_definition):
    import app as app_module

    class TestConfig(AdminTestConfig):
        TEMP_DIR = tmp_path / "tmp"
        DATABASE_URL = f"sqlite:///{tmp_path / 'admin.db'}"

    flask_app = app_module.create_app(config_object=TestConfig, storage_override=InMemoryStorage(form_definition))
    flask_app.config.update(TESTING=True, TEMP_DIR=tmp_path / "tmp")
    Path(flask_app.config["TEMP_DIR"]).mkdir(parents=True, exist_ok=True)
    yield flask_app


@pytest.fixture()
def admin_client(admin_app):
    return admin_app.test_client()


def create_user(app, email="admin@example.com", password="secret", role="super_admin"):
    session_factory = create_session_factory(app.config["DATABASE_URL"])
    with session_factory() as db:
        user = User(
            email=email,
            password_hash=generate_password_hash(password),
            role=role,
            is_active=True,
            is_blocked=False,
        )
        db.add(user)
        db.commit()
        return user.id


def create_form(app, slug="sample_form", name="Sample", user_id=None, **kwargs):
    session_factory = create_session_factory(app.config["DATABASE_URL"])
    with session_factory() as db:
        form = Form(
            slug=slug,
            name=name,
            title=kwargs.pop("title", name),
            definition_json=kwargs.pop("definition_json", {"title": name, "fields": []}),
            created_by_id=user_id,
            **kwargs,
        )
        db.add(form)
        db.flush()
        if user_id:
            db.add(FormPermission(user_id=user_id, form_id=form.id, can_manage=True))
        db.commit()
        return form.id


def login(client, email="admin@example.com", password="secret"):
    response = client.get("/admin/")
    html = response.get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    return client.post("/admin/", data={"email": email, "password": password, "csrf_token": token})


def grant_rbac(app, user_id, permission_keys, *, form_id=None, global_scope=False):
    factory = create_session_factory(app.config["DATABASE_URL"])
    with factory() as db:
        scope = "global" if global_scope else "form"
        role = AccessRole(key=f"test_{user_id}_{scope}_{'_'.join(permission_keys)}", name="Test role", scope=scope, is_active=True)
        db.add(role)
        db.flush()
        for key in permission_keys:
            permission = db.execute(select(Permission).where(Permission.key == key)).scalar_one_or_none()
            if permission is None:
                permission = Permission(key=key, name=key, scope=scope, category=scope, is_active=True)
                db.add(permission)
                db.flush()
            db.add(AccessRolePermission(role_id=role.id, permission_id=permission.id))
        db.flush()
        if global_scope:
            db.add(UserGlobalRole(user_id=user_id, role_id=role.id))
        else:
            db.add(FormUserRole(user_id=user_id, form_id=form_id, role_id=role.id))
        db.commit()
        return role.id


def test_rbac_backend_separates_view_decision_and_mail(admin_app, admin_client):
    user_id = create_user(admin_app, email="rbac@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="rbac_form")
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        db.add(FormField(form_id=form_id, name="email", label="E-mail", type="email", data_classification="personal"))
        submission = FormSubmission(
            submission_id="rbac-submission", form_slug="rbac_form",
            process_status="WAITING_FOR_OFFICER_DECISION", email="participant@example.org",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    grant_rbac(admin_app, user_id, ["can_view_submissions"], form_id=form_id)
    login(admin_client, email="rbac@example.com")

    detail = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}")
    assert detail.status_code == 200
    assert b"participant@example.org" not in detail.data
    variables = admin_client.get(f"/admin/submissions/{submission_pk}/mail-variables")
    assert variables.status_code == 200
    assert b"participant@example.org" not in variables.data
    assert admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").status_code == 403
    assert admin_client.get(f"/admin/forms/{form_id}/submissions/export.csv").status_code == 403
    assert admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": admin_csrf(admin_client), "officer_decision": "accepted"},
    ).status_code == 403

    grant_rbac(admin_app, user_id, ["can_send_email", "can_make_decision", "can_export_data"], form_id=form_id)
    assert admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").status_code == 200
    export = admin_client.get(f"/admin/forms/{form_id}/submissions/export.csv")
    assert export.status_code == 200
    assert b"participant@example.org" not in export.data
    assert admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": admin_csrf(admin_client), "officer_decision": "accepted"},
    ).status_code == 302


def test_rbac_sensitive_attachment_download_requires_permission(admin_app, admin_client):
    user_id = create_user(admin_app, email="files@example.com", role="form_manager")
    form_id, owner_pk, _, file_id, _ = create_participant_attachment(admin_app)
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        db.get(SubmissionFile, file_id).data_classification = "sensitive"
        db.commit()
    grant_rbac(admin_app, user_id, ["can_view_submissions"], form_id=form_id)
    login(admin_client, email="files@example.com")
    url = f"/admin/forms/{form_id}/submissions/{owner_pk}/attachments/{file_id}/download"
    assert admin_client.get(url).status_code == 403
    grant_rbac(admin_app, user_id, ["can_view_sensitive_data"], form_id=form_id)
    assert admin_client.get(url).status_code == 200


def test_global_manage_users_role_is_enforced_by_backend(admin_app, admin_client):
    user_id = create_user(admin_app, email="global-admin@example.com", role="form_manager")
    grant_rbac(admin_app, user_id, ["can_manage_users"], global_scope=True)
    login(admin_client, email="global-admin@example.com")
    assert admin_client.get("/admin/users").status_code == 200


def test_super_admin_assigns_submission_and_route_writes_audit(admin_app, admin_client):
    admin_id = create_user(admin_app)
    assignee_id = create_user(admin_app, email="officer@example.com", role="form_manager")
    form_id = create_form(admin_app)
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        db.add(FormPermission(
            user_id=assignee_id, form_id=form_id, can_manage=False,
            can_review=True, can_assign_submissions=True,
        ))
        submission = FormSubmission(submission_id="assigned-route", form_slug="sample_form")
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client)
    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/assignment",
        data={
            "csrf_token": admin_csrf(admin_client), "assigned_to_user_id": str(assignee_id),
            "priority": "high", "assignment_reason": "Podział pracy",
        },
    )
    assert response.status_code == 302
    with factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        history = db.execute(select(SubmissionAssignmentHistory)).scalar_one()
        assert submission.assigned_to_user_id == assignee_id
        assert submission.priority == "high"
        assert history.assigned_by_user_id == admin_id
        assert history.reason == "Podział pracy"


def test_form_manager_without_assignment_permission_gets_403(admin_app, admin_client):
    manager_id = create_user(admin_app, email="manager@example.com", role="form_manager")
    assignee_id = create_user(admin_app, email="officer2@example.com", role="form_manager")
    form_id = create_form(admin_app)
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        db.add_all([
            FormPermission(user_id=manager_id, form_id=form_id, can_manage=False, can_review=True, can_assign_submissions=False),
            FormPermission(user_id=assignee_id, form_id=form_id, can_manage=False, can_review=True),
        ])
        submission = FormSubmission(submission_id="forbidden-assignment", form_slug="sample_form")
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client, email="manager@example.com")
    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/assign-selected",
        data={
            "csrf_token": admin_csrf(admin_client), "submission_pk_ids": str(submission_pk),
            "assigned_to_user_id": str(assignee_id),
        },
    )
    assert response.status_code == 403
    claim = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/claim",
        data={"csrf_token": admin_csrf(admin_client)},
    )
    assert claim.status_code == 403


def test_sla_is_read_only_without_workflow_permission_and_editable_with_it(admin_app, admin_client):
    user_id = create_user(admin_app, email="sla-viewer@example.com", role="form_manager")
    form_id = create_form(
        admin_app,
        slug="sla-rbac",
        definition_json={
            "workflow": {
                "initial_step": "review",
                "steps": [{"id": "review", "admin_label": "Weryfikacja"}],
            }
        },
    )
    grant_rbac(admin_app, user_id, ["can_view_submissions"], form_id=form_id)
    login(admin_client, email="sla-viewer@example.com")
    url = f"/admin/forms/{form_id}/sla"

    read_only = admin_client.get(url)
    assert read_only.status_code == 200
    assert "Zapisz SLA" not in read_only.get_data(as_text=True)
    assert admin_client.post(url, data={"csrf_token": admin_csrf(admin_client)}).status_code == 403

    grant_rbac(admin_app, user_id, ["can_edit_workflow"], form_id=form_id)
    editable = admin_client.get(url)
    assert editable.status_code == 200
    assert "Zapisz SLA" in editable.get_data(as_text=True)


def test_submission_lists_and_dashboard_require_view_permission_per_form(admin_app, admin_client):
    user_id = create_user(admin_app, email="scoped-dashboard@example.com", role="form_manager")
    visible_id = create_form(admin_app, slug="visible-scope")
    hidden_id = create_form(admin_app, slug="hidden-scope")
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        db.add_all([
            FormSubmission(submission_id="VISIBLE-CASE", form_slug="visible-scope"),
            FormSubmission(submission_id="HIDDEN-CASE", form_slug="hidden-scope"),
        ])
        db.commit()
    grant_rbac(admin_app, user_id, ["can_view_submissions"], form_id=visible_id)
    grant_rbac(admin_app, user_id, ["can_edit_workflow"], form_id=hidden_id)
    login(admin_client, email="scoped-dashboard@example.com")

    listing = admin_client.get("/admin/submissions").get_data(as_text=True)
    assert "VISIBLE-CASE" in listing
    assert "HIDDEN-CASE" not in listing
    dashboard = admin_client.get("/admin/dashboard").get_data(as_text=True)
    assert "Nieprzydzielone</span><strong>1</strong>" in dashboard


def test_internal_note_routes_enforce_dedicated_permissions(admin_app, admin_client):
    user_id = create_user(admin_app, email="notes-rbac@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="notes-rbac")
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        submission = FormSubmission(submission_id="NOTES-RBAC", form_slug="notes-rbac")
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    grant_rbac(admin_app, user_id, ["can_view_submissions"], form_id=form_id)
    login(admin_client, email="notes-rbac@example.com")
    url = f"/admin/forms/{form_id}/submissions/{submission_pk}/internal-notes"

    denied = admin_client.post(
        url,
        data={"csrf_token": admin_csrf(admin_client), "content": "Poufna notatka"},
    )
    assert denied.status_code == 403

    grant_rbac(
        admin_app, user_id,
        ["can_view_internal_notes", "can_add_internal_notes"],
        form_id=form_id,
    )
    created = admin_client.post(
        url,
        data={"csrf_token": admin_csrf(admin_client), "content": "Poufna notatka"},
    )
    assert created.status_code == 302
    with factory() as db:
        note = db.execute(select(SubmissionInternalNote)).scalar_one()
        assert note.author_user_id == user_id


def create_participant_attachment(app):
    form_id = create_form(
        app,
        slug="attachment_form",
        definition_json={"title": "Attachments", "fields": [{"type": "file", "name": "proof", "label": "Dowód"}]},
    )
    session_factory = create_session_factory(app.config["DATABASE_URL"])
    with session_factory() as db:
        owner = FormSubmission(submission_id="attachment-owner", form_slug="attachment_form")
        stranger = FormSubmission(submission_id="attachment-stranger", form_slug="attachment_form")
        db.add_all([owner, stranger])
        db.flush()
        attachment = SubmissionFile(
            submission_id=owner.id,
            public_submission_id=owner.submission_id,
            form_slug="attachment_form",
            filename="safe.pdf",
            original_filename="proof.pdf",
            storage_path="output/attachment_form/submissions/attachment-owner/attachments/proof/safe.pdf",
            mime_type="application/pdf",
            field_key="proof",
            attachment_version=1,
            status="active",
        )
        db.add(attachment)
        db.commit()
        result = (form_id, owner.id, stranger.id, attachment.id, attachment.storage_path)
    app.extensions["services"].storage.write_bytes(result[4], b"%PDF-1.4\ntest", "application/pdf")
    return result


def readable_workflow_definition():
    return {
        "title": "Workflow Form",
        "fields": [],
        "workflow": {
            "name": "Obsługa wniosku",
            "initial_step": "submission",
            "legacy_extension": {"preserve": True},
            "steps": [
                {
                    "id": "submission",
                    "status": "application_submitted",
                    "next": "review",
                },
                {
                    "id": "review",
                    "admin_label": "Ocena wniosku",
                    "user_label": "Weryfikacja wniosku",
                    "status": "OFFICER_REVIEW",
                    "next": "completed",
                    "requires_officer_action": True,
                },
                {
                    "id": "completed",
                    "admin_label": "Zakończenie",
                    "user_label": "Proces zakończony",
                    "status": "PROCESS_COMPLETED",
                    "final": True,
                },
            ],
        },
    }


def create_rollback_submission(app, *, form_slug="rollback_form", email="participant@example.com", status="AGREEMENT_WAITING_FOR_SIGNATURE"):
    session_factory = create_session_factory(app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="rollback-public-uuid",
            form_slug=form_slug,
            form_name="Rollback Form",
            email=email,
            process_status=status,
            workflow_step="agreement_signature",
            officer_decision="accepted",
            declaration_required="Tak",
            declaration_generated="Tak",
            declaration_filename="declaration.pdf",
            declaration_signed="Tak",
            declaration_signature_valid="Tak",
            declaration_signed_filename="signed-declaration.pdf",
            agreement_required="Tak",
            agreement_generated="Tak",
            agreement_filename="agreement.pdf",
            training_agreements='[{"id":"one","filename":"agreement.pdf"}]',
        )
        db.add(submission)
        db.commit()
        return submission.id, submission.submission_id


def create_blocked_agreement_submission(
    app,
    *,
    form_slug="blocked_form",
    submission_id="blocked-public-uuid",
):
    session_factory = create_session_factory(app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id=submission_id,
            form_slug=form_slug,
            form_name="Blocked Form",
            email="participant@example.com",
            process_status="AGREEMENT_BLOCKED",
            workflow_step="agreement",
            officer_decision="accepted",
            officer_decision_reason="",
            declaration_required="Tak",
            declaration_generated="Tak",
            declaration_signature_valid="Tak",
            agreement_required="Tak",
            agreement_blocked="Tak",
            agreement_block_reason="Warunki deklaracji nie zostały spełnione.",
            data_json={
                "_agreement_block_source": "Warunki deklaracji",
                "_agreement_blocked_at": "2026-07-28T09:00:00+00:00",
            },
        )
        db.add(submission)
        db.commit()
        return submission.id, submission.submission_id


def admin_csrf(client):
    with client.session_transaction() as session:
        return session["admin_csrf_token"]


def test_participant_attachment_download_is_scoped_to_owning_submission(admin_app, admin_client):
    create_user(admin_app)
    form_id, owner_pk, stranger_pk, file_id, _ = create_participant_attachment(admin_app)
    login(admin_client)

    response = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{owner_pk}/attachments/{file_id}/download"
    )
    detail = admin_client.get(f"/admin/forms/{form_id}/submissions/{owner_pk}")
    foreign = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{stranger_pk}/attachments/{file_id}/download"
    )

    assert response.status_code == 200
    assert detail.status_code == 200
    assert "Załączniki uczestnika" in detail.get_data(as_text=True)
    assert response.data.startswith(b"%PDF")
    assert "attachment" in response.headers["Content-Disposition"]
    assert foreign.status_code == 404


def test_admin_can_request_attachment_replacement_without_deleting_history(admin_app, admin_client):
    create_user(admin_app)
    form_id, owner_pk, _, file_id, _ = create_participant_attachment(admin_app)
    login(admin_client)

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{owner_pk}/attachments/{file_id}/status",
        data={"csrf_token": admin_csrf(admin_client), "attachment_status": "requires_correction", "reason": "Nieczytelny skan"},
    )
    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        attachment = db.get(SubmissionFile, file_id)
        assert attachment.status == "requires_correction"
        assert attachment.rejection_reason == "Nieczytelny skan"


def test_admin_can_create_secure_conditional_file_field(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="file_config_form")
    login(admin_client)
    response = admin_client.post(
        f"/admin/forms/{form_id}/fields",
        data={
            "csrf_token": admin_csrf(admin_client), "action": "add",
            "new_name": "employment_certificate", "new_label": "Zaświadczenie", "new_type": "file",
            "new_required": "on", "new_allowed_extensions": "pdf", "new_allowed_mime_types": "application/pdf",
            "new_max_size_mb": "10", "new_max_files": "1", "new_document_type": "employment_certificate",
            "new_category": "employment", "new_required_if_field": "employment_status",
            "new_required_if_operator": "equals", "new_required_if_value": "employed",
        },
    )
    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        field = next(item for item in form.definition_json["fields"] if item.get("name") == "employment_certificate")
        assert field["type"] == "file"
        assert field["allowed_extensions"] == ["pdf"]
        assert field["required_if"] == {"field": "employment_status", "operator": "equals", "value": "employed"}


@pytest.mark.parametrize("role,email", [("admin", "rollback-admin@example.com"), ("super_admin", "rollback-super@example.com")])
def test_stage_rollback_button_is_visible_for_admin_roles(admin_app, admin_client, role, email):
    user_id = create_user(admin_app, email=email, role=role)
    form_id = create_form(admin_app, slug="rollback_form", name="Rollback Form", user_id=user_id)
    submission_pk, submission_id = create_rollback_submission(admin_app)
    login(admin_client, email=email)

    response = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{submission_pk}",
        environ_overrides={"SCRIPT_NAME": "/aplikacja"},
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Cofnij etap" in html
    assert f'/aplikacja/admin/submissions/{submission_id}/rollback-stage' in html
    assert "Powód cofnięcia" in html
    assert "Wyślij powiadomienie do użytkownika" in html


def test_stage_rollback_is_hidden_and_endpoint_returns_403_for_form_manager(admin_app, admin_client):
    manager_id = create_user(admin_app, email="rollback-manager@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="rollback_form", name="Rollback Form", user_id=manager_id)
    submission_pk, submission_id = create_rollback_submission(admin_app)
    login(admin_client, email="rollback-manager@example.com")

    detail = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}")
    response = admin_client.post(
        f"/admin/submissions/{submission_id}/rollback-stage",
        data={
            "csrf_token": admin_csrf(admin_client),
            "target_status": "DECLARATION_SIGNED",
            "rollback_reason": "Nie powinno się udać",
        },
    )

    assert "Cofnij etap" not in detail.get_data(as_text=True)
    assert response.status_code == 403


def test_stage_rollback_endpoint_updates_status_history_instruction_and_sends_mail(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    instruction_config = {
        "stages": [
            {
                "key": "review-accepted",
                "label": "Wniosek zaakceptowany",
                "status_codes": ["OFFICER_ACCEPTED"],
                "next_action": "Uzupełnij ponownie wymagane dokumenty.",
            }
        ]
    }
    form_id = create_form(
        admin_app,
        slug="rollback_form",
        name="Rollback Form",
        user_instruction_config=instruction_config,
    )
    submission_pk, submission_id = create_rollback_submission(admin_app)
    login(admin_client)
    captured = {}

    def fake_dispatch(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(status="sent", recipient=kwargs["to_email"], error_message="")

    monkeypatch.setattr(
        admin_app.extensions["services"].mail_dispatch_service,
        "dispatch_to_submission",
        fake_dispatch,
    )

    response = admin_client.post(
        f"/admin/submissions/{submission_id}/rollback-stage",
        data={
            "csrf_token": admin_csrf(admin_client),
            "target_status": "OFFICER_ACCEPTED",
            "rollback_reason": "Ponowna kontrola danych",
            "send_notification": "on",
        },
    )

    assert response.status_code == 302
    assert response.location.endswith(f"/admin/forms/{form_id}/submissions/{submission_pk}")
    assert captured["event_type"] == "stage_rollback"
    assert captured["extra_context"]["rollback_reason"] == "Ponowna kontrola danych"
    assert captured["extra_context"]["rollback_next_action"] == "Uzupełnij ponownie wymagane dokumenty."

    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        event = db.query(SubmissionWorkflowEvent).filter_by(public_submission_id=submission_id).one()
        assert submission.process_status == "OFFICER_ACCEPTED"
        assert submission.declaration_filename == ""
        assert submission.agreement_filename == ""
        assert event.submission_id == submission_pk
        assert event.previous_status == "AGREEMENT_WAITING_FOR_SIGNATURE"
        assert event.new_status == "OFFICER_ACCEPTED"
        assert event.source == "stage_rollback"

    status_response = admin_client.get(f"/api/submissions/{submission_id}/workflow-status")
    payload = status_response.get_json()
    assert payload["process_status"] == "OFFICER_ACCEPTED"
    assert payload["current_step"] == "declaration"


def test_stage_rollback_without_email_address_does_not_fail(admin_app, admin_client):
    create_user(admin_app)
    create_form(admin_app, slug="rollback_form", name="Rollback Form")
    _, submission_id = create_rollback_submission(admin_app, email="")
    login(admin_client)

    response = admin_client.post(
        f"/admin/submissions/{submission_id}/rollback-stage",
        data={
            "csrf_token": admin_csrf(admin_client),
            "target_status": "DECLARATION_SIGNED",
            "rollback_reason": "Ponowienie podpisu",
            "send_notification": "on",
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "nie ma adresu e-mail" in response.get_data(as_text=True)


def test_stage_rollback_endpoint_requires_reason_and_keeps_status(admin_app, admin_client):
    create_user(admin_app)
    create_form(admin_app, slug="rollback_form", name="Rollback Form")
    submission_pk, submission_id = create_rollback_submission(admin_app)
    login(admin_client)

    response = admin_client.post(
        f"/admin/submissions/{submission_id}/rollback-stage",
        data={
            "csrf_token": admin_csrf(admin_client),
            "target_status": "DECLARATION_SIGNED",
            "rollback_reason": "",
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "Powód cofnięcia jest wymagany" in response.get_data(as_text=True)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(FormSubmission, submission_pk).process_status == "AGREEMENT_WAITING_FOR_SIGNATURE"
        assert db.query(SubmissionWorkflowEvent).count() == 0


def test_submission_detail_rejects_manual_process_status_bypass(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="rollback_form", name="Rollback Form")
    submission_pk, _ = create_rollback_submission(admin_app)
    login(admin_client)

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}",
        data={
            "csrf_token": admin_csrf(admin_client),
            "process_status": "FORM_SUBMITTED",
            "officer_decision": "accepted",
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "Statusu nie można zmieniać ręcznie" in response.get_data(as_text=True)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(FormSubmission, submission_pk).process_status == "AGREEMENT_WAITING_FOR_SIGNATURE"
        assert db.query(SubmissionWorkflowEvent).count() == 0


def test_return_for_correction_action_is_visible_and_uses_application_prefix(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="rollback_form", name="Rollback Form")
    submission_pk, submission_id = create_rollback_submission(admin_app)
    login(admin_client)

    response = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{submission_pk}",
        environ_overrides={"SCRIPT_NAME": "/aplikacja"},
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Wyślij do poprawy" in html
    assert f'/aplikacja/admin/submissions/{submission_id}/return-for-correction' in html
    assert 'name="clear_submission" checked' in html
    assert 'name="send_email" checked' in html
    assert 'class="admin-modal admin-modal-backdrop"' in html
    assert 'role="dialog" aria-modal="true" aria-labelledby="return-for-correction-title"' in html
    assert 'class="admin-modal__header"' in html
    assert 'class="admin-modal__body"' in html
    assert 'class="admin-modal__checks"' in html
    assert 'class="admin-modal__footer"' in html
    assert 'name="reason" rows="3" required maxlength="5000" autofocus' in html
    assert 'document.body.classList.toggle("admin-modal-open"' in html
    assert 'if (event.key === "Escape")' in html
    assert 'dialog.returnFocusTarget?.focus()' in html


def test_return_for_correction_endpoint_clears_state_and_sends_email(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="rollback_form",
        name="Rollback Form",
        definition_json={"title": "Rollback Form", "fields": [{"name": "email"}]},
    )
    submission_pk, submission_id = create_rollback_submission(admin_app)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(SubmissionFile(
            submission_id=submission_pk,
            public_submission_id=submission_id,
            form_slug="rollback_form",
            document_id="training_agreement",
            document_type="training_agreement",
            filename="agreement.pdf",
            storage_path="output/rollback_form/pdf/agreement.pdf",
            status="generated",
        ))
        db.commit()
    login(admin_client)
    captured = {}

    def fake_correction_mail(public_id, **kwargs):
        captured.update({"submission_id": public_id, **kwargs})
        return SimpleNamespace(status="sent", error_message="")

    monkeypatch.setattr(
        admin_app.extensions["services"].mail_dispatch_service,
        "dispatch_returned_for_correction",
        fake_correction_mail,
    )
    response = admin_client.post(
        f"/admin/submissions/{submission_id}/return-for-correction",
        data={
            "csrf_token": admin_csrf(admin_client),
            "reason": "Dane wymagają ponownego podania",
            "message_to_user": "Uzupełnij formularz ponownie.",
            "clear_submission": "on",
            "send_email": "on",
        },
    )

    assert response.status_code == 302
    assert response.location.endswith(f"/admin/forms/{form_id}/submissions/{submission_pk}")
    assert captured["submission_id"] == submission_id
    assert captured["recipient"] == "participant@example.com"
    assert captured["cleared"] is True
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.process_status == "RETURNED_FOR_CORRECTION"
        assert submission.email == ""
        assert submission.agreement_filename == ""
        assert db.query(SubmissionFile).one().status == "superseded"
        event = db.query(SubmissionWorkflowEvent).filter_by(source="returned_for_correction").one()
        assert event.previous_status == "AGREEMENT_WAITING_FOR_SIGNATURE"
        assert event.new_status == "RETURNED_FOR_CORRECTION"


def test_return_for_correction_requires_reason_and_is_forbidden_for_form_manager(admin_app, admin_client):
    admin_id = create_user(admin_app, email="correction-admin@example.com", role="admin")
    form_id = create_form(admin_app, slug="rollback_form", name="Rollback Form", user_id=admin_id)
    submission_pk, submission_id = create_rollback_submission(admin_app)
    login(admin_client, email="correction-admin@example.com")

    response = admin_client.post(
        f"/admin/submissions/{submission_id}/return-for-correction",
        data={"csrf_token": admin_csrf(admin_client), "reason": ""},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "Podaj powód" in response.get_data(as_text=True)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(FormSubmission, submission_pk).process_status == "AGREEMENT_WAITING_FOR_SIGNATURE"

    manager_id = create_user(admin_app, email="correction-manager@example.com", role="form_manager")
    with session_factory() as db:
        db.add(FormPermission(user_id=manager_id, form_id=form_id, can_manage=True))
        db.commit()
    admin_client.get("/admin/logout")
    login(admin_client, email="correction-manager@example.com")
    forbidden = admin_client.post(
        f"/admin/submissions/{submission_id}/return-for-correction",
        data={"csrf_token": admin_csrf(admin_client), "reason": "Próba"},
    )
    assert forbidden.status_code == 403


def test_blocked_agreement_section_and_actions_follow_admin_permissions(admin_app, admin_client):
    super_id = create_user(admin_app)
    form_id = create_form(admin_app, slug="blocked_form", name="Blocked Form", user_id=super_id)
    submission_pk, _ = create_blocked_agreement_submission(admin_app)
    login(admin_client)

    super_html = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{submission_pk}"
    ).get_data(as_text=True)

    assert "Umowa zablokowana" in super_html
    assert "Warunki deklaracji nie zostały spełnione." in super_html
    assert "Wyślij do poprawy" in super_html
    assert "Cofnij etap" in super_html
    assert "Odblokuj umowę" in super_html
    assert "Zakończ jako odrzucone" in super_html

    admin_id = create_user(admin_app, email="blocked-admin@example.com", role="admin")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormPermission(user_id=admin_id, form_id=form_id, can_manage=True))
        db.commit()
    admin_client.get("/admin/logout")
    login(admin_client, email="blocked-admin@example.com")
    admin_html = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{submission_pk}"
    ).get_data(as_text=True)

    assert "Wyślij do poprawy" in admin_html
    assert "Cofnij etap" in admin_html
    assert "Zakończ jako odrzucone" in admin_html
    assert 'data-admin-modal-target="unblock-agreement-dialog"' not in admin_html


def test_unblock_agreement_requires_csrf_and_super_admin(admin_app, admin_client):
    super_id = create_user(admin_app)
    form_id = create_form(admin_app, slug="blocked_form", name="Blocked Form", user_id=super_id)
    _, submission_id = create_blocked_agreement_submission(admin_app)
    login(admin_client)

    missing_csrf = admin_client.post(
        f"/admin/submissions/{submission_id}/agreement/unblock",
        data={"reason": "Decyzja po dodatkowej weryfikacji"},
    )
    assert missing_csrf.status_code == 400

    admin_id = create_user(admin_app, email="no-unblock@example.com", role="admin")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormPermission(user_id=admin_id, form_id=form_id, can_manage=True))
        db.commit()
    admin_client.get("/admin/logout")
    login(admin_client, email="no-unblock@example.com")
    forbidden = admin_client.post(
        f"/admin/submissions/{submission_id}/agreement/unblock",
        data={"csrf_token": admin_csrf(admin_client), "reason": "Próba"},
    )
    assert forbidden.status_code == 403


def test_super_admin_unblocks_agreement_and_preserves_audit_history(admin_app, admin_client, monkeypatch):
    super_id = create_user(admin_app)
    form_id = create_form(admin_app, slug="blocked_form", name="Blocked Form", user_id=super_id)
    submission_pk, submission_id = create_blocked_agreement_submission(admin_app)
    login(admin_client)
    captured = {}

    monkeypatch.setattr(
        admin_app.extensions["services"].audit_log_service,
        "log_event",
        lambda *args, **kwargs: captured.update({"args": args, "kwargs": kwargs}),
    )
    response = admin_client.post(
        f"/admin/submissions/{submission_id}/agreement/unblock",
        data={
            "csrf_token": admin_csrf(admin_client),
            "reason": "Potwierdzono wyjątek po dodatkowej weryfikacji.",
        },
    )

    assert response.status_code == 302
    assert response.location.endswith(f"/admin/forms/{form_id}/submissions/{submission_pk}")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.process_status == "AGREEMENT_READY"
        assert submission.agreement_blocked == ""
        assert submission.agreement_block_reason == ""
        history = submission.data_json["_agreement_block_history"]
        assert history[-1]["original_block_reason"] == "Warunki deklaracji nie zostały spełnione."
        assert history[-1]["actor_role"] == "super_admin"
        event = db.query(SubmissionWorkflowEvent).filter_by(source="manual_agreement_unblock").one()
        assert event.previous_status == "AGREEMENT_BLOCKED"
        assert event.new_status == "AGREEMENT_READY"
        assert event.reason == "Potwierdzono wyjątek po dodatkowej weryfikacji."
    assert captured["args"][0] == "AGREEMENT_UNBLOCKED"

    public_status = admin_client.get(
        f"/api/submissions/{submission_id}/acceptance-status"
    ).get_json()
    assert public_status["process_status"] == "AGREEMENT_READY"
    assert public_status["agreement_blocked"] is False
    assert public_status["can_generate_agreement"] is True


def test_blocked_agreement_can_be_rolled_back_by_admin(admin_app, admin_client):
    admin_id = create_user(admin_app, email="blocked-rollback@example.com", role="admin")
    form_id = create_form(
        admin_app,
        slug="blocked_form",
        name="Blocked Form",
        user_id=admin_id,
    )
    submission_pk, submission_id = create_blocked_agreement_submission(admin_app)
    login(admin_client, email="blocked-rollback@example.com")

    response = admin_client.post(
        f"/admin/submissions/{submission_id}/rollback-stage",
        data={
            "csrf_token": admin_csrf(admin_client),
            "target_status": "DECLARATION_SIGNED",
            "rollback_reason": "Ponowna weryfikacja deklaracji",
        },
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.process_status == "DECLARATION_SIGNED"
        assert submission.agreement_blocked == ""
        event = db.query(SubmissionWorkflowEvent).filter_by(source="stage_rollback").one()
        assert event.previous_status == "AGREEMENT_BLOCKED"


def test_admin_rejects_blocked_agreement_with_reason_and_public_status(admin_app, admin_client):
    admin_id = create_user(admin_app, email="blocked-reject@example.com", role="admin")
    form_id = create_form(
        admin_app,
        slug="blocked_form",
        name="Blocked Form",
        user_id=admin_id,
    )
    submission_pk, submission_id = create_blocked_agreement_submission(admin_app)
    login(admin_client, email="blocked-reject@example.com")

    response = admin_client.post(
        f"/admin/submissions/{submission_id}/reject-final",
        data={
            "csrf_token": admin_csrf(admin_client),
            "reason": "Warunki udziału nie zostały spełnione.",
        },
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.process_status == "OFFICER_REJECTED"
        assert submission.officer_decision == "rejected"
        assert submission.officer_decision_reason == "Warunki udziału nie zostały spełnione."
        assert submission.agreement_blocked == ""
        decision = db.query(SubmissionDecision).filter_by(target_status="OFFICER_REJECTED").one()
        assert decision.justification == "Warunki udziału nie zostały spełnione."
        event = db.query(SubmissionWorkflowEvent).filter_by(source="agreement_block_final_rejection").one()
        assert event.reason == decision.justification

    public_status = admin_client.get(
        f"/api/submissions/{submission_id}/acceptance-status"
    ).get_json()
    assert public_status["process_status"] == "OFFICER_REJECTED"
    assert public_status["status_reason"] == "Warunki udziału nie zostały spełnione."
    assert public_status["can_sign_documents"] is False
    assert public_status["can_generate_agreement"] is False


def test_participant_can_refill_returned_submission_and_conditions_are_evaluated_again(
    admin_app, admin_client, form_definition, valid_form_data, monkeypatch, tmp_path
):
    definition = dict(form_definition)
    definition["qualification_conditions"] = {
        "enabled": True,
        "conditions": [{
            "id": "age-rule",
            "field_name": "wiek",
            "field_label": "Wiek",
            "operator": "greater_than_or_equal",
            "expected_value": "99",
            "failure_action": "auto_reject",
            "user_message": "Nie spełniasz warunku wieku.",
            "officer_message": "Wiek poniżej wymaganego progu.",
            "is_active": True,
        }],
    }
    form_id = create_form(
        admin_app,
        slug="correction_form",
        name="Correction Form",
        definition_json=definition,
        is_active=True,
        is_public=True,
    )
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        submission = FormSubmission(
            submission_id="correction-public-uuid",
            form_slug="correction_form",
            form_name="Correction Form",
            access_token="correction-secret",
            process_status="RETURNED_FOR_CORRECTION",
            workflow_step="returned_for_correction",
            correction_required="Tak",
            correction_message="Uzupełnij dane ponownie.",
            data_json={
                "_qualification": {"passed": False, "evaluated_at": "2026-01-01T00:00:00+00:00"},
                "_qualification_history": [{"passed": False, "evaluated_at": "2026-01-01T00:00:00+00:00"}],
                "_correction_history": [{"reason": "Ponowna próba"}],
            },
        )
        db.add(submission)
        db.commit()

    def fake_generate_pdf(*, output_path, **_kwargs):
        Path(output_path).write_bytes(b"%PDF-1.4\n% correction test\n")

    monkeypatch.setattr("services.submission_service.generate_pdf", fake_generate_pdf)
    path = "/form/correction_form/correction/correction-public-uuid?token=correction-secret"
    get_response = admin_client.get(path, environ_overrides={"SCRIPT_NAME": "/aplikacja"})

    assert get_response.status_code == 200
    get_html = get_response.get_data(as_text=True)
    assert "Zgłoszenie zostało wysłane do poprawy" in get_html
    assert "/aplikacja/form/correction_form/correction/correction-public-uuid?token=correction-secret" in get_html

    post_response = admin_client.post(path, data=valid_form_data)
    assert post_response.status_code == 200
    assert "nie spełnia" in post_response.get_data(as_text=True).lower()
    with session_factory() as db:
        submission = db.query(FormSubmission).filter_by(submission_id="correction-public-uuid").one()
        assert submission.process_status == "AUTO_REJECTED"
        assert submission.data_json["_qualification"]["passed"] is False
        assert len(submission.data_json["_qualification_history"]) == 2
        assert submission.data_json["_correction_history"][0]["reason"] == "Ponowna próba"
        events = db.query(SubmissionWorkflowEvent).filter_by(public_submission_id=submission.submission_id).all()
        assert any(event.source == "auto_rejected_by_condition" for event in events)


def test_admin_requires_login(admin_client):
    response = admin_client.get("/admin/dashboard")

    assert response.status_code == 302
    assert "/admin/" in response.location


def test_admin_login(admin_app, admin_client):
    create_user(admin_app)

    response = login(admin_client)

    assert response.status_code == 302
    assert response.location.endswith("/admin/dashboard")


def test_super_admin_sees_all_forms(admin_app, admin_client):
    create_user(admin_app)
    create_form(admin_app, slug="one", name="One")
    create_form(admin_app, slug="two", name="Two")
    login(admin_client)

    response = admin_client.get("/admin/forms")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "One" in html
    assert "Two" in html


def test_form_manager_sees_only_assigned_forms(admin_app, admin_client):
    manager_id = create_user(admin_app, email="manager@example.com", role="form_manager")
    create_form(admin_app, slug="owned", name="Owned", user_id=manager_id)
    create_form(admin_app, slug="hidden", name="Hidden")
    login(admin_client, email="manager@example.com")

    response = admin_client.get("/admin/forms")
    html = response.get_data(as_text=True)

    assert "Owned" in html
    assert "Hidden" not in html


def test_delete_button_visible_only_for_super_admin(admin_app, admin_client):
    create_user(admin_app)
    create_form(admin_app, slug="one", name="One")
    login(admin_client)

    response = admin_client.get("/admin/forms")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Usuń" in html
    assert "/admin/forms/1/delete" in html


def test_delete_button_hidden_for_regular_admin(admin_app, admin_client):
    admin_id = create_user(admin_app, email="regular@example.com", role="admin")
    create_form(admin_app, slug="owned", name="Owned", user_id=admin_id)
    login(admin_client, email="regular@example.com")

    response = admin_client.get("/admin/forms")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Usuń" not in html


def test_regular_admin_cannot_delete_form(admin_app, admin_client):
    admin_id = create_user(admin_app, email="regular@example.com", role="admin")
    form_id = create_form(admin_app, slug="owned", name="Owned", user_id=admin_id)
    login(admin_client, email="regular@example.com")
    token = admin_client.get("/admin/forms").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(f"/admin/forms/{form_id}/delete", data={"csrf_token": token})

    assert response.status_code == 403


def test_workflow_status_tile_has_no_border():
    stylesheet = Path("static/css/documents_to_sign.css").read_text(encoding="utf-8")
    status_block = stylesheet.split(".status-tile {", 1)[1].split("}", 1)[0]

    assert "border: 0;" in status_block
    assert "border: 1px solid var(--border)" not in status_block

def test_super_admin_delete_removes_form_from_database(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="delete_me", name="Delete me", is_active=True, is_public=True)
    login(admin_client)
    token = admin_client.get("/admin/forms").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(f"/admin/forms/{form_id}/delete", data={"csrf_token": token})

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(Form, form_id) is None
    list_html = admin_client.get("/admin/forms").get_data(as_text=True)
    assert "Delete me" not in list_html
    assert admin_client.get("/form/delete_me").status_code == 404


def test_super_admin_delete_blocks_form_with_submissions(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="with_submission", name="With submission")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormSubmission(submission_id="abc", form_slug="with_submission", form_name="With submission"))
        db.commit()
    login(admin_client)
    token = admin_client.get("/admin/forms").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(f"/admin/forms/{form_id}/delete", data={"csrf_token": token})

    assert response.status_code == 302
    with session_factory() as db:
        assert db.get(Form, form_id) is not None


def test_logo_upload_rejects_invalid_image_content(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    token = admin_client.get("/admin/logos").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/logos",
        data={
            "csrf_token": token,
            "name": "Bad logo",
            "logo_file": (io.BytesIO(b"not an image"), "logo.png"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.query(Logo).count() == 0


def test_public_forms_use_database_visibility_and_label(admin_app, admin_client):
    create_form(admin_app, slug="public_form", name="Public Form", label_text="PROJEKT 6.08", sort_order=1)
    create_form(admin_app, slug="hidden_form", name="Hidden Form", is_public=False, sort_order=2)
    create_form(admin_app, slug="inactive_form", name="Inactive Form", is_active=False, sort_order=3)

    response = admin_client.get("/")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Public Form" in html
    assert "PROJEKT 6.08" in html
    assert "Hidden Form" not in html
    assert "Inactive Form" not in html


def test_public_form_page_uses_active_database_fields(admin_app, admin_client):
    form_id = create_form(admin_app, slug="db_form", name="Database Form")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add_all(
            [
                FormField(form_id=form_id, name="email", label="E-mail", type="email", active=True),
                FormField(form_id=form_id, name="archived", label="Archived", type="text", active=False),
            ]
        )
        db.commit()

    response = admin_client.get("/form/db_form")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'name="email"' in html
    assert 'name="archived"' not in html


def test_public_database_form_renders_selected_logo_under_title_with_alignment(admin_app, tmp_path):
    logo_path = tmp_path / "logo.png"
    logo_path.write_bytes(b"\x89PNG\r\n\x1a\n")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        logo = Logo(
            name="Admin logo",
            filename="admin-logo.png",
            storage_path=str(logo_path),
            mime_type="image/png",
            active=True,
        )
        db.add(logo)
        db.flush()
        form = Form(
            slug="logo_form",
            name="Logo Form",
            title="Logo Form",
            description="Opis formularza",
            definition_json={"title": "Logo Form", "fields": [], "header_image": "Logo/static-big-logo.png"},
            is_active=True,
            is_public=True,
            logo_id=logo.id,
            logo_alignment="center",
        )
        db.add(form)
        db.commit()

    response = admin_app.test_client().get("/form/logo_form")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert html.index("<h2>Logo Form</h2>") < html.index("/assets/logos/")
    assert "form-logo-row form-logo-row--center" in html
    assert "admin-logo.png" in html
    assert "form-header-image" not in html
    assert "Logo/static-big-logo.png" not in html


def test_form_edit_saves_logo_alignment(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="alignment_form", name="Alignment Form")
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Alignment Form",
            "slug": "alignment_form",
            "title": "Alignment Form",
            "sort_order": "0",
            "workflow_json": "{}",
            "logo_alignment": "right",
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(Form, form_id).logo_alignment == "right"


def test_upload_form_detects_fields(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    payload = {
        "title": "Uploaded",
        "fields": [
            {"type": "section", "label": "Dane"},
            {"type": "text", "name": "imiona", "label": "Imiona", "required": True},
            {"type": "select", "name": "status", "label": "Status", "options": ["A", "B"]},
        ],
    }
    token = admin_client.get("/admin/forms/upload").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/forms/upload",
        data={"csrf_token": token, "slug": "uploaded", "form_file": (io.BytesIO(json.dumps(payload).encode()), "uploaded.json")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.query(Form).filter_by(slug="uploaded").one()
        assert [field.name for field in form.fields] == ["imiona", "status"]


def test_upload_form_preserves_polish_characters_and_workflow(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    payload = {
        "title": "Zażółć gęślą jaźń",
        "fields": [{"type": "text", "name": "imiona", "label": "Imię i nazwisko"}],
        "workflow": {
            "name": "Ścieżka",
            "initial_step": "submission",
            "steps": [{"id": "submission", "type": "end", "triggers": ["application_submitted"]}],
        },
    }
    token = admin_client.get("/admin/forms/upload").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/forms/upload",
        data={
            "csrf_token": token,
            "slug": "polskie",
            "form_file": (io.BytesIO(json.dumps(payload, ensure_ascii=False).encode("utf-8")), "polskie.json"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.query(Form).filter_by(slug="polskie").one()
        assert form.title == "Zażółć gęślą jaźń"
        assert form.definition_json["fields"][0]["label"] == "Imię i nazwisko"
        assert form.definition_json["workflow"]["name"] == "Ścieżka"


def test_upload_html_form_redirects_to_fields(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html_form = b'<form><input name="email" type="email" required><textarea name="opis"></textarea></form>'
    token = admin_client.get("/admin/forms/upload").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/forms/upload",
        data={
            "csrf_token": token,
            "slug": "html_form",
            "name": "HTML Form",
            "title": "Publiczny formularz HTML",
            "description": "Opis z kreatora",
            "label_text": "PROJEKT TEST",
            "form_file": (io.BytesIO(html_form), "html_form.html"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    assert "/fields" in response.location
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.query(Form).filter_by(slug="html_form").one()
        assert [field.name for field in form.fields] == ["email", "opis"]
        assert all(field.required for field in form.fields)
        assert form.name == "HTML Form"
        assert form.title == "Publiczny formularz HTML"
        assert form.description == "Opis z kreatora"
        assert form.label_text == "PROJEKT TEST"


def test_upload_form_saves_basic_mail_and_appearance_settings_for_public_views(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    payload = {
        "title": "Tytuł z pliku",
        "fields": [{"type": "text", "name": "imiona", "label": "Imiona"}],
    }
    token = admin_client.get("/admin/forms/upload").get_data(as_text=True).split(
        'name="csrf_token" value="', 1
    )[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/forms/upload",
        data={
            "csrf_token": token,
            "name": "Nazwa administratora",
            "title": "Tytuł publiczny administratora",
            "description": "Opis publiczny administratora",
            "slug": "ustawienia_importu",
            "sort_order": "17",
            "label_text": "PROJEKT IMPORT",
            "label_color": "#123456",
            "label_background": "#abcdef",
            "logo_alignment": "right",
            "mail_mode": "custom",
            "form_smtp_host": "smtp.example.test",
            "form_smtp_port": "465",
            "form_smtp_mail_from": "formularz@example.test",
            "form_smtp_sender_name": "Formularz testowy",
            "form_smtp_reply_to": "odpowiedz@example.test",
            "form_smtp_use_ssl": "on",
            "is_active": "on",
            "is_public": "on",
            "form_file": (io.BytesIO(json.dumps(payload).encode()), "ustawienia.json"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.query(Form).filter_by(slug="ustawienia_importu").one()
        assert form.name == "Nazwa administratora"
        assert form.title == "Tytuł publiczny administratora"
        assert form.description == "Opis publiczny administratora"
        assert form.sort_order == 17
        assert form.label_text == "PROJEKT IMPORT"
        assert form.label_color == "#123456"
        assert form.label_background == "#abcdef"
        assert form.logo_alignment == "right"
        assert form.mail_mode == "custom"
        assert form.smtp_config == {
            "host": "smtp.example.test",
            "port": 465,
            "user": "",
            "mail_from": "formularz@example.test",
            "sender_name": "Formularz testowy",
            "use_tls": False,
            "use_ssl": True,
            "timeout": 10,
            "reply_to": "odpowiedz@example.test",
        }
        version_service = admin_app.extensions["services"].form_version_service
        draft = version_service.editable_draft(db, form.id)
        assert draft is not None
        version_service.publish(
            db,
            form,
            draft,
            actor_id=None,
            change_summary="Publikacja po imporcie",
        )
        db.commit()

    index_html = admin_client.get("/").get_data(as_text=True)
    public_form_html = admin_client.get("/form/ustawienia_importu").get_data(as_text=True)
    assert "Tytuł publiczny administratora" in index_html
    assert "PROJEKT IMPORT" in index_html
    assert "color: #123456; background: #abcdef;" in index_html
    assert "Tytuł publiczny administratora" in public_form_html
    assert "Opis publiczny administratora" in public_form_html
    assert "color: #123456; background: #abcdef;" in public_form_html


def test_imported_fields_default_to_required_and_preserve_explicit_false():
    definition = normalize_form_definition(
        {
            "title": "Wymagalnosc",
            "fields": [
                {"type": "section", "label": "Sekcja"},
                {"type": "text", "name": "default_required"},
                {"type": "email", "name": "optional_email", "required": False},
                {"type": "text", "name": "system_value", "system": True},
            ],
        }
    )

    assert definition["fields"][0]["required"] is False
    assert definition["fields"][1]["required"] is True
    assert definition["fields"][2]["required"] is False
    assert definition["fields"][3]["required"] is False


def test_html_import_reports_missing_names_duplicates_and_unsupported_types():
    with pytest.raises(ValueError) as exc_info:
        build_definition_from_html(
            '<form><input type="text"><input name="email"><input name="email"><input name="avatar" type="file"></form>',
            "bledny.html",
        )

    message = str(exc_info.value)
    assert "nie ma atrybutu 'name'" in message
    assert "Duplikat pola HTML" in message
    assert "nieobsługiwany typ HTML 'file'" in message


def test_json_validation_reports_duplicate_field_names():
    definition = normalize_form_definition(
        {
            "title": "Duplikaty",
            "fields": [
                {"type": "text", "name": "email"},
                {"type": "email", "name": "email"},
            ],
        }
    )

    with pytest.raises(ValueError, match="Duplikat pola"):
        validate_form_definition(definition)


def test_super_admin_manages_import_instruction_and_download_is_protected(admin_app, admin_client):
    create_user(admin_app)
    create_user(admin_app, email="manager-instruction@example.com", role="form_manager")
    login(admin_client)
    documents_page = admin_client.get("/admin/site/documents")
    token = documents_page.get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/site/documents",
        data={
            "csrf_token": token,
            "document_type": "form_import_instruction",
            "title": "Instrukcja przygotowania plikow formularza",
            "document_file": (io.BytesIO(b"# Instrukcja\n\nPola formularza."), "instrukcja.md", "text/markdown"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "instrukcja.md" in response.get_data(as_text=True)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        instruction = db.query(ServiceDocument).filter_by(document_type="form_import_instruction").one()
        instruction_id = instruction.id
        stored_path = Path(instruction.storage_path)
        assert stored_path.exists()

    upload_page = admin_client.get("/admin/forms/upload").get_data(as_text=True)
    assert "Pobierz instrukcję przygotowania plików formularza" in upload_page
    assert admin_client.get(f"/admin/site/documents/{instruction_id}/file").status_code == 200

    manager_client = admin_app.test_client()
    login(manager_client, email="manager-instruction@example.com")
    assert manager_client.get(f"/admin/site/documents/{instruction_id}/file").status_code == 403

    delete_token = admin_client.get("/admin/site/documents").get_data(as_text=True).split(
        'name="csrf_token" value="', 1
    )[1].split('"', 1)[0]
    delete_response = admin_client.post(
        "/admin/site/documents/form-import-instruction/delete",
        data={"csrf_token": delete_token},
        follow_redirects=True,
    )
    assert delete_response.status_code == 200
    assert not stored_path.exists()
    with session_factory() as db:
        assert db.query(ServiceDocument).filter_by(document_type="form_import_instruction").one_or_none() is None


def test_workflow_can_be_edited_after_json_import(admin_app, admin_client):
    create_user(admin_app)
    definition = {
        "title": "Workflow Form",
        "fields": [{"type": "text", "name": "first_name", "label": "Imię"}],
        "workflow": {
            "name": "Import",
            "initial_step": "submission",
            "steps": [
                {"id": "submission", "type": "form_submit", "next": "completed", "triggers": ["application_submitted"]},
                {"id": "completed", "type": "end"},
            ],
        },
    }
    form_id = create_form(admin_app, slug="workflow_form", name="Workflow Form", definition_json=definition)
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Workflow Form",
            "slug": "workflow_form",
            "title": "Workflow Form",
            "sort_order": "0",
            "workflow_name": "Po imporcie",
            "workflow_initial_step": "submission",
            "requires_declaration": "on",
            "declaration_template_html": "<p>Deklaracja {{ first_name }}</p>",
            "workflow_json": json.dumps(definition["workflow"], ensure_ascii=False),
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.get(Form, form_id)
        workflow = form.definition_json["workflow"]
        assert workflow["name"] == "Po imporcie"
        assert workflow["requires_declaration"] is True
        assert workflow["declaration_template_html"] == "<p>Deklaracja {{ first_name }}</p>"
        declaration = next(document for document in form.definition_json["documents"] if document["id"] == "declaration")
        assert declaration["template_html"] == "<p>Deklaracja {{ first_name }}</p>"


def test_workflow_edit_shows_tooltips_and_conditional_html_fields(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="workflow_form",
        name="Workflow Form",
        definition_json={
            "title": "Workflow Form",
            "fields": [],
            "workflow": {
                "name": "Workflow",
                "initial_step": "submission",
                "requires_contract": True,
                "contract_template_html": "<p>Umowa</p>",
                "steps": [{"id": "submission", "type": "end", "triggers": ["application_submitted"]}],
            },
        },
    )
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/edit")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "application_submitted" in html
    assert "Uruchamiany po wysłaniu formularza przez użytkownika." in html
    assert "additional_fields_completed" in html
    assert "Szablon umowy HTML" in html
    assert "Szablon deklaracji HTML" in html
    assert 'data-workflow-template="contract"' in html
    assert 'name="contract_generation_mode"' in html
    assert 'name="contract_filename_pattern"' in html
    assert 'name="contract_number_pattern"' in html
    assert "Wspólny CSS dokumentów zostanie dołączony automatycznie." in html


def test_admin_submission_list_renders_workflow_status_label(admin_app, admin_client):
    user_id = create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="status_form",
        name="Status Form",
        user_id=user_id,
        definition_json={
            "title": "Status Form",
            "fields": [],
            "workflow": {
                "initial_step": "submission",
                "statuses": [{"id": "custom_status", "label": "Przyjazny status"}],
                "steps": [{"id": "submission", "type": "end"}],
            },
        },
    )
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormSubmission(submission_id="abc", form_slug="status_form", form_name="Status Form", process_status="custom_status"))
        db.commit()
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/submissions")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Przyjazny status" in html


def test_admin_submission_list_renders_unknown_status_fallback(admin_app, admin_client):
    user_id = create_user(admin_app)
    form_id = create_form(admin_app, slug="status_form", name="Status Form", user_id=user_id)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormSubmission(submission_id="abc", form_slug="status_form", form_name="Status Form", process_status="MISSING_STATUS"))
        db.commit()
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/submissions")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Nieznany status: MISSING_STATUS" in html


def test_form_fields_can_be_edited(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormField(form_id=form_id, name="email", label="E-mail", type="email", required=True, section="Kontakt", sort_order=1))
        db.commit()
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    assert "Dostępność w procesie" in html
    assert f'field_' in html and "_availability_0_visible" in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    with session_factory() as db:
        field = db.query(FormField).filter_by(form_id=form_id, name="email").one()
        field_id = field.id

    response = admin_client.post(
        f"/admin/forms/{form_id}/fields",
        data={
            "csrf_token": token,
            f"field_{field_id}_label": "Adres e-mail",
            f"field_{field_id}_document_label": "Adres e-mail uczestnika",
            f"field_{field_id}_type": "text",
            f"field_{field_id}_required": "on",
            f"field_{field_id}_section": "Dane kontaktowe",
            f"field_{field_id}_sort_order": "3",
        },
    )

    assert response.status_code == 302
    with session_factory() as db:
        field = db.query(FormField).filter_by(id=field_id).one()
        assert field.label == "Adres e-mail"
        assert field.type == "text"
        assert field.section == "Dane kontaktowe"
        assert field.sort_order == 3
        assert db.get(Form, form_id).definition_json["document_field_labels"]["email"] == "Adres e-mail uczestnika"
    reloaded = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    assert 'value="Adres e-mail uczestnika"' in reloaded


def test_visual_form_builder_saves_layout_and_reopens_it(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="visual_builder",
        name="Visual Builder",
        definition_json={
            "title": "Visual Builder",
            "fields": [
                {"name": "imie", "label": "Imię", "type": "text", "custom_setting": "zachowaj"},
                {"name": "nazwisko", "label": "Nazwisko", "type": "text"},
            ],
        },
    )
    factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with factory() as db:
        sync_form_fields(db, db.get(Form, form_id), db.get(Form, form_id).definition_json)
        db.commit()
        fields = {field.name: field for field in db.query(FormField).filter_by(form_id=form_id).all()}
        state = [
            {"id": fields["nazwisko"].id, "name": "nazwisko", "label": "Nazwisko kandydata", "type": "text", "required": True, "width": "half", "width_span": 6, "placeholder": "", "section": "Dane", "document_label": "", "options": []},
            {"id": fields["imie"].id, "name": "imie", "label": "Imię kandydata", "type": "text", "required": True, "width": "half", "width_span": 6, "placeholder": "Wpisz imię", "section": "Dane", "document_label": "", "options": []},
            {"id": None, "name": "email", "label": "Adres e-mail", "type": "email", "required": False, "width": "full", "width_span": 12, "placeholder": "name@example.org", "section": "Kontakt", "document_label": "", "options": []},
        ]
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    assert "Wizualny układ formularza" in html
    assert "data-form-canvas" in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        f"/admin/forms/{form_id}/fields",
        data={"csrf_token": token, "action": "builder_save", "builder_state": json.dumps(state)},
    )
    assert response.status_code == 302
    with factory() as db:
        form = db.get(Form, form_id)
        active = db.query(FormField).filter_by(form_id=form_id, active=True).order_by(FormField.sort_order).all()
        assert [field.name for field in active] == ["nazwisko", "imie", "email"]
        configs = {item["name"]: item for item in form.definition_json["fields"]}
        assert configs["nazwisko"]["width"] == configs["imie"]["width"] == "half"
        assert configs["imie"]["custom_setting"] == "zachowaj"
    reopened = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    assert '"width_span": 6' in reopened
    assert "Adres e-mail" in reopened


def test_form_field_availability_uses_configured_workflow_steps(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="availability_form", name="Availability")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        form = db.get(Form, form_id)
        form.definition_json = {
            **(form.definition_json or {}),
            "workflow": {
                "initial_step": "submission",
                "steps": [
                    {"id": "submission", "admin_label": "Złożenie", "next": "declaration"},
                    {"id": "declaration", "admin_label": "Deklaracja", "requires_user_action": True},
                ],
            },
        }
        field = FormField(form_id=form_id, name="account", label="Rachunek", type="text", required=False, sort_order=1)
        db.add(field)
        db.commit()
        field_id = field.id
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        f"/admin/forms/{form_id}/fields",
        data={
            "csrf_token": token,
            f"field_{field_id}_label": "Numer rachunku",
            f"field_{field_id}_type": "text",
            f"field_{field_id}_section": "Dane",
            f"field_{field_id}_sort_order": "1",
            f"field_{field_id}_availability_present": "1",
            f"field_{field_id}_availability_1_visible": "on",
            f"field_{field_id}_availability_1_editable": "on",
            f"field_{field_id}_availability_1_required": "on",
        },
    )
    assert response.status_code == 302
    with session_factory() as db:
        saved = db.get(FormField, field_id)
        assert saved.availability_json == [
            {"step": "submission", "visible": False, "editable": False, "required": False},
            {"step": "declaration", "visible": True, "editable": True, "required": True},
        ]


def test_form_field_can_be_added_and_deactivated_without_removing_history(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="history",
            form_slug="sample_form",
            form_name="Sample",
            data_json={"kurs": "Python"},
        )
        db.add(submission)
        db.commit()
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/fields",
        data={
            "csrf_token": token,
            "action": "add",
            "new_name": "kurs",
            "new_label": "Kurs",
            "new_type": "select",
            "new_options": "Python\nExcel",
        },
    )

    assert response.status_code == 302
    with session_factory() as db:
        field = db.query(FormField).filter_by(form_id=form_id, name="kurs").one()
        field_id = field.id
        assert field.active is True
        assert field.options == ["Python", "Excel"]

    html = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        f"/admin/forms/{form_id}/fields",
        data={"csrf_token": token, "action": f"delete:{field_id}"},
    )

    assert response.status_code == 302
    with session_factory() as db:
        field = db.get(FormField, field_id)
        submission = db.query(FormSubmission).filter_by(submission_id="history").one()
        assert field.active is False
        assert submission.data_json["kurs"] == "Python"


def test_submissions_filter_and_sort_dynamic_data(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add_all(
            [
                FormSubmission(submission_id="a", form_slug="sample_form", form_name="Sample", email="a@example.com", nazwisko="Zed", data_json={"kurs": "Excel"}),
                FormSubmission(submission_id="b", form_slug="sample_form", form_name="Sample", email="b@example.com", nazwisko="Ann", data_json={"kurs": "Python"}),
            ]
        )
        db.commit()
    login(admin_client)

    filtered = admin_client.get(f"/admin/forms/{form_id}/submissions?field=kurs&value=Python").get_data(as_text=True)
    sorted_html = admin_client.get(f"/admin/forms/{form_id}/submissions?sort=nazwisko&direction=asc").get_data(as_text=True)

    assert "b@example.com" in filtered
    assert "a@example.com" not in filtered
    assert sorted_html.index("Ann") < sorted_html.index("Zed")


def test_submissions_filter_supports_dynamic_operator(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add_all(
            [
                FormSubmission(submission_id="a", form_slug="sample_form", form_name="Sample", email="a@example.com", data_json={"kurs": "Excel"}),
                FormSubmission(submission_id="b", form_slug="sample_form", form_name="Sample", email="b@example.com", data_json={"kurs": "Python"}),
            ]
        )
        db.commit()
    login(admin_client)

    html = admin_client.get(
        f"/admin/forms/{form_id}/submissions?field=kurs&operator=equals&value=Python"
    ).get_data(as_text=True)

    assert "b@example.com" in html
    assert "a@example.com" not in html


def test_officer_decision_visible_and_quick_update(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(submission_id="abc", form_slug="sample_form", form_name="Sample", email="a@example.com")
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": token, "officer_decision": "accepted", "officer_decision_reason": "OK"},
    )

    assert response.status_code == 302
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.officer_decision == "accepted"
        assert submission.officer_decision_reason == ""
        decision = db.query(SubmissionDecision).filter_by(public_submission_id="abc").one()
        assert decision.decision == "accepted"
        assert decision.justification == ""
        assert decision.previous_status == "FORM_SUBMITTED"
        assert decision.target_status in {"OFFICER_ACCEPTED", "accepted_waiting_for_additional_fields"}
        assert decision.email_requested is False
        workflow_event = db.query(SubmissionWorkflowEvent).filter_by(public_submission_id="abc").one()
        assert workflow_event.new_status in {"REVIEW_ACCEPTED", "ACCEPTED_WAITING_FOR_ADDITIONAL_FIELDS"}
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    assert "Decyzja urzednika" in html
    assert "Wniosek zaakceptowany" in html
    assert f'name="officer_decision_{submission_pk}"' not in html


def test_form_manager_can_edit_arbitrary_instruction_stages(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="instruction_form", name="Instruction form")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(
            FormSubmission(
                submission_id="instruction-admin",
                form_slug="instruction_form",
                form_name="Instruction form",
                officer_decision="TAK",
                process_status="OFFICER_ACCEPTED",
            )
        )
        db.commit()
    login(admin_client)
    edit_html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = edit_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    assert "Instrukcja dalszego postępowania" in edit_html
    assert "Dodaj etap" in edit_html
    assert "Lista zawiera wyłącznie etapy użyte w aktualnym workflow" in edit_html
    assert "Podgląd instrukcji" in edit_html
    assert "OFFICER_ACCEPTED" in edit_html

    instruction_workflow = {
        "name": "Instrukcje",
        "initial_step": "submitted",
        "steps": [
            {
                "id": "submitted",
                "admin_label": "Etap pierwszy",
                "user_label": "Etap pierwszy",
                "status": "FORM_SUBMITTED",
                "description": "Opis pierwszego etapu",
                "next_action": "Poczekaj na kontakt.",
                "next": "accepted",
            },
            {
                "id": "accepted",
                "admin_label": "Etap zaakceptowany",
                "user_label": "Etap zaakceptowany",
                "status": "OFFICER_ACCEPTED",
                "description": "Opis drugiego etapu",
                "next_action": "Wykonaj własną czynność.",
                "final": True,
            },
        ],
    }

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Instruction form",
            "slug": "instruction_form",
            "title": "Instruction form",
            "sort_order": "0",
            "is_active": "on",
            "is_public": "on",
            "instruction_title": "Moja instrukcja",
            "user_instruction": "  Pierwszy krok.\nDrugi krok.  ",
            "workflow_builder_json": json.dumps(instruction_workflow, ensure_ascii=False),
            "user_instruction_config": json.dumps(
                {
                    "stages": [
                        {
                            "label": "Etap pierwszy",
                            "status_codes": ["FORM_SUBMITTED"],
                            "description": "Opis pierwszego etapu",
                            "next_action": "Poczekaj na kontakt.",
                        },
                        {
                            "label": "Etap zaakceptowany",
                            "status_codes": ["OFFICER_ACCEPTED", "REVIEW_ACCEPTED"],
                            "description": "Opis drugiego etapu",
                            "next_action": "Wykonaj własną czynność.",
                            "final": True,
                        },
                    ]
                }
            ),
        },
    )
    assert response.status_code == 302
    with session_factory() as db:
        saved_form = db.get(Form, form_id)
        assert saved_form.user_instruction == "Pierwszy krok.\nDrugi krok."
        assert saved_form.user_instruction_config["title"] == "Moja instrukcja"
        assert len(saved_form.user_instruction_config["stages"]) == 2
        assert saved_form.user_instruction_config["stages"][1]["status_codes"] == ["OFFICER_ACCEPTED"]
        assert "user_instruction" not in FormSubmission.__table__.columns
    payload = admin_client.get("/api/submissions/instruction-admin/acceptance-status").get_json()
    assert payload["form_instruction"] == "Pierwszy krok.\nDrugi krok."
    assert payload["has_form_instruction"] is True
    assert payload["instruction"]["title"] == "Moja instrukcja"
    assert payload["instruction"]["current_stage_label"] == "Etap zaakceptowany"
    assert payload["instruction"]["next_action"] == "Wykonaj własną czynność."

    edit_html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = edit_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Instruction form",
            "slug": "instruction_form",
            "title": "Instruction form",
            "sort_order": "0",
            "is_active": "on",
            "is_public": "on",
            "instruction_title": "",
            "user_instruction": "",
            "user_instruction_config": '{"stages": []}',
        },
    )
    assert response.status_code == 302
    with session_factory() as db:
        saved_form = db.get(Form, form_id)
        assert saved_form.user_instruction is None
        assert saved_form.user_instruction_config["stages"] == []


def test_officer_decision_does_not_send_automatic_mail(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            process_status="FORM_SUBMITTED",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    calls = []
    admin_app.extensions["services"].mail_dispatch_service.dispatch_decision_email = lambda submission_id, decision: calls.append((submission_id, decision))
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    first = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": token, "officer_decision": "accepted", "officer_decision_reason": "OK"},
    )
    second = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": token, "officer_decision": "accepted", "officer_decision_reason": "OK again"},
    )

    assert first.status_code == 302
    assert second.status_code == 302
    assert calls == []
    with session_factory() as db:
        decisions = db.query(SubmissionDecision).filter_by(public_submission_id="abc").order_by(SubmissionDecision.id).all()
        assert [decision.email_requested for decision in decisions] == [False]


def test_officer_decision_update_survives_missing_decision_audit_table(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            process_status="FORM_SUBMITTED",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
        db.execute(text("DROP TABLE submission_decisions"))
        db.commit()
    calls = []
    admin_app.extensions["services"].mail_dispatch_service.dispatch_decision_email = lambda submission_id, decision: calls.append((submission_id, decision))
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": token, "officer_decision": "rejected", "officer_decision_reason": "Braki"},
    )

    assert response.status_code == 302
    assert calls == []
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.officer_decision == "rejected"
        assert submission.officer_decision_reason == "Braki"
        assert submission.process_status == "OFFICER_REJECTED"


def test_bulk_officer_decision_update_saves_all_rows(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        first = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            process_status="FORM_SUBMITTED",
        )
        second = FormSubmission(
            submission_id="def",
            form_slug="sample_form",
            form_name="Sample",
            email="ewa@example.com",
            process_status="FORM_SUBMITTED",
        )
        db.add_all([first, second])
        db.commit()
        first_pk = first.id
        second_pk = second.id
    calls = []
    admin_app.extensions["services"].mail_dispatch_service.dispatch_decision_email = lambda submission_id, decision: calls.append((submission_id, decision))
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/decisions",
        data={
            "csrf_token": token,
            "submission_row_ids": [str(first_pk), str(second_pk)],
            f"officer_decision_{first_pk}": "accepted",
            f"officer_decision_reason_{first_pk}": "",
            f"officer_decision_{second_pk}": "correction",
            f"officer_decision_reason_{second_pk}": "Popraw PESEL",
        },
    )

    assert response.status_code == 302
    assert calls == []
    with session_factory() as db:
        first = db.get(FormSubmission, first_pk)
        second = db.get(FormSubmission, second_pk)
        assert first.officer_decision == "accepted"
        assert first.officer_decision_reason == ""
        assert second.officer_decision == "correction"
        assert second.officer_decision_reason == "Popraw PESEL"
        assert second.correction_required == "Tak"
        assert second.correction_message == "Popraw PESEL"
        assert second.process_status == "WAITING_FOR_CORRECTION"


def test_bulk_officer_decision_update_requires_reason_for_rejection_or_correction(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            process_status="FORM_SUBMITTED",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/decisions",
        data={
            "csrf_token": token,
            "submission_row_ids": [str(submission_pk)],
            f"officer_decision_{submission_pk}": "rejected",
            f"officer_decision_reason_{submission_pk}": "",
        },
        follow_redirects=True,
    )

    assert "Pominieto 1 decyzji" in response.get_data(as_text=True)
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.officer_decision == ""


def test_submission_detail_survives_missing_submission_file_alignment_columns(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            process_status="FORM_SUBMITTED",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
        db.execute(text("ALTER TABLE submission_files DROP COLUMN original_filename"))
        db.commit()
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}")

    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert "dokumentow" in html


def test_dashboard_survives_missing_submission_file_alignment_columns(admin_app, admin_client):
    create_user(admin_app)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.execute(text("ALTER TABLE submission_files DROP COLUMN original_filename"))
        db.commit()
    login(admin_client)

    response = admin_client.get("/admin/dashboard")

    assert response.status_code == 200
    assert "Dashboard" in response.get_data(as_text=True)


def test_dashboard_reads_email_statistics_and_recent_errors_from_email_logs(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="mail_stats", name="Mail stats")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add_all(
            [
                EmailLog(form_id=form_id, public_submission_id="sent-1", status="sent", to_email="ok@example.com"),
                EmailLog(
                    form_id=form_id,
                    public_submission_id="failed-1",
                    status="failed",
                    to_email="fail@example.com",
                    subject="Test",
                    error_message="Błąd testowy SMTP",
                ),
            ]
        )
        db.commit()
    login(admin_client)

    html = admin_client.get("/admin/dashboard").get_data(as_text=True)

    assert "Udane wysyłki e-mail" in html
    assert "Nieudane wysyłki e-mail" in html
    assert "Ostatnia próba wysyłki" in html
    assert "Błąd testowy SMTP" in html


def test_smtp_test_failure_is_safely_logged_and_saved_in_email_logs(
    admin_app, admin_client, monkeypatch, caplog
):
    create_user(admin_app)
    service = admin_app.extensions["services"].mail_settings_service
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        settings = SystemMailSettings(
            smtp_config={
                "host": "smtp.failure.test",
                "port": 2525,
                "user": "smtp-user",
                "mail_from": "sender@example.com",
                "use_tls": True,
                "use_ssl": False,
                "timeout": 7,
            },
            smtp_password_encrypted=service.encrypt_password("smtp-super-secret"),
            layout_config={},
        )
        db.add(settings)
        db.commit()
    login(admin_client)

    def raise_timeout(config):
        assert config["smtp_password"] == "smtp-super-secret"
        raise TimeoutError("smtp-super-secret must never be logged")

    monkeypatch.setattr(service, "test_connection", raise_timeout)
    with caplog.at_level(logging.WARNING):
        response = admin_client.post(
            "/admin/mail-settings/test",
            data={"csrf_token": admin_csrf(admin_client)},
            follow_redirects=True,
        )

    html = response.get_data(as_text=True)
    assert "Nie udało się połączyć z serwerem SMTP w wyznaczonym czasie." in html
    assert "host=smtp.failure.test" in caplog.text
    assert "port=2525" in caplog.text
    assert "error_type=TimeoutError" in caplog.text
    assert "smtp-super-secret" not in caplog.text
    assert "smtp-user" not in caplog.text

    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        log = db.query(EmailLog).filter_by(event_type="smtp_test").one()
        assert log.status == "failed"
        assert log.error_type == "TimeoutError"
        assert log.administrator_message.startswith(
            "Nie udało się połączyć z serwerem SMTP w wyznaczonym czasie."
        )
        assert "smtp-super-secret" not in log.error_message
        assert log.created_at is not None


def test_successful_smtp_test_is_saved_and_visible_on_dashboard(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    service = admin_app.extensions["services"].mail_settings_service
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            SystemMailSettings(
                smtp_config={
                    "host": "smtp.success.test",
                    "port": 587,
                    "mail_from": "sender@example.com",
                    "timeout": 10,
                },
                layout_config={},
            )
        )
        db.commit()
    login(admin_client)
    monkeypatch.setattr(service, "test_connection", lambda config: None)

    admin_client.post(
        "/admin/mail-settings/test",
        data={"csrf_token": admin_csrf(admin_client)},
    )
    html = admin_client.get("/admin/dashboard").get_data(as_text=True)

    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        log = db.query(EmailLog).filter_by(event_type="smtp_test").one()
        assert log.status == "sent"
        assert log.error_type == ""
        assert log.created_at is not None
    assert "Udane testy SMTP" in html
    assert "Ostatnia próba SMTP" in html


def test_successful_smtp_test_reports_log_failure_without_claiming_smtp_failure(
    admin_app, admin_client, monkeypatch
):
    create_user(admin_app)
    service = admin_app.extensions["services"].mail_settings_service
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            SystemMailSettings(
                smtp_config={
                    "host": "smtp.success.test",
                    "port": 587,
                    "mail_from": "sender@example.com",
                },
                layout_config={},
            )
        )
        db.commit()
    login(admin_client)
    monkeypatch.setattr(service, "test_connection", lambda config: None)
    monkeypatch.setattr(
        "routes.admin.mail_settings._save_smtp_attempt",
        lambda db, **values: False,
    )

    response = admin_client.post(
        "/admin/mail-settings/test",
        data={"csrf_token": admin_csrf(admin_client)},
        follow_redirects=True,
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert (
        "Połączenie SMTP działa, ale nie udało się zapisać logu testu. "
        "Sprawdź migracje bazy danych."
    ) in html
    assert "Nie udało się połączyć z serwerem SMTP" not in html


def test_failed_smtp_test_keeps_primary_error_when_log_write_also_fails(
    admin_app, admin_client, monkeypatch
):
    create_user(admin_app)
    service = admin_app.extensions["services"].mail_settings_service
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            SystemMailSettings(
                smtp_config={
                    "host": "smtp.failure.test",
                    "port": 587,
                    "mail_from": "sender@example.com",
                },
                layout_config={},
            )
        )
        db.commit()
    login(admin_client)
    monkeypatch.setattr(
        service,
        "test_connection",
        lambda config: (_ for _ in ()).throw(TimeoutError("timeout")),
    )
    monkeypatch.setattr(
        "routes.admin.mail_settings._save_smtp_attempt",
        lambda db, **values: False,
    )

    response = admin_client.post(
        "/admin/mail-settings/test",
        data={"csrf_token": admin_csrf(admin_client)},
        follow_redirects=True,
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Nie udało się połączyć z serwerem SMTP w wyznaczonym czasie." in html
    assert "Połączenie SMTP działa, ale" not in html


def test_smtp_test_email_contains_footer_logo_as_inline_cid(admin_app, admin_client):
    create_user(admin_app)
    logo_path = Path(admin_app.config["TEMP_DIR"]) / "smtp-test-footer.svg"
    logo_path.write_bytes(b"<svg></svg>")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(
            name="Logo Lubuskie",
            filename="smtp-test-footer.svg",
            storage_path=str(logo_path),
            mime_type="image/svg+xml",
            active=True,
        )
        db.add(logo)
        db.flush()
        db.add(MailFooter(
            form_id=None,
            name="Stopka testowa",
            html_body="<p>Treść stopki testowej</p>",
            logo_id=logo.id,
            logo_width=220,
            logo_alignment="center",
            logo_position="top",
            is_active=True,
            is_default=True,
        ))
        db.add(SystemMailSettings(
            smtp_config={
                "host": "smtp.test",
                "port": 587,
                "mail_from": "sender@example.com",
                "use_tls": True,
                "timeout": 10,
            },
            layout_config={},
        ))
        db.commit()
    login(admin_client)
    sent = []
    dispatch = admin_app.extensions["services"].mail_dispatch_service
    dispatch.smtp_sender = lambda **kwargs: sent.append(kwargs)

    response = admin_client.post(
        "/admin/mail-settings/test-email",
        data={"csrf_token": admin_csrf(admin_client)},
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "Wysłano wiadomość testową SMTP" in response.get_data(as_text=True)
    assert len(sent) == 1
    assert sent[0]["to_emails"] == ["admin@example.com"]
    assert 'src="cid:footer-logo"' in sent[0]["html_body"]
    assert "width:220px" in sent[0]["html_body"]
    assert "text-align:center" in sent[0]["html_body"]
    assert str(logo_path) not in sent[0]["html_body"]
    assert sent[0]["inline_images"] == [{
        "cid": "footer-logo",
        "content": b"<svg></svg>",
        "mime_type": "image/svg+xml",
        "filename": "smtp-test-footer.svg",
    }]


def test_send_mail_logs_email(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            nazwisko="Kowalski",
            imiona="Jan",
            data_json={"imiona": "Jan"},
        )
        db.add(submission)
        db.flush()
        db.add(MailTemplate(form_id=form_id, name="Info", subject="Witaj {{ imiona }}", html_body="<p>{{ submission_id }}</p>"))
        db.commit()
        submission_pk = submission.id
    sent = []
    admin_app.extensions["notification_service"].smtp_sender = lambda **kwargs: sent.append(kwargs)
    login(admin_client)
    token = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/mail",
        data={"csrf_token": token, "to_email": "jan@example.com", "subject": "Witaj {{ imiona }}", "html_body": "<p>{{ submission_id }}</p>"},
    )

    assert response.status_code == 302
    assert sent[0]["subject"] == "Witaj Jan"
    assert "platform-mail-card" in sent[0]["html_body"]
    assert "abc" in sent[0]["html_body"]
    with session_factory() as db:
        log = db.query(EmailLog).one()
        assert log.status == "sent"
        assert log.to_email == "jan@example.com"


def test_send_mail_view_populates_editor_from_selected_template(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            nazwisko="Kowalski",
            imiona="Jan",
        )
        db.add(submission)
        db.flush()
        db.add(
            MailTemplate(
                form_id=form_id,
                name="Wniosek zaakceptowany",
                subject="Temat {{ imiona }}",
                content_html="<p>HTML {{ submission_id }}</p>",
                content_text="TXT {{ submission_id }}",
                html_body="<p>HTML {{ submission_id }}</p>",
                text_body="TXT {{ submission_id }}",
            )
        )
        db.commit()
        submission_pk = submission.id
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").get_data(as_text=True)

    payload_text = html.split("const templates = ", 1)[1].split(";\n", 1)[0]
    payload = json.loads(payload_text)
    assert payload[0]["subject"] == "Temat {{ imiona }}"
    assert payload[0]["html_body"] == "<p>HTML {{ submission_id }}</p>"
    assert payload[0]["text_body"] == "TXT {{ submission_id }}"
    assert "fillFromTemplate()" in html
    assert "Mail systemowy" not in html


def test_send_mail_uses_first_template_when_manual_fields_are_empty(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            nazwisko="Kowalski",
            imiona="Jan",
            data_json={"imiona": "Jan"},
        )
        db.add(submission)
        db.flush()
        db.add(MailTemplate(form_id=form_id, name="Info", subject="Witaj {{ imiona }}", html_body="<p>{{ submission_id }}</p>"))
        db.commit()
        submission_pk = submission.id
    sent = []
    admin_app.extensions["notification_service"].smtp_sender = lambda **kwargs: sent.append(kwargs)
    login(admin_client)
    token = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/mail",
        data={"csrf_token": token, "to_email": "jan@example.com"},
    )

    assert response.status_code == 302
    assert sent[0]["subject"] == "Witaj Jan"
    assert "abc" in sent[0]["html_body"]
    with session_factory() as db:
        log = db.query(EmailLog).one()
        assert log.status == "sent"


def test_send_mail_without_template_is_skipped(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            nazwisko="Kowalski",
            imiona="Jan",
            data_json={"imiona": "Jan"},
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    sent = []
    admin_app.extensions["notification_service"].smtp_sender = lambda **kwargs: sent.append(kwargs)
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    assert "Mail systemowy" in html

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/mail",
        data={"csrf_token": token, "to_email": "jan@example.com"},
    )

    assert response.status_code == 302
    assert sent == []


def test_send_mail_failure_flash_includes_error_reason(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="abc",
            form_slug="sample_form",
            form_name="Sample",
            email="jan@example.com",
            nazwisko="Kowalski",
            imiona="Jan",
            data_json={"imiona": "Jan"},
        )
        db.add(submission)
        db.flush()
        db.add(MailTemplate(form_id=form_id, name="Info", subject="Test", html_body="<p>Test</p>"))
        db.commit()
        submission_pk = submission.id

    def fail_smtp(**kwargs):
        raise RuntimeError("SMTP auth failed")

    admin_app.extensions["notification_service"].smtp_sender = fail_smtp
    login(admin_client)
    token = admin_client.get(f"/admin/forms/{form_id}/submissions/{submission_pk}/mail").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/mail",
        data={"csrf_token": token, "to_email": "jan@example.com"},
        follow_redirects=True,
    )

    html = response.get_data(as_text=True)
    assert "Nie udalo sie wyslac maila: RuntimeError: SMTP auth failed" in html
    with session_factory() as db:
        log = db.query(EmailLog).one()
        assert log.status == "failed"
        assert log.error_message == "RuntimeError: SMTP auth failed"


def test_send_bulk_mail_to_selected_submissions_uses_matching_template(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(
            FormSubmission(
                submission_id="abc",
                form_slug="sample_form",
                form_name="Sample",
                email="jan@example.com",
                imiona="Jan",
                officer_decision="accepted",
                data_json={"imiona": "Jan"},
            )
        )
        db.add(
            MailTemplate(
                form_id=form_id,
                name="Accepted",
                subject="Witaj {{ imiona }}",
                html_body="<p>{{ officer_decision }}</p>",
                text_body="{{ submission_id }}",
                trigger_event="manual_bulk",
                trigger_decision="accepted",
            )
        )
        db.commit()
    sent = []
    admin_app.extensions["notification_service"].smtp_sender = lambda **kwargs: sent.append(kwargs)
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/mail-selected",
        data={"csrf_token": token, "submission_ids": ["abc"]},
    )

    assert response.status_code == 302
    assert sent[0]["subject"] == "Witaj Jan"
    assert "platform-mail-card" in sent[0]["html_body"]
    assert "accepted" in sent[0]["html_body"]
    assert sent[0]["text_body"].strip().startswith("Accepted")
    with session_factory() as db:
        log = db.query(EmailLog).one()
        assert log.status == "sent"
        assert log.public_submission_id == "abc"


def test_bulk_mail_opens_composer_with_preview_variables_and_missing_addresses(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="bulk_composer", name="Bulk composer")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add_all(
            [
                FormSubmission(submission_id="with-email", form_slug="bulk_composer", form_name="Bulk", email="ok@example.com"),
                FormSubmission(submission_id="without-email", form_slug="bulk_composer", form_name="Bulk", email=""),
            ]
        )
        db.commit()
    login(admin_client)
    list_html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)
    token = list_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/mail-selected",
        data={
            "csrf_token": token,
            "compose": "1",
            "submission_ids": ["with-email", "without-email"],
        },
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Wyślij wiadomość" in html
    assert "without-email" in html
    assert "Dostępne zmienne" in html
    assert "data-mail-preview" in html


def test_admin_mail_module_keeps_endpoint_names(admin_app):
    import routes.admin.mail as admin_mail
    from routes.admin import bp

    assert admin_mail is not None
    assert bp.name == "admin"
    with admin_app.test_request_context():
        assert url_for("admin.submission_mail", form_id=1, submission_pk=2) == "/admin/forms/1/submissions/2/mail"
        assert url_for("admin.submissions_mail_selected", form_id=1) == "/admin/forms/1/submissions/mail-selected"
        assert url_for("admin.mail_templates_list", form_id=1) == "/admin/forms/1/mail-templates"
        assert url_for("admin.mail_template_delete", form_id=1, template_id=2) == "/admin/forms/1/mail-templates/2/delete"
        assert url_for("admin.mail_footers_list", form_id=1) == "/admin/forms/1/mail-footers"


@pytest.mark.parametrize("role,email", [("super_admin", "admin@example.com"), ("admin", "admin-role@example.com")])
def test_mail_template_delete_allowed_for_all_managing_roles(admin_app, admin_client, role, email):
    user_id = create_user(admin_app, email=email, role=role)
    form_id = create_form(admin_app, slug=f"form_{role}", name="Sample", user_id=None if role == "super_admin" else user_id)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        template = MailTemplate(form_id=form_id, name="Delete me", subject="Temat", html_body="<p>Test</p>")
        db.add(template)
        db.commit()
        template_id = template.id
    login(admin_client, email=email)
    html = admin_client.get(f"/admin/forms/{form_id}/mail-templates").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(f"/admin/forms/{form_id}/mail-templates/{template_id}/delete", data={"csrf_token": token})

    assert response.status_code == 302
    with session_factory() as db:
        assert db.get(MailTemplate, template_id) is None


def test_form_manager_sees_mail_templates_read_only_and_cannot_delete(admin_app, admin_client):
    manager_id = create_user(admin_app, email="manager-readonly@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="readonly_mail", name="Readonly", user_id=manager_id)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        template = MailTemplate(form_id=form_id, name="Keep me", subject="Temat", html_body="<p>Test</p>")
        db.add(template)
        db.commit()
        template_id = template.id
    login(admin_client, email="manager-readonly@example.com")

    response = admin_client.get(f"/admin/forms/{form_id}/mail-templates")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Widok tylko do odczytu" in html
    assert "Importuj HTML" not in html
    with admin_client.session_transaction() as session:
        token = session["admin_csrf_token"]
    delete_response = admin_client.post(
        f"/admin/forms/{form_id}/mail-templates/{template_id}/delete",
        data={"csrf_token": token},
    )
    assert delete_response.status_code == 403


def test_mail_template_delete_requires_manage_permission(admin_app, admin_client):
    user_id = create_user(admin_app, email="viewer@example.com", role="admin")
    form_id = create_form(admin_app, slug="view_only", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(FormPermission(user_id=user_id, form_id=form_id, can_manage=False))
        template = MailTemplate(form_id=form_id, name="Keep me", subject="Temat", html_body="<p>Test</p>")
        db.add(template)
        db.commit()
        template_id = template.id
    login(admin_client, email="viewer@example.com")
    token = admin_client.get("/admin/forms").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(f"/admin/forms/{form_id}/mail-templates/{template_id}/delete", data={"csrf_token": token})

    assert response.status_code == 403
    with session_factory() as db:
        assert db.get(MailTemplate, template_id) is not None


def test_zip_mail_template_import_reads_content_text_and_assets(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    login(admin_client)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr(
            "zaakceptowany.html",
            "<html><head><style>.x{color:red}</style></head><body>Wniosek zaakceptowany<br>{{ imiona }}<br>Instrukcja<br>Przejdz dalej</body></html>",
        )
        archive.writestr("zaakceptowany.txt", "Wniosek zaakceptowany\n\n{{ imiona }}\n\nInstrukcja\nPrzejdz dalej")
        archive.writestr("images/logo.png", b"png-bytes")
    zip_buffer.seek(0)
    html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/import-zip").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-templates/import-zip",
        data={
            "csrf_token": token,
            "name": "ZIP",
            "subject": "Temat",
            "zip_file": (zip_buffer, "mail.zip"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        template = db.query(MailTemplate).filter_by(name="ZIP").one()
        assert template.content_title == "Wniosek zaakceptowany"
        assert "<style>" not in template.html_body
        assert "{{ imiona }}" in template.html_body
        assert "{{ imiona }}" in template.text_body
        assert "Przejdz dalej" in template.instruction_text
        asset = db.query(MailTemplateAsset).one()
        assert asset.filename == "logo.png"
        assert asset.content == b"png-bytes"


def test_html_mail_template_import_is_primary_path(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    login(admin_client)
    list_html = admin_client.get(f"/admin/forms/{form_id}/mail-templates").get_data(as_text=True)

    assert "Importuj HTML" in list_html
    assert "Dodaj szablon" in list_html
    assert "Import ZIP (zaawansowane)" in list_html

    html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/import-html").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-templates/import-html",
        data={
            "csrf_token": token,
            "name": "HTML",
            "subject": "Temat {{ submission_id }}",
            "body_html": "<h2>Potwierdzenie</h2><p>{{ submission_id }}</p>",
            "body_text": "Potwierdzenie {{ submission_id }}",
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        template = db.query(MailTemplate).filter_by(name="HTML").one()
        assert template.content_title == "Potwierdzenie"
        assert "{{ submission_id }}" in template.content_html
        assert template.content_text == "Potwierdzenie {{ submission_id }}"


def test_form_manager_cannot_use_technical_zip_import(admin_app, admin_client):
    manager_id = create_user(admin_app, email="manager@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="sample_form", name="Sample", user_id=manager_id)
    login(admin_client, email="manager@example.com")

    response = admin_client.get(f"/admin/forms/{form_id}/mail-templates/import-zip")

    assert response.status_code == 403


def test_platform_mail_layout_has_expected_defaults():
    from types import SimpleNamespace
    from services.mail_template_service import MAIL_LAYOUT, render_platform_mail_html

    template = SimpleNamespace(
        name="Wniosek zaakceptowany",
        content_title="Wniosek zaakceptowany",
        html_body="<p>Dzien dobry {{ imiona }}</p>",
        text_body="Dzien dobry {{ imiona }}",
        instruction_html="<p>Krok</p>",
        instruction_text="Instrukcja\nKrok",
        footer_note="Numer: <strong>{{ submission_id }}</strong>",
    )
    html = render_platform_mail_html(
        template,
        {"imiona": "Jan", "submission_id": "abc", "form_name": "Sample", "status_label": "Zaakceptowano"},
    )

    assert MAIL_LAYOUT["container_width"] == "600px"
    assert "width:600px" in html
    assert "#f0f1f5" in html
    assert "#1d2e5b" in html
    assert "#c8a35d" in html
    assert "Dzien dobry Jan" in html
    assert "Numer zgloszenia" in html


def test_mail_template_edit_does_not_expose_style_controls(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new").get_data(as_text=True)

    assert "Wgraj plik HTML" in html
    assert "Wklej HTML" in html
    assert "Treść TXT" in html
    assert "Instrukcja" in html
    assert "Kolor" not in html
    assert 'name="label_color"' not in html
    assert 'name="label_background"' not in html
    assert "Margines" not in html


def test_admin_form_and_mail_editor_explain_required_fields(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="required_notes", name="Required Notes")
    login(admin_client)

    form_html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    mail_html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new").get_data(as_text=True)

    for html in (form_html, mail_html):
        assert "Pola oznaczone" in html
        assert "są obowiązkowe." in html
        assert 'class="required-marker" aria-hidden="true">*</span>' in html
        assert 'aria-required="true"' in html


def test_mail_template_preview_renders_platform_layout(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new").get_data(as_text=True)

    assert "platform-mail-card" in html
    assert "#f0f1f5" in html
    assert "Numer zgloszenia" in html


def test_mail_template_editor_uses_catalog_dropdowns_and_all_form_variables(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="mail_catalog",
        name="Katalog maila",
        definition_json={
            "title": "Katalog maila",
            "fields": [{"name": "pesel", "label": "Numer PESEL", "type": "text"}],
            "workflow": {
                "initial_step": "submission",
                "steps": [
                    {
                        "id": "submission",
                        "admin_label": "Złożenie wniosku",
                        "status": "FORM_SUBMITTED",
                        "triggers": ["application_submitted"],
                    }
                ],
            },
            "documents": [
                {
                    "id": "declaration",
                    "fields": [
                        {
                            "name": "selected_trainings",
                            "label": "Wybrane szkolenia",
                            "type": "training_selection",
                            "catalog": [{"id": "excel", "name": "Excel"}],
                        }
                    ],
                }
            ],
        },
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(FormField(form_id=form_id, name="telefon", label="Numer telefonu", type="phone", active=True))
        db.commit()
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert '<select name="trigger_event">' in html
    assert '<select name="trigger_status">' in html
    assert ">Wniosek złożony</option>" in html
    assert "Wniosek złożony — FORM_SUBMITTED" not in html
    assert "{{ app_name }}" in html
    assert "{{ public_submission_id }}" in html
    assert "{{ pesel }}" in html
    assert "{{ telefon }}" in html
    assert "{{ selected_trainings_count }}" in html
    assert "{{ available_trainings }}" in html
    assert "{{ available_trainings_table }}" in html
    assert "{{ available_trainings_list }}" in html
    assert "{{ available_trainings_text }}" in html
    assert 'data-insert-variable="{{ pesel }}"' in html
    assert 'class="admin-card mail-variable-sidebar"' in html
    assert 'data-variable-collapse' in html
    assert 'data-variable-expand' in html
    assert html.index("5. Podgląd wiadomości") < html.index('class="admin-card mail-variable-sidebar"')
    assert "js/mail_variables_panel.js" in html


def test_mail_template_live_preview_endpoint_uses_fallback_context_without_submissions(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="mail_preview_fallback",
        name="Formularz podglądu",
        title="Tytuł formularza",
        definition_json={"title": "Tytuł formularza", "fields": [{"name": "imie_firmy", "label": "Nazwa firmy"}]},
    )
    login(admin_client)
    editor_response = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new")
    editor_html = editor_response.get_data(as_text=True)
    token = editor_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-templates/preview",
        data={
            "csrf_token": token,
            "name": "Powitanie",
            "subject": "Witaj w {{ form_title }}",
            "content_title": "Zgłoszenie {{ submission_id }}",
            "body_html": "<p>{{ imiona }} — {{ process_status_label }}</p><script>alert('x')</script>",
            "body_text": "{{ imiona }} / {{ public_submission_id }}",
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["ok"] is True
    assert payload["subject"] == "Witaj w Tytuł formularza"
    assert "Jan" in payload["html"]
    assert "<script" not in payload["html"]
    assert "6ef64b28-530b" in payload["text"]
    assert "Brak zgłoszeń dla formularza" in editor_html
    assert 'setTimeout(refreshPreview, 400)' in editor_html
    assert "data-preview-frame" in editor_html
    assert " sandbox " in editor_html


def test_mail_variable_endpoints_return_form_and_submission_context(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="mail_variables_api",
        name="Zmienne API",
        definition_json={
            "title": "Zmienne API",
            "fields": [{"name": "firma", "label": "Nazwa firmy", "type": "text"}],
        },
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        submission = FormSubmission(
            submission_id="mail-variables-submission",
            form_slug="mail_variables_api",
            form_name="Zmienne API",
            email="real@example.com",
            data_json={"firma": "Realna firma"},
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client)

    form_response = admin_client.get(f"/admin/forms/{form_id}/mail-variables")
    submission_response = admin_client.get(f"/admin/submissions/{submission_pk}/mail-variables")

    assert form_response.status_code == 200
    assert submission_response.status_code == 200
    form_payload = form_response.get_json()
    submission_payload = submission_response.get_json()
    assert [group["name"] for group in form_payload["categories"]] == [
        "Systemowe",
        "Formularz",
        "Zgłoszenie",
        "Pola formularza",
        "Workflow",
        "Decyzje urzędnika",
        "Dokumenty",
        "E-mail",
    ]
    form_variables = {
        variable["key"]: variable
        for group in form_payload["categories"]
        for variable in group["variables"]
    }
    submission_variables = {
        variable["key"]: variable
        for group in submission_payload["categories"]
        for variable in group["variables"]
    }
    assert form_variables["firma"]["placeholder"] == "{{ firma }}"
    assert form_variables["firma"]["available_in_html"] is True
    assert form_variables["firma"]["available_in_txt"] is True
    assert submission_variables["firma"]["example"] == "Realna firma"
    assert submission_variables["email"]["example"] == "real@example.com"


def test_mail_preview_reports_unknown_variable_without_server_error(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="unknown_mail_variable", name="Nieznana zmienna")
    login(admin_client)
    editor_html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new").get_data(as_text=True)
    token = editor_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-templates/preview",
        data={
            "csrf_token": token,
            "subject": "Test {{ nieznana_zmienna }}",
            "content_title": "Tytuł",
            "body_html": "<p>{{ nieznana_zmienna }}</p>",
            "body_text": "{{ nieznana_zmienna }}",
        },
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["ok"] is True
    assert payload["unknown_variables"] == ["nieznana_zmienna"]
    assert payload["warnings"] == ["Nierozpoznana zmienna: {{ nieznana_zmienna }}"]


def test_custom_submission_mail_uses_reusable_collapsible_variable_panel(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="custom_mail_panel", name="Własna wiadomość")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        submission = FormSubmission(
            submission_id="custom-mail-panel-submission",
            form_slug="custom_mail_panel",
            form_name="Własna wiadomość",
            email="participant@example.com",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client)

    html = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/mail"
    ).get_data(as_text=True)

    assert 'data-mail-variable-scope' in html
    assert 'class="admin-card mail-variable-sidebar"' in html
    assert 'data-variable-storage-key="submission-mail-variables-collapsed"' in html
    assert 'data-variable-target' in html
    assert 'data-variable-default-target' in html
    assert "js/mail_variables_panel.js" in html


def test_mail_template_live_preview_renders_available_training_formats_from_form_config(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="mail_training_preview",
        name="Szkolenia",
        definition_json={
            "title": "Szkolenia",
            "fields": [],
            "documents": [
                {
                    "id": "declaration",
                    "fields": [
                        {
                            "type": "training_selection",
                            "name": "selected_trainings",
                            "currency": "PLN",
                            "catalog": [
                                {
                                    "id": "excel",
                                    "name": "Excel <script>alert(1)</script>",
                                    "description": "Arkusze i raporty",
                                    "price": "1234.50",
                                    "capacity": 2,
                                    "dates": [
                                        {
                                            "start_date": "2026-08-01",
                                            "location": "Zielona Góra",
                                        }
                                    ],
                                }
                            ],
                        }
                    ],
                }
            ],
        },
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            FormSubmission(
                submission_id="training-mail-1",
                form_slug="mail_training_preview",
                form_name="Szkolenia",
                process_status="FORM_SUBMITTED",
                selected_trainings='[{"id":"excel","name":"Excel","price":"1234.50"}]',
            )
        )
        db.commit()
    login(admin_client)
    editor_html = admin_client.get(f"/admin/forms/{form_id}/mail-templates/new").get_data(as_text=True)
    token = editor_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-templates/preview",
        data={
            "csrf_token": token,
            "name": "Oferta",
            "subject": "Dostępne szkolenia",
            "content_title": "Oferta",
            "body_html": "{{ available_trainings_table }}",
            "body_text": "{{ available_trainings_text }}",
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["ok"] is True
    assert "1 234,50 zł" in payload["html"]
    assert "Dostępne: 2" in payload["html"]
    assert "01.08.2026" in payload["html"]
    assert "Zielona Góra" in payload["html"]
    assert "<script>alert(1)</script>" not in payload["html"]
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in payload["html"]
    assert "Dostępne szkolenia:" in payload["text"]
    assert "Dostępne miejsca: 2" in payload["text"]


def test_simple_html_full_document_extracts_body_and_sanitizes():
    from services.mail_template_service import parse_mail_content

    parsed = parse_mail_content(
        '<!doctype html><html><head><title>x</title></head><body><h2>Odrzucony</h2><p>{{ decision }}</p><script>alert(1)</script><a href="javascript:alert(1)">x</a></body></html>'
    )

    assert parsed.title == "Odrzucony"
    assert "<body" not in parsed.body_html
    assert "<script" not in parsed.body_html
    assert "javascript:" not in parsed.body_html
    assert "{{ decision }}" in parsed.body_html


def test_mail_import_inlines_safe_embedded_css_and_keeps_inline_styles():
    from services.mail_template_service import parse_mail_content

    parsed = parse_mail_content(
        """
        <!doctype html>
        <html>
        <head>
            <style>
                .lead { color: #1d2e5b; font-weight: 700; background-image: url("https://bad.test/x.png"); }
                #cta { background-color: #c8a35d; padding: 12px; position: absolute; }
                p { line-height: 1.5; }
                .lead strong { color: red; }
            </style>
        </head>
        <body>
            <h2>Potwierdzenie</h2>
            <p class="lead" style="font-size: 18px; position: fixed;">Witaj {{ imie }}</p>
            <div id="cta">Dalej</div>
        </body>
        </html>
        """
    )

    assert 'style="' in parsed.body_html
    assert "color:#1d2e5b" in parsed.body_html
    assert "font-weight:700" in parsed.body_html
    assert "line-height:1.5" in parsed.body_html
    assert "font-size:18px" in parsed.body_html
    assert 'style="background-color:#c8a35d;padding:12px"' in parsed.body_html
    assert "background-image" not in parsed.body_html
    assert "position" not in parsed.body_html
    assert "class=" not in parsed.body_html
    assert "id=" not in parsed.body_html
    assert "<style" not in parsed.body_html


def test_simple_html_fragment_is_preserved_as_content():
    from services.mail_template_service import parse_mail_content

    parsed = parse_mail_content("<p>Potwierdzenie {{ submission_id }}</p>")

    assert parsed.body_html == "<p>Potwierdzenie {{ submission_id }}</p>"


def test_platform_mail_renders_jinja_submission_and_agreement():
    from types import SimpleNamespace
    from services.mail_template_service import build_mail_context, render_platform_mail_html

    template = SimpleNamespace(
        name="Umowa",
        content_title="Umowa {{ agreement.get(\"number\") }}",
        content_html="<p>{{ submission.get(\"imie\") }} / {{ agreement.get(\"training_name\") }}</p>",
        content_text="",
        html_body="<p>{{ submission.get(\"imie\") }} / {{ agreement.get(\"training_name\") }}</p>",
        text_body="",
        instruction_html="",
        instruction_text="",
        footer_note="",
    )
    form = SimpleNamespace(name="Sample", slug="sample")
    submission = SimpleNamespace(
        data_json={"imie": "Jan"},
        __table__=SimpleNamespace(columns=[]),
    )
    context = build_mail_context(form, submission, extra={"agreement": {"number": "1/2026", "training_name": "Excel"}})
    html = render_platform_mail_html(template, context)

    assert "Umowa 1/2026" in html
    assert "Jan / Excel" in html
    assert "platform-instruction-title" not in html
    assert "Instrukcja" not in html


def test_platform_mail_renders_html_escaped_jinja_quotes():
    from types import SimpleNamespace
    from services.mail_template_service import render_platform_mail_html

    template = SimpleNamespace(
        name="Umowa",
        content_title="Umowa",
        content_html="<p>{{ agreement.get(&quot;number&quot;) }}</p>",
        content_text="",
        html_body="",
        text_body="",
        instruction_html="",
        instruction_text="",
        footer_note="",
    )

    html = render_platform_mail_html(template, {"agreement": {"number": "1/2026"}, "form_name": "Sample", "submission_id": "abc"})

    assert "1/2026" in html
    assert "&quot;number&quot;" not in html


def test_auto_template_selection_by_type_for_decisions(admin_app):
    from routes.admin import select_mail_template

    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    accepted = MailTemplate(form_id=form_id, name="Accepted", template_type="accepted", subject="", html_body="")
    rejected = MailTemplate(form_id=form_id, name="Rejected", template_type="rejected", subject="", html_body="")
    submission = FormSubmission(
        submission_id="abc",
        form_slug="sample_form",
        form_name="Sample",
        officer_decision="rejected",
    )

    assert select_mail_template([accepted, rejected], submission, "manual_bulk").name == "Rejected"


def test_super_admin_can_upload_logo(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/logos").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
        b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
        b"\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01"
        b"\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    response = admin_client.post(
        "/admin/logos",
        data={
            "csrf_token": token,
            "name": "Logo",
            "logo_file": (io.BytesIO(png_bytes), "logo.png"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = db.query(Logo).one()
        assert logo.name == "Logo"
        assert logo.active is True
        assert logo.size_bytes == len(png_bytes)
        assert logo.checksum_sha256
        logo_id = logo.id

    edit_html = admin_client.get(f"/admin/logos/{logo_id}/edit").get_data(as_text=True)
    edit_token = edit_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        f"/admin/logos/{logo_id}/edit",
        data={"csrf_token": edit_token, "name": "Logo po zmianie"},
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = db.get(Logo, logo_id)
        assert logo.name == "Logo po zmianie"
        assert logo.active is False


def test_super_admin_can_toggle_logo_and_fetch_asset(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    logo_path = Path(admin_app.config["TEMP_DIR"]) / "logo.png"
    logo_path.parent.mkdir(parents=True, exist_ok=True)
    logo_path.write_bytes(b"logo-bytes")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(
            name="Logo",
            filename="logo.png",
            storage_path=str(logo_path),
            mime_type="image/png",
            size_bytes=10,
            checksum_sha256="abc",
            active=True,
        )
        db.add(logo)
        db.commit()
        logo_id = logo.id

    html = admin_client.get("/admin/logos").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    assert "Logo" in html
    assert admin_client.get(f"/admin/logos/{logo_id}/asset").status_code == 200

    response = admin_client.post(f"/admin/logos/{logo_id}/toggle", data={"csrf_token": token})

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = db.get(Logo, logo_id)
        assert logo.active is False
    assert admin_client.get(f"/admin/logos/{logo_id}/asset").status_code == 200


def test_regular_admin_cannot_toggle_logo(admin_app, admin_client):
    create_user(admin_app, role="admin")
    logo_path = Path(admin_app.config["TEMP_DIR"]) / "logo.png"
    logo_path.parent.mkdir(parents=True, exist_ok=True)
    logo_path.write_bytes(b"logo-bytes")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(
            name="Logo",
            filename="logo.png",
            storage_path=str(logo_path),
            mime_type="image/png",
            size_bytes=10,
            checksum_sha256="abc",
            active=True,
        )
        db.add(logo)
        db.commit()
        logo_id = logo.id

    login(admin_client)
    with admin_client.session_transaction() as session:
        token = session["admin_csrf_token"]

    response = admin_client.post(f"/admin/logos/{logo_id}/toggle", data={"csrf_token": token})

    assert response.status_code == 403
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = db.get(Logo, logo_id)
        assert logo.active is True


def test_regular_admin_can_select_existing_logo_but_not_upload(admin_app, admin_client):
    manager_id = create_user(admin_app, email="manager@example.com", role="admin")
    form_id = create_form(admin_app, slug="owned", name="Owned", user_id=manager_id)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        logo = Logo(name="Logo", filename="logo.png", storage_path=str(Path(admin_app.config["TEMP_DIR"]) / "logo.png"), mime_type="image/png", active=True)
        db.add(logo)
        db.commit()
        logo_id = logo.id
    login(admin_client, email="manager@example.com")
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Owned",
            "title": "Owned",
            "slug": "owned",
            "logo_id": str(logo_id),
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    with session_factory() as db:
        form = db.get(Form, form_id)
        assert form.logo_id == logo_id

    assert admin_client.get("/admin/forms/upload").status_code == 403
    with admin_client.session_transaction() as session:
        upload_token = session["admin_csrf_token"]
    response = admin_client.post(
        "/admin/logos",
        data={
            "csrf_token": upload_token,
            "name": "Nope",
            "logo_file": (io.BytesIO(b"png"), "logo.png"),
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 403


def test_mail_footer_uses_logo_library_instead_of_manual_path(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        logo = Logo(
            name="Logo UMWL",
            filename="logo.png",
            storage_path=str(Path(admin_app.config["TEMP_DIR"]) / "logo.png"),
            mime_type="image/png",
            active=True,
        )
        db.add(logo)
        db.commit()
        logo_id = logo.id

    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/mail-footers/new").get_data(as_text=True)
    assert "Logo z Nextcloud lub URL" not in html
    assert "Logo w stopce" in html
    assert "Brak logo" in html
    assert "Logo UMWL" in html
    assert "Logo można dodać tylko z poziomu konta super_admin." in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-footers/new",
        data={
            "csrf_token": token,
            "name": "Stopka",
            "logo_id": str(logo_id),
            "logo_path": "https://example.com/manual.png",
            "html_body": "<p>Kontakt</p>",
            "is_active": "on",
            "is_default": "on",
        },
    )

    assert response.status_code == 302
    with session_factory() as db:
        footer = db.query(MailFooter).one()
        assert footer.logo_id == logo_id
        assert footer.logo_path == ""
        assert footer.html_body == "<p>Kontakt</p>"
        assert footer.is_active is True
        assert footer.is_default is True


def test_mail_footer_rejects_inactive_logo(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="sample_form", name="Sample")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        active_logo = Logo(
            name="Aktywne logo",
            filename="active.png",
            storage_path=str(Path(admin_app.config["TEMP_DIR"]) / "active.png"),
            mime_type="image/png",
            active=True,
        )
        inactive_logo = Logo(
            name="Nieaktywne logo",
            filename="inactive.png",
            storage_path=str(Path(admin_app.config["TEMP_DIR"]) / "inactive.png"),
            mime_type="image/png",
            active=False,
        )
        db.add_all([active_logo, inactive_logo])
        db.commit()
        inactive_logo_id = inactive_logo.id

    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/mail-footers/new").get_data(as_text=True)
    assert "Aktywne logo" in html
    assert "Nieaktywne logo" not in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-footers/new",
        data={
            "csrf_token": token,
            "name": "Stopka",
            "logo_id": str(inactive_logo_id),
            "html_body": "<p>Kontakt</p>",
            "is_active": "on",
        },
    )

    assert response.status_code == 403
    with session_factory() as db:
        assert db.query(MailFooter).count() == 0


def test_user_without_form_access_cannot_edit_mail_footer(admin_app, admin_client):
    owner_id = create_user(admin_app, email="owner@example.com", role="form_manager")
    create_user(admin_app, email="blocked@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="owned", name="Owned", user_id=owner_id)

    login(admin_client, email="blocked@example.com")
    response = admin_client.get(f"/admin/forms/{form_id}/mail-footers/new")

    assert response.status_code == 403


def test_submission_files_store_metadata_only(admin_app):
    create_user(admin_app)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(submission_id="abc", form_slug="sample_form", form_name="Sample")
        db.add(submission)
        db.flush()
        db.add(
            SubmissionFile(
                submission_id=submission.id,
                public_submission_id="abc",
                form_slug="sample_form",
                filename="doc.pdf",
                storage_path="output/sample_form/pdf/doc.pdf",
                mime_type="application/pdf",
                size_bytes=100,
                checksum_sha256="abc",
                signed=False,
            )
        )
        db.commit()
        file_row = db.query(SubmissionFile).one()
        assert hasattr(file_row, "storage_path")
        assert hasattr(file_row, "signature_validation_result")
        assert hasattr(file_row, "generated_at")
        assert not hasattr(file_row, "content")


def test_super_admin_can_block_user(admin_app, admin_client):
    create_user(admin_app)
    user_id = create_user(admin_app, email="blocked@example.com", role="form_manager")
    login(admin_client)
    token = admin_client.get("/admin/users").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(f"/admin/users/{user_id}/toggle-block", data={"csrf_token": token})

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        user = db.get(User, user_id)
        assert user.is_blocked is True


def test_super_admin_can_change_user_password(admin_app, admin_client):
    from werkzeug.security import check_password_hash

    create_user(admin_app)
    user_id = create_user(admin_app, email="password@example.com", role="form_manager")
    login(admin_client)
    html = admin_client.get(f"/admin/users/{user_id}/password").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/users/{user_id}/password",
        data={"csrf_token": token, "password": "new-secret", "password_confirm": "new-secret"},
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        user = db.get(User, user_id)
        assert check_password_hash(user.password_hash, "new-secret")


def test_super_admin_can_delete_user(admin_app, admin_client):
    create_user(admin_app)
    user_id = create_user(admin_app, email="delete@example.com", role="form_manager")
    form_id = create_form(admin_app, slug="delete_user_form", name="Delete User Form", user_id=user_id)
    login(admin_client)
    token = admin_client.get("/admin/users").get_data(as_text=True).split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(f"/admin/users/{user_id}/delete", data={"csrf_token": token})

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(User, user_id) is None
        assert db.query(FormPermission).filter_by(user_id=user_id).count() == 0
        assert db.get(Form, form_id).created_by_id is None


def test_form_training_catalog_can_be_edited_in_admin(admin_app, admin_client):
    create_user(admin_app)
    definition = {
        "title": "Training Form",
        "fields": [],
        "documents": [
            {
                "id": "declaration",
                "label": "Deklaracja",
                "kind": "generated_pdf",
                "enabled": True,
                "template_html": "<p>Deklaracja</p>",
                "fields": [
                    {"type": "section", "label": "Wybór szkoleń"},
                    {"type": "section", "label": "Oświadczenia uczestnika"},
                    {"type": "checkbox", "name": "osw_rodo"},
                ],
            }
        ],
        "workflow": {
            "name": "Workflow",
            "initial_step": "submission",
            "steps": [{"id": "submission", "type": "end"}],
        },
    }
    form_id = create_form(admin_app, slug="training_form", name="Training Form", definition_json=definition)
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Training Form",
            "slug": "training_form",
            "title": "Training Form",
            "sort_order": "0",
            "workflow_name": "Workflow",
            "workflow_initial_step": "submission",
            "workflow_json": json.dumps(definition["workflow"]),
            "is_active": "on",
            "is_public": "on",
            "training_selection_enabled": "on",
            "training_selection_name": "selected_trainings",
            "training_selection_label": "Wybierz szkolenia",
            "training_selection_max_total": "7000",
            "training_selection_currency": "PLN",
            "training_selection_required": "on",
            "training_active_present": "1",
            "training_item_id": ["s1", "s2"],
            "training_item_name": ["Excel zaawansowany", "Kadry i płace"],
            "training_item_price": ["6200", "12"],
            "training_item_capacity": ["10", "5"],
            "training_item_active": ["0", "1"],
            "training_item_sort_order": ["1", "2"],
            "training_item_description": ["Arkusze i raporty", "Prawo pracy w praktyce"],
            "training_item_admin_comment": ["Przynieś własny laptop.", ""],
            "training_item_low_seats_comment": ["Tego komentarza nie pokazuj.", "Zostało niewiele miejsc."],
            "training_date_training_index": ["0"],
            "training_date_start_date": ["2026-09-01"],
            "training_date_end_date": [""],
            "training_date_start_time": ["09:00"],
            "training_date_end_time": ["12:00"],
            "training_date_location": ["Zielona Góra"],
            "training_date_description": ["Warsztat stacjonarny"],
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        declaration = next(document for document in form.definition_json["documents"] if document["id"] == "declaration")
        training_field = next(field for field in declaration["fields"] if field["type"] == "training_selection")
        assert [field.get("name") or field.get("label") for field in declaration["fields"]] == [
            "Wybór szkoleń",
            "selected_trainings",
            "Oświadczenia uczestnika",
            "osw_rodo",
        ]
        assert training_field["max_total_amount"] == "7000.00"
        assert training_field["currency"] == "PLN"
        assert training_field["catalog"] == [
            {
                "id": "s1",
                "name": "Excel zaawansowany",
                "price": "6200.00",
                "capacity": 10,
                "description": "Arkusze i raporty",
                "admin_comment": "Przynieś własny laptop.",
                "low_seats_comment": "Tego komentarza nie pokazuj.",
                "dates": [
                    {
                        "start_date": "2026-09-01",
                        "end_date": "",
                        "start_time": "09:00",
                        "end_time": "12:00",
                        "location": "Zielona Góra",
                        "description": "Warsztat stacjonarny",
                    }
                ],
                "active": True,
                "sort_order": 1,
            },
            {
                "id": "s2",
                "name": "Kadry i płace",
                "price": "12.00",
                "capacity": 5,
                "description": "Prawo pracy w praktyce",
                "low_seats_comment": "Zostało niewiele miejsc.",
                "dates": [],
                "active": True,
                "sort_order": 2,
            },
        ]
        db.add(
            FormSubmission(
                submission_id="training-declaration-1",
                form_slug="training_form",
                form_name="Training Form",
                officer_decision="accepted",
                acceptance_required="Tak",
                declaration_required="Tak",
                declaration_generated="Tak",
                declaration_signed="Tak",
                declaration_signature_valid="Tak",
                access_token="training-secret",
            )
        )
        db.commit()

    edit_html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    assert 'name="training_item_code"' not in edit_html
    assert "RRRR-MM-DD|" not in edit_html
    assert 'type="date" name="training_date_start_date" value="2026-09-01"' in edit_html
    assert "Dodaj termin" in edit_html
    assert '<details class="admin-training-item admin-training-card"' in edit_html
    assert "Rozwiń wszystkie" in edit_html
    assert "Zwiń wszystkie" in edit_html
    assert "Mało miejsc" in edit_html
    assert "Brak terminów" in edit_html
    assert "data-training-summary-name" in edit_html

    declaration_response = admin_client.get("/declaration/training_form/training-declaration-1")
    declaration_html = declaration_response.get_data(as_text=True)

    assert declaration_response.status_code == 200
    assert "Excel zaawansowany" not in declaration_html
    assert "Kadry i płace" not in declaration_html

    training_response = admin_client.get(
        "/submissions/training-declaration-1/trainings?token=training-secret"
    )
    training_html = training_response.get_data(as_text=True)
    assert training_response.status_code == 200
    assert "Excel zaawansowany" in training_html
    assert "Kadry i płace" in training_html
    assert "7 000,00 zł" in training_html
    assert "6 200,00 zł" in training_html
    assert "12,00 zł" in training_html
    assert "Przynieś własny laptop." in training_html
    assert "Dostępne" in training_html
    assert "Zajęte" in training_html
    assert "Limit" in training_html
    assert ">10</dd>" in training_html
    assert ">5</dd>" in training_html
    assert "Zostało niewiele miejsc." in training_html
    assert "Tego komentarza nie pokazuj." not in declaration_html
    assert "01.09.2026, 09:00–12:00" in training_html
    assert "Zielona Góra" in training_html

    save_response = admin_client.post(
        "/submissions/training-declaration-1/trainings?token=training-secret",
        data={"selected_trainings": ["s1", "s2"]},
    )
    assert save_response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        selected = db.query(SubmissionTraining).order_by(SubmissionTraining.training_id).all()
        assert [item.training_id for item in selected] == ["s1", "s2"]

    edit_html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=trainings").get_data(as_text=True)
    token = edit_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    close_response = admin_client.post(
        f"/admin/forms/{form_id}/training-selection/toggle",
        data={"csrf_token": token},
    )
    assert close_response.status_code == 302
    closed_save = admin_client.post(
        "/submissions/training-declaration-1/trainings?token=training-secret",
        data={"selected_trainings": ["s1"]},
    )
    assert closed_save.status_code == 409
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        assert form.training_selection_open is False
        selected = db.query(SubmissionTraining).all()
        assert {item.training_id for item in selected} == {"s1", "s2"}


def test_declaration_download_disabled_before_declaration_form_is_completed(admin_app, admin_client):
    create_form(
        admin_app,
        slug="declaration_flow",
        name="Declaration Flow",
        definition_json={
            "title": "Declaration Flow",
            "fields": [],
            "documents": [
                {
                    "id": "declaration",
                    "label": "Deklaracja",
                    "kind": "generated_pdf",
                    "enabled": True,
                    "template_html": "<p>Deklaracja</p>",
                    "fields": [{"type": "radio", "name": "deklaracja_18_lat", "label": "18 lat", "options": ["Tak", "Nie"]}],
                }
            ],
        },
    )
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        db.add(
            FormSubmission(
                submission_id="decl-1",
                form_slug="declaration_flow",
                form_name="Declaration Flow",
                officer_decision="accepted",
                acceptance_required="Tak",
                declaration_required="Tak",
                declaration_generated="",
                declaration_filename="",
                agreement_required="Nie",
            )
        )
        db.commit()

    response = admin_client.get("/do-podpisania?submission_id=decl-1")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "/declaration/declaration_flow/decl-1" in html
    assert "/downloads/pdfs/declaration_flow" not in html
    assert "Pobierz deklarac" in html
    assert "disabled" in html
    with session_factory() as db:
        submission = db.query(FormSubmission).filter_by(submission_id="decl-1").one()
        assert submission.declaration_generated != "Tak"


def test_submission_delete_button_is_hidden_until_selection(admin_app, admin_client):
    user_id = create_user(admin_app)
    form_id = create_form(admin_app, slug="delete_ui", name="Delete UI", user_id=user_id)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(FormSubmission(submission_id="abc", form_slug="delete_ui", form_name="Delete UI"))
        db.commit()
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/submissions").get_data(as_text=True)

    assert "Zgłoszenia formularza" in html
    assert "data-delete-button hidden disabled" in html
    assert "Czy na pewno chcesz usunąć zaznaczone zgłoszenia? Tej operacji nie można cofnąć." in html


def test_bulk_delete_removes_submission_related_rows(admin_app, admin_client):
    user_id = create_user(admin_app)
    form_id = create_form(admin_app, slug="delete_rows", name="Delete Rows", user_id=user_id)
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(submission_id="abc", form_slug="delete_rows", form_name="Delete Rows")
        db.add(submission)
        db.flush()
        db.add(SubmissionFile(submission_id=submission.id, public_submission_id="abc", form_slug="delete_rows", filename="a.pdf", storage_path="x"))
        db.add(SubmissionDecision(submission_id=submission.id, public_submission_id="abc", form_slug="delete_rows", decision="accepted"))
        db.add(SubmissionWorkflowEvent(submission_id=submission.id, public_submission_id="abc", form_slug="delete_rows", new_status="FORM_SUBMITTED"))
        db.add(EmailLog(submission_id=submission.id, public_submission_id="abc", status="sent"))
        db.commit()
        submission_pk = submission.id
    login(admin_client)
    with admin_client.session_transaction() as session:
        token = session["admin_csrf_token"]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/delete-selected",
        data={"csrf_token": token, "submission_pk_ids": str(submission_pk)},
    )

    assert response.status_code == 302
    with session_factory() as db:
        assert db.query(FormSubmission).count() == 0
        assert db.query(SubmissionFile).count() == 0
        assert db.query(SubmissionDecision).count() == 0
        assert db.query(SubmissionWorkflowEvent).count() == 0
        assert db.query(EmailLog).count() == 0


def test_regular_admin_cannot_delete_submission_without_form_permission(admin_app, admin_client):
    create_user(admin_app, email="regular@example.com", role="admin")
    form_id = create_form(admin_app, slug="blocked_delete", name="Blocked Delete")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        submission = FormSubmission(submission_id="abc", form_slug="blocked_delete", form_name="Blocked Delete")
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client, email="regular@example.com")
    with admin_client.session_transaction() as session:
        token = session["admin_csrf_token"]

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/delete-selected",
        data={"csrf_token": token, "submission_pk_ids": str(submission_pk)},
    )

    assert response.status_code == 403


def test_form_regulation_upload_and_public_link(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="regulated",
        name="Regulated",
        definition_json={
            "title": "Regulated",
            "fields": [
                {
                    "type": "checkbox",
                    "name": "accept_regulamin",
                    "label": "Regulamin",
                    "required": True,
                    "options": [{"value": "Tak", "label": "Akceptuję regulamin"}],
                }
            ],
        },
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            FormField(
                form_id=form_id,
                name="accept_regulamin",
                label="Regulamin",
                type="checkbox",
                required=True,
                options=[{"value": "Tak", "label": "Akceptuję regulamin"}],
                active=True,
            )
        )
        db.commit()
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    assert "Ten formularz nie ma wgranego regulaminu" in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Regulated",
            "slug": "regulated",
            "title": "Regulated",
            "sort_order": "0",
            "workflow_json": "{}",
            "is_active": "on",
            "is_public": "on",
            "regulation_file": (io.BytesIO(b"%PDF-1.4\nregulamin"), "regulamin.pdf"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    page = admin_client.get("/form/regulated").get_data(as_text=True)
    assert '<a href="/form/regulated/regulamin" target="_blank" rel="noopener noreferrer">regulamin</a>' in page
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(FormRegulation).one().original_filename == "regulamin.pdf"


def test_public_form_does_not_render_dead_regulation_link(admin_app, admin_client):
    form_id = create_form(
        admin_app,
        slug="no_regulation",
        name="No Regulation",
        definition_json={
            "title": "No Regulation",
            "fields": [
                {
                    "type": "checkbox",
                    "name": "accept_regulamin",
                    "required": True,
                    "options": [{"value": "Tak", "label": "Akceptuję regulamin"}],
                }
            ],
        },
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            FormField(
                form_id=form_id,
                name="accept_regulamin",
                label="Regulamin",
                type="checkbox",
                required=True,
                options=[{"value": "Tak", "label": "Akceptuję regulamin"}],
                active=True,
            )
        )
        db.commit()

    html = admin_client.get("/form/no_regulation").get_data(as_text=True)

    assert "/form/no_regulation/regulamin" not in html
    assert "Akceptuję regulamin" in html


def test_super_admin_can_edit_contact_page_and_regular_admin_cannot(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/site/contact").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/site/contact",
        data={
            "csrf_token": token,
            "title": "Kontakt",
            "content_html": "<p>Treść kontaktowa</p>",
            "contact_details": "Dane kontaktowe",
            "email": "kontakt@example.com",
            "phone": "123",
        },
    )

    assert response.status_code == 302
    assert "Treść kontaktowa" in admin_client.get("/kontakt").get_data(as_text=True)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(ContactPage).one().email == "kontakt@example.com"

    create_user(admin_app, email="regular2@example.com", role="admin")
    admin_client.get("/admin/logout")
    login(admin_client, email="regular2@example.com")
    blocked = admin_client.get("/admin/site/contact")
    assert blocked.status_code == 403
    assert "Nie masz uprawnień do edycji tej strony." in blocked.get_data(as_text=True)


def test_contact_menu_footer_and_multiple_phones(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/site/contact").get_data(as_text=True)
    assert "Dodaj telefon" in html
    assert "Usuń telefon" in html
    assert "Podgląd" in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/site/contact",
        data={
            "csrf_token": token,
            "title": "Kontakt",
            "content_html": "<h2>Treść kontaktowa</h2><script>alert(1)</script><p><strong>Biuro</strong></p>",
            "address": "ul. Testowa 1\n65-001 Zielona Góra",
            "email": "kontakt@example.com",
            "phone_label": ["sekretariat", "infolinia", "fax"],
            "phone_number": ["111", "222", "333"],
        },
    )

    assert response.status_code == 302
    contact_html = admin_client.get("/kontakt").get_data(as_text=True)
    assert "Treść kontaktowa" in contact_html
    assert "alert(1)" not in contact_html
    assert "sekretariat 111" in contact_html
    assert "infolinia 222" in contact_html
    assert "fax 333" in contact_html
    assert "ul. Testowa 1" in contact_html
    assert 'href="/kontakt"' in contact_html
    assert 'class="is-active"' in contact_html

    footer_html = admin_client.get("/").get_data(as_text=True)
    assert "Strona kontaktowa" not in footer_html
    assert "sekretariat 111" in footer_html
    assert "infolinia 222" in footer_html
    assert "fax 333" in footer_html
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        page = db.query(ContactPage).one()
        assert page.email == "kontakt@example.com"
        assert [item["number"] for item in page.phones] == ["111", "222", "333"]

    html = admin_client.get("/admin/site/contact").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        "/admin/site/contact",
        data={
            "csrf_token": token,
            "title": "Kontakt",
            "content_html": "<p>Po zmianie</p>",
            "address": "",
            "email": "",
            "phone_label": ["sekretariat", ""],
            "phone_number": ["111", ""],
        },
    )
    assert response.status_code == 302
    html_after_remove = admin_client.get("/kontakt").get_data(as_text=True)
    assert "sekretariat 111" in html_after_remove
    assert "infolinia 222" not in html_after_remove
    assert "fax 333" not in html_after_remove
    assert 'href="mailto:' not in html_after_remove


def test_service_documents_show_links_in_footer(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/site/documents").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/site/documents",
        data={
            "csrf_token": token,
            "document_type": "privacy",
            "title": "Polityka prywatności",
            "content_html": "<p>Prywatność</p>",
            "document_file": (io.BytesIO(b"%PDF-1.4\nprivacy"), "privacy.pdf"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 302
    footer_html = admin_client.get("/").get_data(as_text=True)
    assert "Polityka prywatności" in footer_html
    assert "/dokumenty/privacy" in footer_html
    assert "Regulamin serwisu" not in footer_html
    doc_html = admin_client.get("/dokumenty/privacy").get_data(as_text=True)
    assert "Prywatność" in doc_html
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(ServiceDocument).one().original_filename == "privacy.pdf"


def test_footer_contains_creator_credit_but_mail_footer_does_not(admin_app, admin_client):
    create_form(admin_app, slug="credit", name="Credit")

    html = admin_client.get("/").get_data(as_text=True)

    assert "Created by Witold Grzesiak" in html
    assert "Created by Witold Grzesiak" not in admin_app.extensions["services"].mail_dispatch_service.build_footer(None)


def test_public_footer_has_no_hardcoded_or_form_logo_fallback(admin_app, admin_client):
    form_id = create_form(admin_app, slug="footer-logo-separation", name="Footer logo separation")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        form.logo_id = None
        db.commit()

    html = admin_client.get("/").get_data(as_text=True)
    footer_html = html.split('<footer class="site-footer">', 1)[1].split("</footer>", 1)[0]

    assert "images/logo.png" not in footer_html
    assert "site-brand__logo" not in footer_html
    assert "form-logo" not in footer_html


def test_mail_footer_uses_form_footer_then_global_fallback(admin_app, caplog):
    from routes.admin.mail import select_default_footer

    form_id = create_form(admin_app, slug="footer_form", name="Footer Form")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form_footer = MailFooter(form_id=form_id, name="Form", html_body="<p>Form</p>", is_active=True, is_default=True)
        global_footer = MailFooter(form_id=None, name="Global", html_body="<p>Global</p>", is_active=True, is_default=True)
        db.add_all([form_footer, global_footer])
        db.commit()
        caplog.clear()
        with admin_app.app_context():
            assert select_default_footer([global_footer, form_footer], form_id=form_id).name == "Form"
            assert select_default_footer([global_footer], form_id=form_id).name == "Global"
            assert select_default_footer([], form_id=form_id) is None

    assert "scope=form" in caplog.text
    assert "scope=global" in caplog.text
    assert "scope=none" in caplog.text


def test_mail_footer_resolver_enforces_global_initial_mail_and_form_process_footer(admin_app):
    from services.mail_footer_resolver import MailFooterResolver

    form_id = create_form(admin_app, slug="resolver_form", name="Resolver form")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        form_footer = MailFooter(form_id=form_id, name="Form", html_body="<p>Form</p>", is_active=True, is_default=True)
        global_footer = MailFooter(form_id=None, name="Global", html_body="<p>Global</p>", is_active=True, is_default=True)
        db.add_all([form_footer, global_footer])
        db.flush()
        resolver = MailFooterResolver()

        assert resolver.resolve(db, mail_type="agreement_generated", form=form).name == "Form"
        assert resolver.resolve(db, mail_type="submission_received", form=form).name == "Global"
        form_footer.use_global = True
        assert resolver.resolve(db, mail_type="agreement_uploaded", form=form).name == "Global"
        form_footer.use_global = False
        form_footer.is_active = False
        assert resolver.resolve(db, mail_type="agreement_uploaded", form=form).name == "Global"


def test_global_and_form_footer_editors_render_preview_and_sanitize_html(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="footer_preview", name="Footer preview")
    login(admin_client)
    global_html = admin_client.get("/admin/mail-footer").get_data(as_text=True)
    assert "Podgląd stopki e-mail" in global_html
    assert "Przykładowa wiadomość" in global_html
    assert 'name="logo_width"' in global_html
    assert 'name="logo_height"' in global_html
    assert 'name="logo_position"' in global_html
    assert "{{ footer_logo }}" in global_html
    assert "Logo w stopce jest niezależne od logo formularza i logo platformy." in global_html
    token = global_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(
        "/admin/mail-footer",
        data={
            "csrf_token": token,
            "name": "Ogólna",
            "html_body": '<p onclick="x"><strong>Kontakt</strong><img src=x></p><script>alert(1)</script>',
            "contact_html": "<p>kontakt@example.com</p>",
            "links_text": "Serwis|https://example.com",
            "legal_text": "<p>Tekst prawny</p>",
            "logo_alignment": "center",
            "logo_position": "right",
            "logo_width": "260",
            "logo_height": "90",
            "is_active": "on",
        },
    )
    assert response.status_code == 302
    form_html = admin_client.get(f"/admin/forms/{form_id}/mail-footers/new").get_data(as_text=True)
    assert "Używana jest stopka formularza" in form_html
    assert "Mail początkowy po rejestracji wniosku zawsze używa stopki ogólnej" in form_html
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        footer = db.query(MailFooter).filter(MailFooter.form_id.is_(None)).one()
        assert footer.html_body == "<p><strong>Kontakt</strong></p>"
        assert footer.logo_alignment == "center"
        assert footer.logo_position == "right"
        assert footer.logo_width == 260
        assert footer.logo_height == 90
        assert footer.links == [{"label": "Serwis", "url": "https://example.com"}]


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("logo_width", "19", "Szerokość logo musi mieć wartość od 20 do 800 px."),
        ("logo_width", "801", "Szerokość logo musi mieć wartość od 20 do 800 px."),
        ("logo_width", "duże", "Szerokość logo musi być liczbą całkowitą od 20 do 800 px."),
        ("logo_height", "19", "Wysokość logo musi mieć wartość od 20 do 400 px."),
        ("logo_height", "401", "Wysokość logo musi mieć wartość od 20 do 400 px."),
        ("logo_height", "wysokie", "Wysokość logo musi być liczbą całkowitą od 20 do 400 px."),
        ("logo_alignment", "justify", "Wybierz prawidłowe wyrównanie logo."),
        ("logo_position", "floating", "Wybierz prawidłowe położenie logo."),
    ],
)
def test_mail_footer_rejects_invalid_logo_layout(admin_app, admin_client, field, value, message):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/mail-footer").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    data = {
        "csrf_token": token,
        "name": "Stopka",
        "html_body": "<p>Treść</p>",
        "logo_alignment": "left",
        "logo_position": "top",
        "logo_width": "",
        "logo_height": "",
        "is_active": "on",
    }
    data[field] = value

    response = admin_client.post("/admin/mail-footer", data=data)

    assert response.status_code == 400
    assert message in response.get_data(as_text=True)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(MailFooter).count() == 0


def test_form_admin_can_save_footer_logo_size_and_position(admin_app, admin_client):
    manager_id = create_user(admin_app, email="footer-manager@example.com", role="admin")
    form_id = create_form(admin_app, slug="footer_layout", user_id=manager_id)
    login(admin_client, email="footer-manager@example.com")
    html = admin_client.get(f"/admin/forms/{form_id}/mail-footers/new").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/mail-footers/new",
        data={
            "csrf_token": token,
            "name": "Stopka formularza",
            "html_body": "<p>Treść {{ footer_logo }}</p>",
            "logo_alignment": "right",
            "logo_position": "inline",
            "logo_width": "320",
            "logo_height": "",
            "is_active": "on",
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        footer = db.query(MailFooter).filter(MailFooter.form_id == form_id).one()
        assert footer.logo_position == "inline"
        assert footer.logo_alignment == "right"
        assert footer.logo_width == 320
        assert footer.logo_height is None
        assert "{{ footer_logo }}" in footer.html_body


def test_form_manager_saves_custom_smtp_with_encrypted_hidden_password(admin_app, admin_client):
    manager_id = create_user(admin_app, email="manager@example.com", role="admin")
    form_id = create_form(admin_app, slug="mail_form", name="Mail form", user_id=manager_id)
    login(admin_client, email="manager@example.com")
    edit_html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = edit_html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Mail form",
            "slug": "mail_form",
            "title": "Mail form",
            "is_active": "on",
            "is_public": "on",
            "user_instruction_config": '{"stages": []}',
            "mail_mode": "custom",
            "form_smtp_host": "smtp.form.test",
            "form_smtp_port": "465",
            "form_smtp_user": "form-user",
            "form_smtp_password": "form-secret",
            "form_smtp_mail_from": "form@example.com",
            "form_smtp_sender_name": "Form sender",
            "form_smtp_reply_to": "reply@example.com",
            "form_smtp_timeout": "20",
            "form_smtp_use_ssl": "on",
        },
    )

    assert response.status_code == 302
    service = admin_app.extensions["services"].mail_settings_service
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        db.add(
            SystemMailSettings(
                smtp_config={"host": "smtp.system.test", "mail_from": "system@example.com"},
                layout_config={},
            )
        )
        db.flush()
        assert form.mail_mode == "custom"
        assert form.smtp_config["host"] == "smtp.form.test"
        assert form.smtp_password_encrypted != "form-secret"
        assert service.decrypt_password(form.smtp_password_encrypted) == "form-secret"
        resolved = service.resolve_smtp(db, form, admin_app.config)
        assert resolved["smtp_host"] == "smtp.form.test"
        assert resolved["smtp_password"] == "form-secret"

    edit_html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    assert "form-secret" not in edit_html


def test_system_mail_settings_are_superadmin_only_and_password_is_hidden(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/mail-settings").get_data(as_text=True)
    assert "Położenie logo w mailu" in html
    assert "Wyrównanie logo" in html
    assert "Wysokość logo" in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/mail-settings",
        data={
            "csrf_token": token,
            "system_smtp_host": "smtp.system.test",
            "system_smtp_port": "587",
            "system_smtp_user": "system-user",
            "system_smtp_password": "system-secret",
            "system_smtp_mail_from": "system@example.com",
            "system_smtp_sender_name": "System sender",
            "system_smtp_reply_to": "help@example.com",
            "system_smtp_timeout": "30",
            "system_smtp_use_tls": "on",
            "template_is_active": "on",
            "template_subject": "Otrzymano {{ submission_id }}",
            "template_html_body": "<p>Witaj {{ imie }}</p><script>alert(1)</script>",
            "template_text_body": "Witaj {{ imie }}",
            "layout_platform_name": "Moja platforma",
            "layout_primary_color": "#123456",
            "layout_accent_color": "#abcdef",
            "layout_footer_html": "<p>Stopka</p>",
            "layout_logo_position": "before_content",
            "layout_logo_alignment": "right",
            "layout_logo_height_px": "92",
        },
    )

    assert response.status_code == 302
    service = admin_app.extensions["services"].mail_settings_service
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        settings = db.query(SystemMailSettings).one()
        template = db.query(PlatformMailTemplate).filter_by(template_type="submission_received").one()
        assert settings.smtp_config["host"] == "smtp.system.test"
        assert settings.smtp_password_encrypted != "system-secret"
        assert service.decrypt_password(settings.smtp_password_encrypted) == "system-secret"
        assert settings.layout_config["platform_name"] == "Moja platforma"
        assert settings.layout_config["logo_position"] == "before_content"
        assert settings.layout_config["logo_alignment"] == "right"
        assert settings.layout_config["logo_height_px"] == 92
        assert template.is_active is True
        assert template.subject == "Otrzymano {{ submission_id }}"
        assert "<script" not in template.html_body

    assert "system-secret" not in admin_client.get("/admin/mail-settings").get_data(as_text=True)
    admin_client.get("/admin/logout")
    create_user(admin_app, email="regular@example.com", role="admin")
    login(admin_client, email="regular@example.com")
    assert admin_client.get("/admin/mail-settings").status_code == 403


def test_superadmin_deletes_used_logo_with_safe_detach_and_regular_admin_is_blocked(admin_app, admin_client):
    create_user(admin_app)
    logo_path = Path(admin_app.config["TEMP_DIR"]) / "kept-logo.png"
    logo_path.write_bytes(b"logo-data")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(name="Used logo", filename="kept-logo.png", storage_path=str(logo_path), mime_type="image/png", active=True)
        db.add(logo)
        db.flush()
        form = Form(slug="logo_form", name="Logo form", title="Logo form", definition_json={"fields": []}, logo_id=logo.id)
        db.add(form)
        db.add(SystemMailSettings(smtp_config={}, layout_config={"logo_id": logo.id}))
        db.commit()
        logo_id = logo.id
        form_id = form.id

    login(admin_client)
    html = admin_client.get("/admin/logos").get_data(as_text=True)
    assert "Przypisane formularze: 1" in html
    assert f"/admin/logos/{logo_id}/delete" in html
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    response = admin_client.post(f"/admin/logos/{logo_id}/delete", data={"csrf_token": token, "detach": "1"})
    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.get(Logo, logo_id) is None
        assert db.get(Form, form_id).logo_id is None
        assert db.query(SystemMailSettings).one().layout_config["logo_id"] is None
    assert logo_path.exists()

    blocked_path = Path(admin_app.config["TEMP_DIR"]) / "blocked-logo.png"
    blocked_path.write_bytes(b"logo-data")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        blocked_logo = Logo(name="Blocked", filename="blocked-logo.png", storage_path=str(blocked_path), mime_type="image/png", active=True)
        db.add(blocked_logo)
        db.commit()
        blocked_logo_id = blocked_logo.id
    admin_client.get("/admin/logout")
    create_user(admin_app, email="regular@example.com", role="admin")
    login(admin_client, email="regular@example.com")
    with admin_client.session_transaction() as session:
        token = session["admin_csrf_token"]
    assert admin_client.post(
        f"/admin/logos/{blocked_logo_id}/delete",
        data={"csrf_token": token, "detach": "1"},
    ).status_code == 403
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.get(Logo, blocked_logo_id) is not None


def test_submission_received_uses_global_template_layout_and_logs_missing_recipient(admin_app):
    form_id = create_form(admin_app, slug="receipt_form", name="Receipt form", mail_mode="system")
    logo_path = Path(admin_app.config["TEMP_DIR"]) / "mail-logo.png"
    logo_path.write_bytes(b"inline-logo")
    footer_logo_path = Path(admin_app.config["TEMP_DIR"]) / "configured-footer-logo.png"
    footer_logo_path.write_bytes(b"footer-logo")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(
            name="Platform logo",
            filename="mail-logo.png",
            storage_path=str(logo_path),
            mime_type="image/png",
            active=True,
        )
        footer_logo = Logo(
            name="Configured footer logo",
            filename="configured-footer-logo.png",
            storage_path=str(footer_logo_path),
            mime_type="image/png",
            active=True,
        )
        db.add_all([logo, footer_logo])
        db.flush()
        db.add(MailFooter(
            form_id=None,
            name="Global footer",
            html_body="<p>Configured footer</p>",
            logo_id=footer_logo.id,
            logo_position="right",
            is_active=True,
            is_default=True,
        ))
        db.add(
            SystemMailSettings(
                smtp_config={
                    "host": "smtp.system.test",
                    "port": 587,
                    "user": "",
                    "mail_from": "system@example.com",
                    "sender_name": "System sender",
                    "use_tls": True,
                    "use_ssl": False,
                    "timeout": 30,
                    "reply_to": "help@example.com",
                },
                layout_config={
                    "platform_name": "Global brand",
                    "primary_color": "#123456",
                    "accent_color": "#abcdef",
                    "footer_html": "<p>Global footer</p>",
                    "logo_id": logo.id,
                    "logo_position": "footer",
                    "logo_alignment": "center",
                    "logo_height_px": 64,
                },
            )
        )
        db.add(
            PlatformMailTemplate(
                template_type="submission_received",
                name="Receipt",
                subject="Odebrano {{ submission_id }}",
                html_body="<p>Witaj {{ imie }}</p>",
                text_body="Witaj {{ imie }}",
                is_active=True,
            )
        )
        db.add_all(
            [
                FormSubmission(
                    submission_id="receipt-1",
                    form_slug="receipt_form",
                    form_name="Receipt form",
                    email="jan@example.com",
                    imiona="Jan",
                    nazwisko="Kowalski",
                    data_json={"imiona": "Jan"},
                ),
                FormSubmission(
                    submission_id="receipt-no-email",
                    form_slug="receipt_form",
                    form_name="Receipt form",
                    email="",
                ),
            ]
        )
        db.commit()

    sent = []
    dispatch = admin_app.extensions["services"].mail_dispatch_service
    dispatch.smtp_sender = lambda **kwargs: sent.append(kwargs)
    with admin_app.test_request_context("/"):
        result = dispatch.dispatch_submission_received("receipt-1")
        missing = dispatch.dispatch_submission_received("receipt-no-email")

    assert result.status == "sent"
    assert missing.status == "skipped"
    assert len(sent) == 1
    assert sent[0]["smtp_host"] == "smtp.system.test"
    assert sent[0]["subject"] == "Odebrano receipt-1"
    assert "Witaj Jan" in sent[0]["html_body"]
    assert "Global brand" in sent[0]["html_body"]
    assert "Configured footer" in sent[0]["html_body"]
    assert "#123456" in sent[0]["html_body"]
    assert 'src="cid:platform-logo-' not in sent[0]["html_body"]
    assert 'data-logo-position="footer"' not in sent[0]["html_body"]
    assert 'src="cid:footer-logo"' in sent[0]["html_body"]
    assert "mail-logo.png" not in sent[0]["html_body"]
    assert "configured-footer-logo.png" not in sent[0]["html_body"]
    assert sent[0]["html_body"].count('src="cid:footer-logo"') == 1
    assert len(sent[0]["inline_images"]) == 1
    assert sent[0]["inline_images"][0] == {
        "cid": "footer-logo",
        "content": b"footer-logo",
        "mime_type": "image/png",
        "filename": "configured-footer-logo.png",
    }
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        sent_log = db.query(EmailLog).filter_by(public_submission_id="receipt-1").one()
        missing_log = db.query(EmailLog).filter_by(public_submission_id="receipt-no-email").one()
        assert sent_log.status == "sent"
        assert missing_log.status == "skipped"
        assert missing_log.error_message == "Brak odbiorcy."


def test_submission_received_without_smtp_is_logged_and_does_not_fail(admin_app):
    create_form(
        admin_app,
        slug="no_smtp_form",
        name="No SMTP form",
        mail_mode="custom",
        smtp_config={"host": "", "mail_from": ""},
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(
            PlatformMailTemplate(
                template_type="submission_received",
                name="Receipt",
                subject="Odebrano {{ submission_id }}",
                html_body="<p>{{ submission_id }}</p>",
                is_active=True,
            )
        )
        db.add(
            FormSubmission(
                submission_id="no-smtp",
                form_slug="no_smtp_form",
                form_name="No SMTP form",
                email="jan@example.com",
            )
        )
        db.commit()

    sent = []
    dispatch = admin_app.extensions["services"].mail_dispatch_service
    dispatch.smtp_sender = lambda **kwargs: sent.append(kwargs)
    with admin_app.test_request_context("/"):
        result = dispatch.dispatch_submission_received("no-smtp")

    assert result.status == "skipped"
    assert result.error_message == "Brak konfiguracji SMTP."
    assert sent == []
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        log = db.query(EmailLog).filter_by(public_submission_id="no-smtp").one()
        assert log.status == "skipped"
        assert log.error_message == "Brak konfiguracji SMTP."


def test_submission_received_omits_unavailable_logo_without_failing(admin_app):
    create_form(admin_app, slug="missing_logo_form", name="Missing logo form", mail_mode="system")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(
            name="Missing logo",
            filename="missing.png",
            storage_path=str(Path(admin_app.config["TEMP_DIR"]) / "does-not-exist.png"),
            mime_type="image/png",
            active=True,
        )
        db.add(logo)
        db.flush()
        db.add(SystemMailSettings(
            smtp_config={"host": "smtp.test", "mail_from": "sender@example.com"},
            layout_config={
                "logo_id": logo.id,
                "logo_position": "footer",
                "logo_alignment": "center",
                "logo_height_px": 64,
                "footer_html": "<p>Footer remains</p>",
            },
        ))
        db.add(PlatformMailTemplate(
            template_type="submission_received",
            name="Receipt",
            subject="Receipt {{ submission_id }}",
            html_body="<p>Body remains</p>",
            is_active=True,
        ))
        db.add(FormSubmission(
            submission_id="missing-logo",
            form_slug="missing_logo_form",
            form_name="Missing logo form",
            email="jan@example.com",
        ))
        db.commit()

    sent = []
    dispatch = admin_app.extensions["services"].mail_dispatch_service
    dispatch.smtp_sender = lambda **kwargs: sent.append(kwargs)
    with admin_app.test_request_context("/"):
        result = dispatch.dispatch_submission_received("missing-logo")

    assert result.status == "sent"
    assert len(sent) == 1
    assert "<img" not in sent[0]["html_body"]
    assert "Body remains" in sent[0]["html_body"]
    assert "Footer remains" in sent[0]["html_body"]
    assert sent[0]["inline_images"] == []


def test_submission_received_omits_missing_mail_footer_logo_and_logs_warning(admin_app, caplog):
    form_id = create_form(admin_app, slug="missing_footer_logo", name="Missing footer logo", mail_mode="system")
    missing_path = Path(admin_app.config["TEMP_DIR"]) / "missing-footer.png"
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        logo = Logo(
            name="Missing footer logo",
            filename="missing-footer.png",
            storage_path=str(missing_path),
            mime_type="image/png",
            active=True,
        )
        db.add(logo)
        db.flush()
        logo_id = logo.id
        db.add(MailFooter(
            form_id=None,
            name="Global footer",
            html_body="<p>Stopka pozostaje</p>",
            logo_id=logo.id,
            is_active=True,
            is_default=True,
        ))
        db.add(SystemMailSettings(
            smtp_config={"host": "smtp.test", "mail_from": "sender@example.com"},
            layout_config={},
        ))
        db.add(PlatformMailTemplate(
            template_type="submission_received",
            name="Receipt",
            subject="Receipt",
            html_body="<p>Treść wiadomości</p>",
            is_active=True,
        ))
        db.add(FormSubmission(
            submission_id="missing-footer-logo",
            form_slug="missing_footer_logo",
            form_name="Missing footer logo",
            email="jan@example.com",
        ))
        db.commit()

    sent = []
    dispatch = admin_app.extensions["services"].mail_dispatch_service
    dispatch.smtp_sender = lambda **kwargs: sent.append(kwargs)
    with caplog.at_level(logging.WARNING), admin_app.test_request_context("/"):
        result = dispatch.dispatch_submission_received("missing-footer-logo")

    assert result.status == "sent"
    assert "<img" not in sent[0]["html_body"]
    assert "Stopka pozostaje" in sent[0]["html_body"]
    assert sent[0]["inline_images"] == []
    assert (
        f"Nie udało się załadować logo stopki mailowej: logo_id={logo_id}, logo_path="
        in caplog.text
    )


def _create_submission_waiting_for_agreement_review(admin_app, form_slug: str, submission_id: str = "agreement-review"):
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id=submission_id,
            form_slug=form_slug,
            form_name="Agreement review",
            email="beneficiary@example.com",
            officer_decision="accepted",
            declaration_required="Tak",
            declaration_signed="Tak",
            declaration_signature_valid="Tak",
            agreement_required="Tak",
            agreement_generated="Tak",
            agreement_signed="Tak",
            agreement_signature_valid="Tak",
            agreement_signed_filename="agreement-signed.pdf",
            process_status="AGREEMENT_WAITING_FOR_OFFICE_SIGNATURE",
        )
        db.add(submission)
        db.flush()
        db.add(
            SubmissionFile(
                submission_id=submission.id,
                public_submission_id=submission.submission_id,
                form_slug=form_slug,
                document_id="agreement",
                document_type="signed_agreement",
                filename="agreement-signed.pdf",
                storage_path=f"output/{form_slug}/agreements/signed/agreement-signed.pdf",
                signed=True,
                status="signed",
            )
        )
        db.commit()
        return submission.id


def test_admin_shows_application_decision_only_during_review_and_agreement_decision_after_upload(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="agreement_review_form", name="Agreement review")
    submission_pk = _create_submission_waiting_for_agreement_review(admin_app, "agreement_review_form")
    login(admin_client)

    response = admin_client.get(
        f"/admin/forms/{form_id}/submissions/{submission_pk}",
        environ_overrides={"SCRIPT_NAME": "/aplikacja"},
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Decyzja o zaakceptowaniu wniosku" not in html
    assert "Umowa podpisana przez urząd" in html
    assert "Zaznacz, czy umowa została podpisana po stronie urzędu" in html
    assert "Umowa podpisana przez beneficjenta" not in html
    assert f'/aplikacja/admin/forms/{form_id}/submissions/{submission_pk}/beneficiary-agreement-decision' in html


def test_admin_confirms_uploaded_agreement_and_finishes_process(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="agreement_confirm_form", name="Agreement confirm")
    submission_pk = _create_submission_waiting_for_agreement_review(admin_app, "agreement_confirm_form", "agreement-confirm")
    login(admin_client)

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/beneficiary-agreement-decision",
        data={
            "csrf_token": admin_csrf(admin_client),
            "agreement_decision": "accepted",
            "agreement_decision_reason": "",
        },
    )

    assert response.status_code == 302
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.process_status == "PROCESS_COMPLETED"
        decision = db.query(SubmissionDecision).filter_by(public_submission_id="agreement-confirm").one()
        assert decision.decision == "office_agreement_accepted"
        assert decision.target_status == "AGREEMENT_SIGNED_BY_OFFICE"
        assert [event.source for event in db.query(SubmissionWorkflowEvent).order_by(SubmissionWorkflowEvent.id)] == [
            "agreement_signed_by_office",
            "process_completed",
        ]


def test_admin_rejects_uploaded_agreement_with_reason_and_blocks_decision_without_file(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="agreement_reject_form", name="Agreement reject")
    submission_pk = _create_submission_waiting_for_agreement_review(admin_app, "agreement_reject_form", "agreement-reject")
    login(admin_client)

    rejected = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/beneficiary-agreement-decision",
        data={
            "csrf_token": admin_csrf(admin_client),
            "agreement_decision": "rejected",
            "agreement_decision_reason": "Brakuje podpisu na ostatniej stronie.",
        },
    )
    assert rejected.status_code == 302

    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = db.get(FormSubmission, submission_pk)
        assert submission.process_status == "AGREEMENT_REJECTED_BY_OFFICE"
        assert submission.agreement_signature_valid == ""
        assert db.query(SubmissionDecision).one().justification == "Brakuje podpisu na ostatniej stronie."

        blocked = FormSubmission(
            submission_id="agreement-no-file",
            form_slug="agreement_reject_form",
            form_name="Agreement reject",
            agreement_required="Tak",
            agreement_signature_valid="Tak",
            process_status="AGREEMENT_WAITING_FOR_OFFICE_SIGNATURE",
        )
        db.add(blocked)
        db.commit()
        blocked_pk = blocked.id

    blocked_response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{blocked_pk}/beneficiary-agreement-decision",
        data={"csrf_token": admin_csrf(admin_client), "agreement_decision": "accepted"},
        follow_redirects=True,
    )
    assert blocked_response.status_code == 200
    assert "Nie znaleziono wgranej podpisanej umowy" in blocked_response.get_data(as_text=True)


def test_application_decision_backend_rejects_second_acceptance_after_declaration_stage(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="application_locked_form", name="Application locked")
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        submission = FormSubmission(
            submission_id="application-locked",
            form_slug="application_locked_form",
            form_name="Application locked",
            officer_decision="accepted",
            declaration_required="Tak",
            declaration_signed="Tak",
            declaration_signature_valid="Tak",
            process_status="DECLARATION_SIGNED",
        )
        db.add(submission)
        db.commit()
        submission_pk = submission.id
    login(admin_client)

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/decision",
        data={"csrf_token": admin_csrf(admin_client), "officer_decision": "accepted"},
        follow_redirects=True,
    )

    assert "tylko na etapie jego weryfikacji" in response.get_data(as_text=True)
    with session_factory() as db:
        assert db.query(SubmissionDecision).count() == 0


def test_form_manager_with_form_permission_cannot_review_uploaded_agreement(admin_app, admin_client):
    email = "agreement-manager@example.com"
    manager_id = create_user(admin_app, email=email, role="form_manager")
    form_id = create_form(
        admin_app,
        slug="agreement_manager_form",
        name="Agreement manager",
        user_id=manager_id,
    )
    submission_pk = _create_submission_waiting_for_agreement_review(
        admin_app,
        "agreement_manager_form",
        "agreement-manager",
    )
    login(admin_client, email=email)

    response = admin_client.post(
        f"/admin/forms/{form_id}/submissions/{submission_pk}/beneficiary-agreement-decision",
        data={"csrf_token": admin_csrf(admin_client), "agreement_decision": "accepted"},
    )

    assert response.status_code == 403
    session_factory = create_session_factory(admin_app.config["DATABASE_URL"])
    with session_factory() as db:
        assert db.get(FormSubmission, submission_pk).process_status != "PROCESS_COMPLETED"
def test_workflow_builder_renders_readable_sections_and_legacy_labels(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="visual_workflow", definition_json=readable_workflow_definition())
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)

    for heading in (
        "Ustawienia procesu",
        "Etapy workflow",
        "Podgląd workflow",
        "Decyzje urzędnika",
        "Dokumenty wymagane w procesie",
        "Instrukcje dla użytkownika",
        "Powiadomienia e-mail",
    ):
        assert heading in html
    assert "Wniosek złożony" in html
    assert "Elementy zaawansowane" in html
    assert "legacy_extension" in html
    assert "data-add-workflow-step" in html
    assert "data-remove-workflow-step" in html
    assert "data-workflow-step-up" in html
    assert "data-edit-workflow-step" in html
    assert "data-edit-step-instruction" in html
    assert "data-workflow-step-editor hidden" in html
    assert "Opis etapu</span><textarea data-step-description data-rich-text" not in html
    assert "Pokaż jak zobaczy to użytkownik" in html
    assert "data-workflow-preview-list" in html
    assert "function renderWorkflowPreview()" in html
    assert "Brak kolejnego etapu." in html


def test_regular_admin_does_not_receive_advanced_json_editor(admin_app, admin_client):
    user_id = create_user(admin_app, role="admin")
    form_id = create_form(
        admin_app,
        slug="regular_admin_workflow",
        user_id=user_id,
        definition_json=readable_workflow_definition(),
    )
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)

    assert "Tryb zaawansowany — edycja JSON" not in html
    assert '<textarea name="workflow_json"' not in html
    assert "Etapy workflow" in html


def test_super_admin_receives_collapsed_advanced_json_editor(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="advanced_workflow", definition_json=readable_workflow_definition())
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)

    assert "Tryb zaawansowany — edycja JSON" in html
    assert "Zmiana JSON-a może uszkodzić workflow" in html
    assert '<details class="admin-section workflow-section workflow-json-advanced"' in html
    assert '<textarea name="workflow_json"' in html


def test_workflow_builder_saves_order_and_generates_user_instructions(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="save_visual_workflow", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    workflow = readable_workflow_definition()["workflow"]
    workflow["steps"] = [workflow["steps"][1], workflow["steps"][0], workflow["steps"][2]]
    workflow["initial_step"] = "submission"
    workflow["steps"][0]["description"] = "Urzędnik sprawdza dane."
    workflow["steps"][0]["next_action"] = "Poczekaj na wynik weryfikacji."

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Workflow Form",
            "slug": "save_visual_workflow",
            "title": "Workflow Form",
            "workflow_name": "Obsługa wniosku",
            "workflow_initial_step": "submission",
            "workflow_builder_json": json.dumps(workflow, ensure_ascii=False),
            "instruction_title": "Co dalej?",
            "user_instruction": "Sprawdź aktualny etap.",
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        assert [step["id"] for step in form.definition_json["workflow"]["steps"]] == [
            "review", "submission", "completed"
        ]
        assert form.definition_json["workflow"]["legacy_extension"] == {"preserve": True}
        assert form.user_instruction_config["stages"][0]["label"] == "Weryfikacja wniosku"
        assert form.user_instruction_config["stages"][0]["next_action"] == "Poczekaj na wynik weryfikacji."

    reopened = admin_client.get(f"/admin/forms/{form_id}/edit?tab=instructions").get_data(as_text=True)
    assert "Urzędnik sprawdza dane." in reopened
    assert "Poczekaj na wynik weryfikacji." in reopened
    assert "syncLinkedInstructionValues();" in reopened


def test_workflow_builder_preserves_manual_diagram_layout_after_save(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="workflow_manual_layout",
        definition_json=readable_workflow_definition(),
    )
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=workflow").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    workflow = readable_workflow_definition()["workflow"]
    workflow["diagram_layout"] = {
        "nodes": {
            "submission": {"x": 180, "y": 60},
            "decision:review:0": {"x": 760.5, "y": 220.25},
        }
    }

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "workflow",
            "name": "Workflow manual layout",
            "slug": "workflow_manual_layout",
            "title": "Workflow manual layout",
            "workflow_name": "Obsługa wniosku",
            "workflow_initial_step": workflow["initial_step"],
            "workflow_builder_json": json.dumps(workflow, ensure_ascii=False),
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        saved_layout = db.get(Form, form_id).definition_json["workflow"]["diagram_layout"]
        assert saved_layout == workflow["diagram_layout"]

    reopened = admin_client.get(f"/admin/forms/{form_id}/edit?tab=workflow").get_data(as_text=True)
    assert '"submission": {"x": 180.0, "y": 60.0}' in reopened
    assert '"decision:review:0": {"x": 760.5, "y": 220.25}' in reopened


def test_workflow_diagram_live_update_drag_zoom_and_reset_in_browser(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    definition = readable_workflow_definition()
    workflow = definition["workflow"]
    workflow["steps"].insert(
        2,
        {
            "id": "waiting_for_correction",
            "admin_label": "Korekta wniosku",
            "user_label": "Popraw wniosek",
            "status": "WAITING_FOR_CORRECTION",
            "next": "review",
            "stage_type": "correction",
        },
    )
    workflow["steps"].insert(
        3,
        {
            "id": "end_rejected",
            "admin_label": "Wniosek odrzucony",
            "user_label": "Wniosek odrzucony",
            "status": "OFFICER_REJECTED",
            "final": True,
            "rejected": True,
        },
    )
    workflow["decision_settings"] = [
        {
            "id": "application_decision",
            "label": "Decyzja o wniosku",
            "step_id": "review",
            "yes_status": "PROCESS_COMPLETED",
            "no_status": "OFFICER_REJECTED",
            "correction_status": "WAITING_FOR_CORRECTION",
            "active": True,
        }
    ]
    workflow["steps"][1]["decisions"] = {
        "accepted": "completed",
        "rejected": "end_rejected",
        "correction": "waiting_for_correction",
    }
    form_id = create_form(
        admin_app,
        slug="workflow_diagram_browser",
        definition_json=definition,
    )
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=workflow").get_data(as_text=True)

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page(viewport={"width": 1400, "height": 900})
        script_errors = []
        console_warnings = []
        page.on("pageerror", lambda exception: script_errors.append(str(exception)))
        page.on(
            "console",
            lambda message: console_warnings.append(message.text)
            if message.type == "warning"
            else None,
        )
        page.set_content(html, wait_until="domcontentloaded")
        page.wait_for_function(
            "document.querySelectorAll('[data-diagram-node-id]').length === 6"
        )

        positions = page.locator("[data-diagram-node-id]").evaluate_all(
            """(nodes) => Object.fromEntries(nodes.map((node) => {
                const matrix = node.transform.baseVal.consolidate().matrix;
                return [node.dataset.diagramNodeId, {x: matrix.e, y: matrix.f}];
            }))"""
        )
        assert len({(position["x"], position["y"]) for position in positions.values()}) == 6
        assert positions["waiting_for_correction"]["x"] < positions["review"]["x"]
        assert positions["decision:review:application_decision"]["x"] > positions["review"]["x"]
        assert positions["end_rejected"]["x"] > positions["decision:review:application_decision"]["x"]
        assert abs(
            positions["decision:review:application_decision"]["y"] - positions["review"]["y"]
        ) < 10
        assert positions["completed"]["y"] > positions["review"]["y"]
        decision_nodes = page.locator('[data-diagram-node-id^="decision:"]')
        assert decision_nodes.count() == 1
        assert page.locator('[data-diagram-node-id*="inline:"]').count() == 0
        assert "Decyzja o wniosku" in decision_nodes.text_content()
        edge_labels = page.locator(".workflow-diagram__edge-label").all_text_contents()
        assert {"Tak", "Nie", "Do poprawy"}.issubset(set(edge_labels))
        assert any("Scalono techniczne przejścia etapu" in warning for warning in console_warnings)
        assert page.locator("[data-workflow-diagram-mode]").inner_text() == "Układ automatyczny"
        inactive_section = page.locator("[data-workflow-inactive-decisions]")
        assert inactive_section.evaluate("(details) => details.open") is False
        review_stage = page.locator("[data-workflow-step]").filter(
            has=page.locator('[data-step-id][value="review"]')
        )
        assert review_stage.locator("[data-workflow-decision-card]").count() == 1
        assert page.locator(
            "[data-workflow-decision-pool] > [data-workflow-decision-card]"
        ).count() == 5

        review_stage.locator("[data-add-workflow-decision]").click()
        assert review_stage.locator("[data-workflow-decision-card]").count() == 2
        assert page.locator("[data-diagram-node-id]").count() == 6
        added_card = review_stage.locator('[data-workflow-decision-card][data-decision-created="true"]')
        added_card.locator('[data-decision-field="yes_status"]').select_option("PROCESS_COMPLETED")
        page.wait_for_function(
            "document.querySelectorAll('[data-diagram-node-id]').length === 7"
        )
        added_workflow = json.loads(page.locator("[data-workflow-builder-json]").input_value())
        added_decision = next(
            decision
            for decision in added_workflow["decision_settings"]
            if decision["id"].startswith("custom_decision_")
        )
        assert added_decision["step_id"] == "review"
        assert added_decision["active"] is True
        assert added_decision["yes_status"] == "PROCESS_COMPLETED"

        first_label = page.locator("[data-step-admin-label]").first
        first_label.evaluate(
            """(field) => {
                field.value = "Złożenie zaktualizowane";
                field.dispatchEvent(new Event("input", {bubbles: true}));
            }"""
        )
        assert "Złożenie zaktualizowane" in page.locator(
            '[data-diagram-node-id="submission"]'
        ).text_content()

        node = page.locator('[data-diagram-node-id="submission"]')
        node.scroll_into_view_if_needed()
        box = node.bounding_box()
        assert box
        page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
        page.mouse.down()
        page.mouse.move(box["x"] + box["width"] / 2 + 80, box["y"] + box["height"] / 2 + 30)
        page.mouse.up()
        assert page.locator("[data-workflow-diagram-mode]").inner_text() == "Układ ręczny"
        saved_workflow = json.loads(page.locator("[data-workflow-builder-json]").input_value())
        assert saved_workflow["diagram_layout"]["nodes"]["submission"]["x"] > 420

        width_before_zoom = page.locator("[data-workflow-diagram]").evaluate(
            "(svg) => parseFloat(svg.style.width)"
        )
        page.locator("[data-workflow-diagram-zoom-in]").click()
        width_after_zoom = page.locator("[data-workflow-diagram]").evaluate(
            "(svg) => parseFloat(svg.style.width)"
        )
        assert width_after_zoom > width_before_zoom

        page.locator("[data-workflow-diagram-layout-reset]").click()
        reset_workflow = json.loads(page.locator("[data-workflow-builder-json]").input_value())
        assert "diagram_layout" not in reset_workflow
        assert page.locator("[data-workflow-diagram-mode]").inner_text() == "Układ automatyczny"
        assert script_errors == []
        browser.close()


def test_existing_instruction_config_is_loaded_into_workflow_instruction_editors(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(
        admin_app,
        slug="existing_workflow_instruction",
        definition_json=readable_workflow_definition(),
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        form.user_instruction_config = {
            "title": "Co dalej?",
            "description": "Instrukcja ogólna",
            "stages": [
                {
                    "key": "review",
                    "label": "Weryfikacja wniosku",
                    "status_codes": ["OFFICER_REVIEW"],
                    "description": "<p>Zażółć gęślą jaźń.</p>",
                    "next_action": "<strong>Poczekaj na kontakt.</strong>",
                }
            ],
        }
        db.commit()
    login(admin_client)

    rendered = admin_client.get(
        f"/admin/forms/{form_id}/edit?tab=instructions"
    ).get_data(as_text=True)

    assert "Zażółć gęślą jaźń." in rendered
    assert "Poczekaj na kontakt." in rendered
    assert "textarea.disabled = false" in rendered
    assert "textarea.readOnly = false" in rendered
    assert "isWorkflowInstructionEditorEvent" in rendered
    assert "if (!isWorkflowInstructionEditorEvent(event)) syncWorkflowBuilder();" in rendered
    assert "if (isWorkflowInstructionEditorEvent(event)) return;" in rendered


def test_workflow_instruction_html_is_sanitized_before_save_and_returned_as_safe_html(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="safe_instruction", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    workflow = readable_workflow_definition()["workflow"]
    workflow["steps"][0]["description"] = '<p onclick="bad()"><strong>Ważny etap</strong><script>bad()</script></p>'
    workflow["steps"][0]["next_action"] = '<a href="https://example.com" onclick="bad()">Czytaj dalej</a><img src=x>'

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Safe instruction",
            "slug": "safe_instruction",
            "title": "Safe instruction",
            "workflow_name": "Workflow",
            "workflow_initial_step": "submission",
            "workflow_builder_json": json.dumps(workflow, ensure_ascii=False),
            "instruction_title": "Instrukcja",
            "user_instruction": '<p><em>Opis</em><iframe src="x">zło</iframe></p>',
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        step = form.definition_json["workflow"]["steps"][0]
        assert step["description"] == "<p><strong>Ważny etap</strong></p>"
        assert step["next_action"] == '<a href="https://example.com" target="_blank" rel="noopener noreferrer">Czytaj dalej</a>'
        assert form.user_instruction == "<p><em>Opis</em></p>"
        db.add(FormSubmission(submission_id="safe-html-submission", form_slug="safe_instruction", form_name="Safe", process_status="application_submitted"))
        db.commit()
    payload = admin_client.get("/api/submissions/safe-html-submission/acceptance-status").get_json()
    assert payload["instruction"]["description"] == "<p><em>Opis</em></p>"
    assert "onclick" not in payload["instruction"]["current_stage_description"]
    assert "<script" not in payload["instruction"]["current_stage_description"]


def test_removed_workflow_instruction_is_shown_as_inactive(admin_app, admin_client):
    create_user(admin_app)
    definition = readable_workflow_definition()
    form_id = create_form(admin_app, slug="inactive_instruction", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        form.user_instruction_config = {
            "title": "Instrukcja",
            "stages": [
                {"key": "submission", "label": "Wniosek", "status_codes": ["application_submitted"]},
                {"key": "removed", "label": "Usunięta instrukcja", "status_codes": ["REMOVED_STATUS"], "description": "Archiwalna treść"},
            ],
        }
        db.commit()
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=instructions").get_data(as_text=True)

    assert "Instrukcje nieaktywne" in html
    assert "Usunięta instrukcja" in html
    assert "Ten etap nie występuje już w workflow." in html
    assert "data-workflow-instruction-list" in html


def test_workflow_builder_rejects_missing_initial_stage(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="invalid_visual_workflow", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "name": "Workflow Form",
            "slug": "invalid_visual_workflow",
            "title": "Workflow Form",
            "workflow_name": "Workflow",
            "workflow_initial_step": "",
            "workflow_builder_json": json.dumps(readable_workflow_definition()["workflow"]),
        },
    )

    assert response.status_code == 400
    assert "Wybierz status początkowy workflow." in response.get_data(as_text=True)


def test_workflow_builder_repairs_office_confirmation_path_before_save(admin_app, admin_client):
    create_user(admin_app)
    definition = readable_workflow_definition()
    workflow = {
        "name": "Umowy szkoleniowe",
        "initial_step": "submission",
        "requires_contract": True,
        "requires_agreement_confirmation": True,
        "steps": [
            {"id": "submission", "status": "FORM_SUBMITTED", "next": "training_agreements_signature"},
            {
                "id": "training_agreements_signature",
                "admin_label": "Umowa oczekuje na podpis beneficjenta",
                "status": "AGREEMENT_WAITING_FOR_BENEFICIARY_SIGNATURE",
                "next": "completed",
            },
            {
                "id": "stage_10",
                "admin_label": "Oczekuje na podpis urzędu",
                "status": "AGREEMENT_WAITING_FOR_OFFICE_SIGNATURE",
                "next": "stage_11",
            },
            {
                "id": "stage_11",
                "admin_label": "Umowa podpisana przez urząd",
                "status": "AGREEMENT_SIGNED_BY_OFFICE",
                "next": "completed",
            },
            {"id": "completed", "status": "PROCESS_COMPLETED", "final": True},
        ],
    }
    definition["workflow"] = workflow
    form_id = create_form(admin_app, slug="repair_confirmation_path", definition_json=definition)
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=workflow").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "workflow",
            "name": "Umowy szkoleniowe",
            "slug": "repair_confirmation_path",
            "title": "Umowy szkoleniowe",
            "workflow_name": "Umowy szkoleniowe",
            "workflow_initial_step": "submission",
            "workflow_builder_json": json.dumps(workflow, ensure_ascii=False),
            "requires_contract": "on",
            "requires_agreement_confirmation": "on",
            "contract_template_html": "<p>Umowa</p>",
            "is_active": "on",
            "is_public": "on",
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        saved = db.get(Form, form_id).definition_json["workflow"]
        by_id = {step["id"]: step for step in saved["steps"]}
        assert by_id["training_agreements_signature"]["next"] == "stage_10"
        assert by_id["stage_10"]["next"] == "stage_11"
        assert by_id["stage_11"]["next"] == "completed"


def test_workflow_builder_action_respects_application_prefix(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="prefixed_workflow", definition_json=readable_workflow_definition())
    login(admin_client)

    response = admin_client.get(
        f"/admin/forms/{form_id}/edit",
        environ_overrides={"SCRIPT_NAME": "/aplikacja"},
    )
    html = response.get_data(as_text=True)

    assert 'href="/aplikacja/admin/forms"' in html
    assert "Zapisz workflow" in html


def test_form_edit_renders_tabbed_configuration(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="tabbed_form", definition_json=readable_workflow_definition())
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)

    for label in (
        "Podstawowe", "Pola formularza", "Szkolenia", "Deklaracja", "Umowa", "Workflow",
        "Instrukcje", "E-mail", "Regulaminy i dokumenty", "Logo i wygląd", "Uprawnienia",
        "Ustawienia zaawansowane",
    ):
        assert label in html
    assert 'data-form-tab-target="basic"' in html
    assert 'data-form-tab-panel="basic"' in html
    assert 'data-form-tab-select' in html
    assert 'data-active-tab-input' in html


def test_agreement_docx_preview_and_example_pdf_are_stateless_and_form_scoped(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    definition = {
        "title": "Umowy",
        "fields": [],
        "workflow": {
            "requires_contract": True,
            "contract_template_source": "docx",
            "contract_docx_template": {
                "html": '<main class="document document--agreement"><h1>{{ agreement_number }}</h1><p>{{ imiona }} {{ nazwisko }} — {{ training_name }}</p></main>',
                "valid": True,
                "unknown_variables": [],
            },
        },
    }
    form_id = create_form(admin_app, slug="preview_form", definition_json=definition)
    other_form_id = create_form(admin_app, slug="other_form", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        preview_submission = FormSubmission(submission_id="preview-submission", form_slug="preview_form", form_name="Umowy", imiona="Anna", nazwisko="Nowak", selected_trainings=json.dumps([{"id": "a", "name": "Excel", "price": 900}]), data_json={})
        db.add(preview_submission)
        db.flush()
        db.add(SubmissionTraining(submission_id=preview_submission.id, training_id="a", training_name_snapshot="Excel", training_price_snapshot="900", status="selected", is_locked=False, agreement_id=""))
        db.add(FormSubmission(submission_id="foreign-submission", form_slug="other_form", form_name="Inny", imiona="Ewa", nazwisko="Obca"))
        db.commit()
    login(admin_client)

    preview_without_training = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=submission&submission_id=preview-submission")
    preview = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=submission&submission_id=preview-submission&training_id=a")
    foreign = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=submission&submission_id=foreign-submission")
    monkeypatch.setattr(admin_app.extensions["services"].document_service.pdf_render_service, "render_document_pdf_bytes", lambda **kwargs: b"%PDF-1.4\npreview")
    pdf = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.pdf?preview_mode=submission&submission_id=preview-submission&training_id=a")

    assert preview_without_training.status_code == 200
    assert preview.status_code == 200
    assert preview.get_json()["ok"] is True
    assert "Anna Nowak" in preview.get_json()["html"]
    assert "Excel" in preview.get_json()["html"]
    assert "@page" in preview.get_json()["html"]
    assert foreign.status_code == 422
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF")
    assert "przykladowa_umowa_preview_form_a.pdf" in pdf.headers["Content-Disposition"]
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(SubmissionFile).count() == 0
        assert db.query(SubmissionWorkflowEvent).count() == 0
        training = db.query(SubmissionTraining).filter_by(submission_id=preview_submission.id, training_id="a").one()
        assert training.status == "selected"
        assert training.is_locked is False
        assert training.agreement_id == ""


def test_agreement_preview_example_needs_no_submission_or_training_and_has_full_context(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    docx_buffer = io.BytesIO()
    docx_document = Document()
    docx_document.add_paragraph("{{ agreement_number }}|{{ training_name }}|{{ training_price_formatted }}|{{ imiona }} {{ nazwisko }}|{{ stanowisko }}")
    docx_document.save(docx_buffer)
    storage_path = "output/example_only/templates/agreement/agreement-template.docx"
    definition = {
        "title": "Example preview",
        "fields": [{"name": "stanowisko", "label": "Stanowisko", "type": "text"}],
        "workflow": {
            "requires_contract": True,
            "contract_template_source": "docx",
            "contract_docx_template": {
                "storage_path": storage_path,
                "html": "<p>Nieaktualny cache</p>",
                "variables": ["agreement_number", "training_name", "training_price_formatted", "imiona", "nazwisko", "stanowisko"],
                "unknown_variables": [],
                "valid": True,
            },
        },
    }
    form_id = create_form(admin_app, slug="example_only", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    monkeypatch.setattr(admin_app.extensions["services"].storage, "read_bytes", lambda path: docx_buffer.getvalue(), raising=False)
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=example")
    monkeypatch.setattr(admin_app.extensions["services"].document_service.pdf_render_service, "render_document_pdf_bytes", lambda **kwargs: b"%PDF-1.4\nexample")
    pdf = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.pdf?preview_mode=example")

    assert response.status_code == 200
    html = response.get_json()["html"]
    assert "UM/PRZYKLAD/2026" in html
    assert "Przykładowe szkolenie" in html
    assert "1 500,00 zł" in html
    assert "Jan Kowalski" in html
    assert "Specjalista ds. projektów" in html
    assert pdf.status_code == 200
    assert pdf.mimetype == "application/pdf"
    assert pdf.data.startswith(b"%PDF")
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(FormSubmission).filter(FormSubmission.form_slug == "example_only").count() == 0
        assert db.query(SubmissionFile).count() == 0
        assert db.query(SubmissionWorkflowEvent).count() == 0
        assert db.query(SubmissionTraining).count() == 0


def test_agreement_example_preview_and_pdf_support_complex_jinja_without_persistence(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    complex_template = """
        <main>
            <p>{{ agreement_number }}|{{ participant_name }}|{{ pesel }}|{{ submission.get("ulica", "") }}</p>
            <p>{{ all_selected_trainings_total }}|{{ generated_date }}|{{ submission_id }}|{{ stanowisko }}|{{ submission.get("stanowisko", "") }}</p>
            <ol>{% for training in selected_trainings %}<li>{{ loop.index }}. {{ training.get("name", "") }}</li>{% endfor %}</ol>
        </main>
    """
    definition = {
        "title": "Złożony example context",
        "fields": [{"name": "stanowisko", "label": "Stanowisko", "type": "text"}],
        "workflow": {
            "requires_contract": True,
            "contract_template_source": "html",
            "contract_template_html": complex_template,
        },
    }
    form_id = create_form(admin_app, slug="complex_example", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    pdf_calls = []
    monkeypatch.setattr(
        admin_app.extensions["services"].document_service.pdf_render_service,
        "render_document_pdf_bytes",
        lambda **kwargs: pdf_calls.append(kwargs) or b"%PDF-1.4\ncomplex-example",
    )
    login(admin_client)

    preview = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=example")
    pdf = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.pdf?preview_mode=example")

    assert preview.status_code == 200
    rendered = preview.get_json()["html"]
    assert "UM/PRZYKLAD/2026|Jan Kowalski|90010112345|Przykładowa" in rendered
    assert "Specjalista ds. projektów" in rendered
    assert "1. Przykładowe szkolenie" in rendered
    assert pdf.status_code == 200
    assert pdf.mimetype == "application/pdf"
    assert pdf.headers["Content-Disposition"].startswith("attachment;")
    assert "przykladowa_umowa_complex_example.pdf" in pdf.headers["Content-Disposition"]
    assert pdf_calls[0]["context"]["submission"]["ulica"] == "Przykładowa"
    assert pdf_calls[0]["context"]["selected_trainings"][0]["name"] == "Przykładowe szkolenie"
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(FormSubmission).count() == 0
        assert db.query(SubmissionFile).count() == 0
        assert db.query(SubmissionWorkflowEvent).count() == 0
        assert db.query(SubmissionTraining).count() == 0
        assert db.query(EmailLog).count() == 0


def test_agreement_template_test_ui_exposes_preview_controls(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="preview_ui", definition_json={"title": "Umowa", "fields": [], "workflow": {"contract_template_source": "docx", "contract_docx_template": {"html": "<p>{{ agreement_number }} {{ imiona }}</p>", "variables": ["agreement_number", "imiona"], "unknown_variables": ["imiona"], "valid": False}}})
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)

    assert "Test szablonu" in html
    assert "Podgląd generowanej umowy" in html
    assert "Pobierz przykładową umowę PDF" in html
    assert 'id="agreement-preview-dialog"' in html
    assert 'aria-label="Zamknij podgląd"' in html
    assert 'class="document-preview__page agreement-preview-page"' in html
    assert "data-agreement-preview-open disabled" not in html
    assert 'data-agreement-preview-pdf href="/admin/forms/' in html
    assert "Aktualny plik: None" not in html
    assert "Pobierz orygina" not in html


def test_agreement_template_test_selects_training_and_opens_preview_in_browser(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    definition = {"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "docx", "contract_docx_template": {"html": "<p>{{ training_name }}</p>", "variables": ["training_name"], "valid": True}}}
    form_id = create_form(admin_app, slug="preview_browser", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(FormSubmission(submission_id="browser-many", form_slug="preview_browser", form_name="Umowa", selected_trainings=json.dumps([{"id": "excel", "name": "Excel"}, {"id": "kadry", "name": "Kadry"}])))
        db.commit()
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page()
        page.route("https://preview.test/**", lambda route: route.fulfill(status=200, content_type="application/json", body=json.dumps({"ok": True, "html": "<main class='document'>Kadry</main>"})))
        page.set_content(html, wait_until="domcontentloaded")
        page.add_style_tag(path=str(Path(__file__).resolve().parents[1] / "static" / "css" / "admin.css"))
        page.locator("[data-agreement-template-test]").evaluate("element => { element.dataset.previewUrl = 'https://preview.test/preview'; element.dataset.pdfUrl = 'https://preview.test/example.pdf'; }")
        assert page.locator("[data-agreement-preview-training-field]").is_hidden()
        assert page.locator("[data-agreement-preview-open]").is_enabled()
        page.locator("[data-agreement-preview-data]").select_option("browser-many")
        assert page.locator("[data-agreement-preview-training-field]").is_visible()
        assert page.locator("[data-agreement-preview-open]").is_disabled()
        page.locator("[data-agreement-preview-training]").select_option("kadry")
        assert page.locator("[data-agreement-preview-open]").is_enabled()
        assert "training_id=kadry" in page.locator("[data-agreement-preview-pdf]").get_attribute("href")
        page.locator("[data-agreement-preview-open]").click()
        assert page.locator("#agreement-preview-dialog").evaluate("dialog => dialog.open") is True
        page.wait_for_function("document.querySelector('[data-agreement-preview-frame]').srcdoc.includes('Kadry')")
        assert page.locator("[data-agreement-preview-page]").is_visible()
        dialog_box = page.locator("#agreement-preview-dialog").bounding_box()
        footer_box = page.locator(".agreement-preview-dialog__footer").bounding_box()
        assert dialog_box and dialog_box["y"] >= 0 and dialog_box["y"] + dialog_box["height"] <= page.viewport_size["height"]
        assert footer_box and footer_box["y"] + footer_box["height"] <= page.viewport_size["height"]
        page.get_by_label("Zamknij podgląd").click()
        assert page.locator("#agreement-preview-dialog").evaluate("dialog => dialog.open") is False
        page.locator("[data-agreement-preview-open]").click()
        page.keyboard.press("Escape")
        assert page.locator("#agreement-preview-dialog").evaluate("dialog => dialog.open") is False
        page.locator("[data-agreement-preview-open]").click()
        page.mouse.click(2, 2)
        assert page.locator("#agreement-preview-dialog").evaluate("dialog => dialog.open") is False
        page.locator("[data-agreement-preview-open]").click()
        page.unroute("https://preview.test/**")
        page.route("https://preview.test/**", lambda route: route.fulfill(status=422, content_type="application/json", body=json.dumps({"ok": False, "error": "Brak danych wymaganych przez szablon.", "reason": "missing_context_variables", "missing_variables": ["participant_address"]})))
        page.locator("[data-agreement-preview-refresh]").click()
        page.wait_for_function("document.querySelector('[data-agreement-preview-state]').textContent.includes('{{ participant_address }}')")
        assert "Brakujące zmienne" in page.locator("[data-agreement-preview-state]").inner_text()
        assert page.locator("[data-agreement-preview-state]").is_visible()
        assert page.locator("[data-agreement-preview-page]").is_hidden()
        current_url = page.url
        page.locator("[data-agreement-preview-modal-pdf]").click()
        page.wait_for_function("document.querySelector('[data-agreement-preview-state]').textContent.includes('{{ participant_address }}')")
        assert page.url == current_url
        page.locator(".agreement-preview-dialog__footer [data-admin-modal-close]").click()
        assert page.locator("#agreement-preview-dialog").evaluate("dialog => dialog.open") is False
        browser.close()


def test_agreement_preview_requires_training_choice_when_submission_has_many_trainings(admin_app, admin_client):
    create_user(admin_app)
    definition = {"title": "Umowa", "fields": [], "workflow": {"contract_template_source": "docx", "contract_docx_template": {"html": "<p>{{ training_name }}</p>", "valid": True, "unknown_variables": []}}}
    form_id = create_form(admin_app, slug="many_trainings", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(FormSubmission(submission_id="many", form_slug="many_trainings", form_name="Umowa", selected_trainings=json.dumps([{"id": "excel", "name": "Excel", "price": 100}, {"id": "kadry", "name": "Kadry", "price": 200}])))
        db.commit()
    login(admin_client)

    missing = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=submission&submission_id=many")
    selected = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=submission&submission_id=many&training_id=kadry")

    assert missing.status_code == 422
    assert "Wybierz konkretne szkolenie" in missing.get_json()["error"]
    assert selected.status_code == 200
    assert "Kadry" in selected.get_json()["html"]
    assert "Excel" not in selected.get_json()["html"].split("</style>")[-1]


def test_agreement_preview_rejects_unknown_mode_and_submission_without_training(admin_app, admin_client, caplog):
    create_user(admin_app)
    definition = {"title": "Umowa", "fields": [], "workflow": {"contract_template_source": "html", "contract_template_html": "<p>{{ training_name }}</p>"}}
    form_id = create_form(admin_app, slug="preview_validation", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        db.add(FormSubmission(submission_id="without-training", form_slug="preview_validation", form_name="Umowa"))
        db.commit()
    login(admin_client)

    with caplog.at_level(logging.WARNING):
        unknown = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=OTHER")
        no_training = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=submission&submission_id=without-training")

    assert unknown.status_code == 422
    assert unknown.get_json()["error"] == "Nieobsługiwany tryb podglądu."
    assert unknown.get_json()["reason"] == "invalid_preview_mode"
    assert no_training.status_code == 422
    assert "nie posiada szkolenia" in no_training.get_json()["error"]
    assert no_training.get_json()["reason"] == "missing_training"
    assert "agreement_preview_422" in caplog.text
    assert "preview_mode=other" in caplog.text


def test_agreement_preview_uses_html_compatibility_and_reports_invalid_docx(admin_app, admin_client):
    create_user(admin_app)
    html_form = create_form(admin_app, slug="html_preview", definition_json={"title": "HTML", "fields": [], "workflow": {"contract_template_source": "html", "contract_template_html": "<p>HTML {{ agreement_number }}</p>"}})
    invalid_form = create_form(admin_app, slug="invalid_docx", definition_json={"title": "DOCX", "fields": [], "workflow": {"contract_template_source": "docx", "contract_docx_template": {"html": "<p>{{ abc }}</p>", "unknown_variables": ["abc"], "valid": False}}})
    missing_form = create_form(admin_app, slug="missing_docx", definition_json={"title": "Brak", "fields": [], "workflow": {"contract_template_source": "docx"}})
    login(admin_client)

    html_preview = admin_client.get(f"/admin/forms/{html_form}/documents/agreement/preview")
    invalid_preview = admin_client.get(f"/admin/forms/{invalid_form}/documents/agreement/preview")
    missing_page = admin_client.get(f"/admin/forms/{missing_form}/edit?tab=agreement").get_data(as_text=True)

    assert html_preview.status_code == 200
    assert "HTML UM/PRZYKLAD/2026" in html_preview.get_json()["html"]
    assert invalid_preview.status_code == 422
    assert "{{ abc }}" in invalid_preview.get_json()["error"]
    assert "Nie wgrano szablonu umowy Word" in missing_page
    assert "data-agreement-preview-open disabled" in missing_page


def test_agreement_preview_reports_exact_missing_context_variable(admin_app, admin_client, caplog):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="missing_context", definition_json={"title": "HTML", "fields": [], "workflow": {"contract_template_source": "html", "contract_template_html": "<p>{{ missing_context_value }}</p>"}})
    login(admin_client)

    with caplog.at_level(logging.WARNING):
        response = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=example")
        pdf = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.pdf?preview_mode=example")

    assert response.status_code == 422
    assert response.get_json() == {
        "ok": False,
        "error": "Szablon używa zmiennej, dla której nie znaleziono danych: {{ missing_context_value }}.",
        "reason": "missing_context_variables",
        "missing_variables": ["missing_context_value"],
    }
    assert pdf.status_code == 422
    assert pdf.is_json
    assert pdf.get_json() == response.get_json()
    assert "agreement_preview_422" in caplog.text
    assert "reason=missing_context_variables" in caplog.text
    assert "missing_context_value" in caplog.text


def test_agreement_preview_classifies_invalid_template_and_missing_template(admin_app, admin_client):
    create_user(admin_app)
    invalid_form = create_form(
        admin_app,
        slug="invalid_jinja_preview",
        definition_json={"title": "Invalid", "fields": [], "workflow": {"contract_template_source": "html", "contract_template_html": "{% for training in selected_trainings %}"}},
    )
    missing_form = create_form(
        admin_app,
        slug="missing_template_preview",
        definition_json={"title": "Missing", "fields": [], "workflow": {"contract_template_source": "html", "contract_template_html": ""}},
    )
    login(admin_client)

    invalid = admin_client.get(f"/admin/forms/{invalid_form}/documents/agreement/preview?preview_mode=example")
    missing = admin_client.get(f"/admin/forms/{missing_form}/documents/agreement/example.pdf?preview_mode=example")

    assert invalid.status_code == 422
    assert invalid.get_json()["reason"] == "invalid_template"
    assert invalid.get_json()["missing_variables"] == []
    assert missing.status_code == 422
    assert missing.is_json
    assert missing.get_json()["reason"] == "missing_template"
    assert missing.get_json()["missing_variables"] == []


def test_agreement_preview_reports_corrupted_stored_docx_without_500(admin_app, admin_client, monkeypatch, caplog):
    create_user(admin_app)
    path = "output/corrupt/templates/agreement/agreement-template.docx"
    definition = {"title": "DOCX", "fields": [], "workflow": {"contract_template_source": "docx", "contract_docx_template": {"storage_path": path, "html": "<p>stare HTML</p>", "variables": [], "valid": True}}}
    form_id = create_form(admin_app, slug="corrupt", definition_json=definition)
    storage = admin_app.extensions["services"].storage
    monkeypatch.setattr(storage, "read_bytes", lambda requested_path: b"not-a-docx", raising=False)
    login(admin_client)

    with caplog.at_level(logging.WARNING):
        response = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview?preview_mode=example")

    assert response.status_code == 422
    assert response.get_json() == {
        "ok": False,
        "error": "Nie udało się odczytać zapisanego szablonu DOCX.",
        "reason": "docx_parse_error",
        "missing_variables": [],
    }
    assert "agreement_preview_422" in caplog.text
    assert "agreement_preview_failed" in caplog.text
    assert f"form_id={form_id}" in caplog.text
    assert "preview_mode=example" in caplog.text
    assert "BadZipFile" in caplog.text


def test_agreement_example_docx_download_is_valid_word_file(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="docx_example")
    login(admin_client)

    response = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.docx")

    assert response.status_code == 200
    assert response.data.startswith(b"PK")
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "agreement_number" in document_xml
    assert "training_name" in document_xml


def test_agreement_preview_endpoints_require_form_management_permission(admin_app, admin_client):
    create_user(admin_app, email="limited@example.com", role="admin")
    form_id = create_form(admin_app, slug="protected_preview", definition_json={"title": "Umowa", "fields": [], "workflow": {"contract_template_source": "html", "contract_template_html": "<p>Umowa</p>"}})
    login(admin_client, email="limited@example.com")

    preview = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/preview")
    pdf = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.pdf")
    docx = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/example.docx")

    assert preview.status_code == 403
    assert pdf.status_code == 403
    assert docx.status_code == 403


def test_agreement_builder_ui_exposes_visual_toolbar_components_and_source_modes(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="builder_ui", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder"}})
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)
    builder = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder").get_data(as_text=True)

    assert "Otwórz kreator wizualny" in html
    assert 'data-agreement-builder-popup' in html
    assert 'class="admin-button document-source-builder-button"' in html
    assert 'data-agreement-builder-popup' in html
    assert f'/admin/forms/{form_id}/documents/agreement/builder' in html
    assert "Kreator wizualny" in html
    assert "Microsoft Word (.docx)" in html
    assert "HTML — tryb zaawansowany" in html
    assert "Umowa: Sample" in builder
    assert 'data-builder-add="agreement_section"' in builder
    assert 'data-builder-add="training_table"' in builder
    assert 'data-builder-add="participant_data"' in builder
    assert 'data-builder-add="signatures"' in builder
    assert "Szukaj zmiennej…" in builder
    assert "Dopasuj szerokość" in builder
    assert "Zamknij kreator" in builder
    assert 'data-builder-view-mode="split"' in builder
    assert 'data-builder-view-mode="editor"' in builder
    assert 'data-builder-view-mode="preview"' in builder
    assert 'data-builder-variables-toggle' in builder
    assert 'data-builder-preview-viewport' in builder
    assert 'data-builder-preview-canvas' in builder
    assert 'data-layout="balanced"' in builder
    assert '<nav class="admin-nav"' not in builder
    assert "agreement_builder.js" in html
    assert "agreement_builder.js" in builder


def test_agreement_builder_endpoint_requires_form_management_permission(admin_app, admin_client):
    create_user(admin_app, email="builder-limited@example.com", role="admin")
    form_id = create_form(admin_app, slug="builder_protected", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True}})
    login(admin_client, email="builder-limited@example.com")

    response = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder")

    assert response.status_code == 403


def test_agreement_builder_draft_does_not_activate_and_activation_is_explicit(admin_app, admin_client):
    create_user(admin_app)
    definition = {"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "html", "contract_template_html": "<p>Legacy HTML</p>"}}
    form_id = create_form(admin_app, slug="builder_save", definition_json=definition)
    login(admin_client)
    builder = {
        "version": 1,
        "blocks": [{
            "type": "paragraph",
            "format": {"alignment": "center"},
            "runs": [
                {"text": "Uczestnik "},
                {"text": "{{ participant_name }}", "bold": True},
            ],
        }],
    }

    draft = admin_client.post(
        f"/admin/forms/{form_id}/agreement-template/builder",
        data={"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(builder), "action": "draft"},
    )
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["contract_template_source"] == "html"
        assert workflow["contract_builder_status"] == "draft"

    invalid = admin_client.post(
        f"/admin/forms/{form_id}/agreement-template/builder",
        data={"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps({"version": 1, "blocks": [{"type": "paragraph", "content": "{{ unknown_builder_value }}"}]}), "action": "activate"},
    )
    activated = admin_client.post(
        f"/admin/forms/{form_id}/agreement-template/builder",
        data={"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(builder), "action": "activate"},
    )

    assert draft.status_code == 200
    assert invalid.status_code == 422
    assert invalid.get_json()["errors"][0]["variable"] == "unknown_builder_value"
    assert activated.status_code == 200
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["contract_template_source"] == "builder"
        assert workflow["contract_builder_status"] == "active"
        assert workflow["contract_template_updated_source"] == "builder"
        assert workflow["contract_template_updated_by"]
        assert workflow["contract_builder_document"]["blocks"][0]["runs"] == [
            {"text": "Uczestnik ", "bold": False, "italic": False, "underline": False},
            {"text": "{{ participant_name }}", "bold": True, "italic": False, "underline": False},
        ]
        assert workflow["contract_builder_active_document"]["blocks"][0]["runs"] == workflow["contract_builder_document"]["blocks"][0]["runs"]
    reopened = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder").get_data(as_text=True)
    assert '"alignment": "center"' in reopened
    assert '"bold": true' in reopened
    assert "{{ participant_name }}" in reopened


def test_agreement_builder_draft_does_not_replace_already_active_builder(admin_app, admin_client):
    create_user(admin_app)
    active = {"version": 1, "blocks": [{"type": "paragraph", "content": "Aktywna treść"}]}
    definition = {"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder", "contract_builder_document": active, "contract_builder_active_document": active, "contract_builder_status": "active"}}
    form_id = create_form(admin_app, slug="builder_active_draft", definition_json=definition)
    login(admin_client)
    draft = {"version": 1, "blocks": [{"type": "paragraph", "content": "Nowa wersja robocza"}]}

    response = admin_client.post(
        f"/admin/forms/{form_id}/agreement-template/builder",
        data={"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(draft), "action": "draft"},
    )

    assert response.status_code == 200
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["contract_builder_document"]["blocks"][0]["content"] == "Nowa wersja robocza"
        assert workflow["contract_builder_active_document"]["blocks"][0]["content"] == "Aktywna treść"


def test_agreement_builder_live_preview_and_pdf_are_stateless(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    definition = {"title": "Umowa", "slug": "builder_preview", "fields": [{"name": "stanowisko", "label": "Stanowisko", "type": "text"}], "workflow": {"requires_contract": True, "contract_template_source": "builder"}}
    form_id = create_form(admin_app, slug="builder_preview", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    login(admin_client)
    builder = {
        "version": 1,
        "blocks": [
            {"type": "paragraph", "content": "{{ participant_name }} — {{ participant_address_inline }} — {{ stanowisko }}"},
            {
                "type": "paragraph",
                "runs": [
                    {"text": "Beneficjent przestrzega "},
                    {"text": "Regulaminu projektu", "bold": True, "italic": True, "underline": True},
                    {"text": "."},
                ],
            },
            {"type": "training_table", "scope": "selected_trainings", "columns": ["index", "name", "price"], "show_total": True},
            {"type": "signatures", "left_label": "Beneficjent", "right_label": "Uczestnik"},
        ],
    }
    payload = {"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(builder), "preview_mode": "example"}

    preview = admin_client.post(f"/admin/forms/{form_id}/documents/agreement/preview", data=payload)
    pdf_calls = []
    monkeypatch.setattr(
        admin_app.extensions["services"].document_service.pdf_render_service,
        "render_document_pdf_bytes",
        lambda **kwargs: pdf_calls.append(kwargs) or b"%PDF-1.4\nbuilder",
    )
    pdf = admin_client.post(f"/admin/forms/{form_id}/documents/agreement/example.pdf", data=payload)

    assert preview.status_code == 200
    html = preview.get_json()["html"]
    assert "Jan Kowalski" in html
    assert "ul. Przykładowa 12/3" in html
    assert "Specjalista ds. projektów" in html
    assert "Przykładowe szkolenie" in html
    assert "document-signatures" in html
    marked_phrase = '<strong><em><span class="document-text-underline">Regulaminu projektu</span></em></strong>'
    assert marked_phrase in html
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF")
    assert marked_phrase in pdf_calls[0]["template_html"]
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        assert db.query(FormSubmission).filter(FormSubmission.form_slug == "builder_preview").count() == 0
        assert db.query(SubmissionFile).count() == 0
        assert db.query(SubmissionWorkflowEvent).count() == 0
        assert db.query(SubmissionTraining).count() == 0


def test_agreement_builder_browser_inserts_blocks_variables_and_debounces_preview(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    form_id = create_form(admin_app, slug="builder_browser", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder"}})
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder").get_data(as_text=True)
    script = (Path(__file__).resolve().parents[1] / "static" / "js" / "agreement_builder.js").read_text(encoding="utf-8")
    window_stub = """<script>
        window.closeCalls = 0;
        window.close = () => { window.closeCalls += 1; };
        window.openerMessages = [];
        Object.defineProperty(window, 'opener', {configurable: true, value: {postMessage: (...args) => window.openerMessages.push(args)}});
    </script>"""
    html = html.replace("<head>", '<head><base href="https://preview.test/">').replace("</body>", window_stub + f"<script>{script}</script></body>")
    requests = []
    preview_error = {"active": False}

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page()

        def route_request(route):
            requests.append(route.request)
            if route.request.url.endswith("/agreement-template/builder"):
                route.fulfill(status=200, content_type="application/json", body=json.dumps({"ok": True, "message": "Wersja robocza zapisana."}))
            elif route.request.url.endswith("/example.pdf"):
                route.fulfill(status=200, content_type="application/pdf", body=b"%PDF")
            elif preview_error["active"]:
                route.fulfill(status=422, content_type="application/json", body=json.dumps({"ok": False, "error": "Brak danych wymaganych przez szablon.", "reason": "missing_context_variables", "missing_variables": ["participant_address"]}))
            else:
                route.fulfill(status=200, content_type="application/json", body=json.dumps({"ok": True, "html": "<main>Jan Kowalski</main>"}))

        page.route("https://preview.test/**", route_request)
        page.goto("https://preview.test/workspace")
        page.set_content(html, wait_until="domcontentloaded")
        builder = page.locator("[data-agreement-builder]")
        assert builder.is_visible()
        initial_count = page.locator("[data-builder-block]").count()
        page.locator('[data-builder-add="paragraph"]').first.click()
        assert page.locator("[data-builder-block]").count() == initial_count + 1
        last_editable = page.locator("[data-builder-block]").last.locator("[contenteditable=true]")
        last_editable.fill("AB")
        last_editable.press("ArrowLeft")
        page.locator('[data-builder-insert-variable="email"]').first.click()
        assert last_editable.inner_text() == "A{{ email }}B"
        page.evaluate("""() => {
            const editable = [...document.querySelectorAll('[data-builder-block] [contenteditable=true]')].at(-1);
            const range = document.createRange();
            range.selectNodeContents(editable);
            const selection = window.getSelection();
            selection.removeAllRanges();
            selection.addRange(range);
        }""")
        assert page.evaluate("""() => {
            const button = document.querySelector('[data-builder-format="bold"]');
            const event = new MouseEvent('mousedown', {bubbles: true, cancelable: true});
            button.dispatchEvent(event);
            return event.defaultPrevented;
        }""") is True

        bold = page.locator('[data-builder-format="bold"]')
        italic = page.locator('[data-builder-format="italic"]')
        underline = page.locator('[data-builder-format="underline"]')
        center = page.locator('[data-builder-align="center"]')
        justify = page.locator('[data-builder-align="justify"]')
        bold.click()
        italic.click()
        underline.click()
        center.click()
        assert center.get_attribute("aria-pressed") == "true"
        justify.click()
        assert justify.get_attribute("aria-pressed") == "true"
        assert center.get_attribute("aria-pressed") == "false"
        assert bold.get_attribute("aria-pressed") == "true"
        assert italic.get_attribute("aria-pressed") == "true"
        assert underline.get_attribute("aria-pressed") == "true"
        assert page.locator("[data-agreement-builder-status]").inner_text() == "Niezapisane zmiany"
        page.locator('[data-builder-action="undo"]').click()
        assert center.get_attribute("aria-pressed") == "true"
        assert justify.get_attribute("aria-pressed") == "false"
        page.locator('[data-builder-action="redo"]').click()
        assert justify.get_attribute("aria-pressed") == "true"
        assert center.get_attribute("aria-pressed") == "false"

        page.locator("[data-builder-close]").click()
        assert page.locator("[data-builder-close-dialog]").evaluate("dialog => dialog.open") is True
        page.locator("[data-builder-close-cancel]").click()
        assert page.locator("[data-builder-close-dialog]").evaluate("dialog => dialog.open") is False
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('Jan Kowalski')")
        preview_requests = [request for request in requests if request.url.endswith("/documents/agreement/preview")]
        assert preview_requests
        assert len(preview_requests) < 5
        assert "builder_json" in (preview_requests[-1].post_data or "")
        preview_error["active"] = True
        request_count_before_error = len(preview_requests)
        page.locator("[data-builder-preview-now]").click()
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview-state]').textContent.includes('{{ participant_address }}')")
        assert "Brakujące zmienne" in page.locator("[data-agreement-builder-preview-state]").inner_text()
        requests_after_error = [request for request in requests if request.url.endswith("/documents/agreement/preview")]
        assert len(requests_after_error) == request_count_before_error + 1
        page.evaluate("document.querySelector('[data-agreement-builder]').dispatchEvent(new CustomEvent('agreement-builder-visible'))")
        page.wait_for_timeout(700)
        requests_after_blocked_retry = [request for request in requests if request.url.endswith("/documents/agreement/preview")]
        assert len(requests_after_blocked_retry) == len(requests_after_error)
        preview_error["active"] = False
        page.locator('[data-builder-save="draft"]').first.click()
        page.wait_for_function("document.querySelector('[data-agreement-builder-status]').textContent.includes('zapisany')")
        saved_format = page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).format")
        saved_runs = page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).runs")
        assert saved_format == {"alignment": "justify"}
        assert saved_runs == [{"text": "A{{ email }}B", "bold": True, "italic": True, "underline": True}]
        messages = page.evaluate("window.openerMessages")
        assert messages[-1][0]["type"] == "agreement-builder-saved"
        assert messages[-1][0]["formId"] == form_id
        assert messages[-1][0]["documentType"] == "agreement"
        assert messages[-1][0]["updatedAt"]
        assert messages[-1][1] == "https://preview.test"
        page.locator("[data-builder-close]").click()
        assert page.evaluate("window.closeCalls") == 1
        browser.close()


def test_agreement_builder_initialization_prefers_backend_and_handles_local_drafts_and_preview_states(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    backend_document = {
        "version": 1,
        "document_type": "agreement",
        "blocks": [
            {"type": "paragraph", "content": "Wartość z backendu: 42"},
            {"type": "ordered_list", "items": [{"content": "Lista z backendu", "level": 0}]},
            *[
                {"type": "paragraph", "content": f"Blok backendu {index}"}
                for index in range(3, 43)
            ],
        ],
    }
    definition = {
        "title": "Umowa",
        "fields": [],
        "workflow": {
            "requires_contract": True,
            "contract_template_source": "builder",
            "contract_builder_document": backend_document,
            "contract_builder_updated_at": "2026-08-12T10:00:00+02:00",
        },
    }
    form_id = create_form(admin_app, slug="builder_initialization", definition_json=definition)
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder").get_data(as_text=True)
    assert f'data-storage-key="document-builder:agreement:form:{form_id}:v1"' in html
    assert f'data-legacy-storage-key="document-builder-agreement-draft-{form_id}"' in html
    assert 'data-backend-updated-at="2026-08-12T10:00:00+02:00"' in html
    assert 'data-backend-is-new="false"' in html
    script = (Path(__file__).resolve().parents[1] / "static" / "js" / "agreement_builder.js").read_text(encoding="utf-8")
    html = re.sub(
        r'<script src="[^"]*agreement_builder\.js"></script>',
        lambda _match: f"<script>{script}</script>",
        html,
    )
    storage_key = f"document-builder:agreement:form:{form_id}:v1"
    legacy_key = f"document-builder-agreement-draft-{form_id}"
    preview_response = {"status": 200, "payload": {"ok": True, "html": "<main>Podgląd działa</main>"}}
    preview_requests = []
    page_errors = []
    console_errors = []

    def route_request(route):
        if route.request.url.endswith("/documents/agreement/preview"):
            preview_requests.append(route.request)
            route.fulfill(
                status=preview_response["status"],
                content_type="application/json",
                body=json.dumps(preview_response["payload"]),
            )
            return
        content_type = "text/css" if route.request.url.endswith(".css") else "text/html"
        route.fulfill(status=200, content_type=content_type, body="<html></html>" if content_type == "text/html" else "")

    def open_builder(browser, *, key=storage_key, value=None):
        page = browser.new_page()
        page.on("pageerror", lambda error: page_errors.append(str(error)))
        page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
        page.route("https://preview.test/**", route_request)
        page.goto("https://preview.test/workspace")
        page.evaluate("localStorage.clear()")
        if value is not None:
            page.evaluate("([storageKey, storageValue]) => localStorage.setItem(storageKey, storageValue)", [key, value])
        page.set_content(html, wait_until="domcontentloaded")
        return page

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")

        empty_draft = json.dumps({
            "schemaVersion": 1,
            "formId": form_id,
            "documentType": "agreement",
            "savedAt": "2099-01-01T00:00:00Z",
            "document": {"version": 1, "document_type": "agreement", "blocks": []},
        })
        page = open_builder(browser, value=empty_draft)
        assert page.locator("[data-builder-block]").count() == 42
        assert "Wartość z backendu: 42" in page.locator("[data-agreement-builder-blocks]").inner_text()
        assert page.locator("[data-agreement-builder-status]").inner_text() == "Wczytano zapisany szablon."
        assert page.evaluate("key => localStorage.getItem(key)", storage_key) is None
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('Podgląd działa')")
        assert console_errors == []
        page.close()

        empty_paragraph_draft = json.dumps({
            "schemaVersion": 1,
            "formId": form_id,
            "documentType": "agreement",
            "savedAt": "2099-01-01T00:00:00Z",
            "document": {
                "version": 1,
                "document_type": "agreement",
                "blocks": [{"type": "paragraph", "content": "", "runs": []}],
            },
        })
        page = open_builder(browser, value=empty_paragraph_draft)
        assert "Wartość z backendu: 42" in page.locator("[data-agreement-builder-blocks]").inner_text()
        assert page.evaluate("key => localStorage.getItem(key)", storage_key) is None
        page.close()

        page = open_builder(browser, value="{uszkodzony-json")
        assert page.locator("[data-builder-block]").count() == 42
        assert page.evaluate("key => localStorage.getItem(key)", storage_key) is None
        page.close()

        local_document = {
            "version": 1,
            "document_type": "agreement",
            "blocks": [{"type": "paragraph", "content": "Nowszy lokalny szkic"}],
        }
        local_draft = json.dumps({
            "schemaVersion": 1,
            "formId": form_id,
            "documentType": "agreement",
            "savedAt": "2099-01-01T00:00:00Z",
            "document": local_document,
        })
        page = open_builder(browser, value=local_draft)
        assert "Wartość z backendu: 42" in page.locator("[data-agreement-builder-blocks]").inner_text()
        assert page.locator("[data-builder-draft-choice]").is_visible()
        assert page.locator("[data-builder-draft-restore]").inner_text() == "Przywróć szkic"
        assert page.locator("[data-builder-draft-discard]").inner_text() == "Użyj zapisanej wersji"
        assert page.locator("[data-agreement-builder-status]").inner_text() == "Znaleziono niezapisany szkic."
        page.locator("[data-builder-draft-restore]").click()
        assert "Nowszy lokalny szkic" in page.locator("[data-agreement-builder-blocks]").inner_text()
        assert page.locator("[data-agreement-builder-status]").inner_text() == "Przywrócono lokalny szkic"

        preview_response.update({
            "status": 422,
            "payload": {
                "ok": False,
                "error": "Brak danych wymaganych przez zapisany szablon.",
                "missing_variables": ["participant_address"],
            },
        })
        page.locator("[data-builder-preview-now]").click()
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview-state]').textContent.includes('{{ participant_address }}')")
        preview_error = page.locator("[data-agreement-builder-preview-state]").inner_text()
        assert "Nie udało się wygenerować podglądu." in preview_error
        assert "Brak danych wymaganych przez zapisany szablon." in preview_error

        preview_response.update({"status": 500, "payload": {"ok": False, "error": "Kontrolowany błąd serwera."}})
        requests_before_change = len(preview_requests)
        page.locator("[data-builder-block] [contenteditable=true]").first.fill("Nowszy lokalny szkic po zmianie")
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview-state]').textContent.includes('Kontrolowany błąd serwera.')")
        assert len(preview_requests) == requests_before_change + 1
        assert "Nie udało się wygenerować podglądu." in page.locator("[data-agreement-builder-preview-state]").inner_text()

        console_errors.clear()
        preview_response.update({"status": 200, "payload": {"ok": True, "html": "<main>Podgląd po błędzie działa</main>"}})
        page.locator("[data-builder-preview-now]").click()
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('Podgląd po błędzie działa')")
        assert page.locator("[data-agreement-builder-preview-state]").is_hidden()
        preview_request = preview_requests[-1]
        preview_body = preview_request.post_data or ""
        assert preview_request.method == "POST"
        assert preview_request.url.endswith("/documents/agreement/preview")
        assert preview_request.headers["content-type"].startswith("multipart/form-data;")
        assert "csrf_token" in preview_body
        assert "preview_mode" in preview_body
        assert "builder_json" in preview_body
        assert 'document_type\\\"%3A\\\"agreement' in preview_body or 'document_type\":\"agreement' in preview_body
        assert "Nowszy lokalny szkic po zmianie" in preview_body
        page.close()

        legacy_draft = json.dumps({"savedAt": "2099-01-01T00:00:00Z", "document": local_document})
        page = open_builder(browser, key=legacy_key, value=legacy_draft)
        assert page.locator("[data-builder-draft-choice]").is_visible()
        assert page.evaluate("key => localStorage.getItem(key)", legacy_key) is None
        assert page.evaluate("key => localStorage.getItem(key)", storage_key) is not None
        page.locator("[data-builder-draft-discard]").click()
        assert "Wartość z backendu: 42" in page.locator("[data-agreement-builder-blocks]").inner_text()
        assert page.evaluate("key => localStorage.getItem(key)", storage_key) is None
        page.close()

        browser.close()

    assert page_errors == []
    assert console_errors == []


def test_agreement_builder_browser_keeps_long_document_inside_independent_scroll_panels(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    form_id = create_form(admin_app, slug="builder_long_layout", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder"}})
    login(admin_client)
    project_root = Path(__file__).resolve().parents[1]
    html = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder").get_data(as_text=True)
    styles = "\n".join(
        (project_root / path).read_text(encoding="utf-8")
        for path in ("static/css/style.css", "static/css/admin.css")
    )
    script = (project_root / "static" / "js" / "agreement_builder.js").read_text(encoding="utf-8")
    html = html.replace("<head>", '<head><base href="https://preview.test/">').replace("</head>", f"<style>{styles}</style></head>").replace("</body>", f"<script>{script}</script></body>")
    preview_render_count = {"value": 0}
    long_paragraphs = "".join(f"<p>Akapit testowy {index}: długa treść podglądu umowy.</p>" for index in range(100))

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page(viewport={"width": 1500, "height": 900})

        def route_request(route):
            if route.request.url.endswith("/documents/agreement/preview"):
                preview_render_count["value"] += 1
                marker = preview_render_count["value"]
                preview_html = f"<!doctype html><html><head><style>html,body{{margin:0}}main{{box-sizing:border-box;padding:18mm;font:14px/1.5 sans-serif}}</style></head><body><main data-render=\"render-{marker}\">{long_paragraphs}</main></body></html>"
                route.fulfill(status=200, content_type="application/json", body=json.dumps({"ok": True, "html": preview_html}))
            else:
                route.fulfill(status=200, content_type="text/plain", body="")

        page.route("https://preview.test/**", route_request)
        page.goto("https://preview.test/workspace")
        page.set_content(html, wait_until="domcontentloaded")
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('render-1')")
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').contentDocument?.body?.textContent.includes('Akapit testowy 99')")
        assert page.locator("[data-agreement-builder-preview]").evaluate(
            "frame => frame.contentDocument.documentElement.classList.contains('agreement-preview-document')"
        ) is True

        page.evaluate("""() => {
            const addParagraph = document.querySelector('[data-builder-add="paragraph"]');
            for (let index = 0; index < 28; index += 1) addParagraph.click();
            const list = document.querySelector('.agreement-builder__variable-list');
            const sample = list.querySelector('.agreement-builder__variable');
            for (let index = 0; index < 18; index += 1) list.append(sample.cloneNode(true));
        }""")
        panel_metrics = page.evaluate("""() => {
            const root = document.querySelector('[data-agreement-builder]');
            const workspace = document.querySelector('.agreement-builder__workspace');
            const editor = document.querySelector('[data-builder-editor-pane]');
            const variables = document.querySelector('[data-agreement-builder-variable-panel]');
            return {
                rootClass: root.className,
                rootHeight: root.clientHeight,
                rootScrollHeight: root.scrollHeight,
                workspaceHeight: workspace.clientHeight,
                editorHeight: editor.clientHeight,
                editorScrollHeight: editor.scrollHeight,
                variablesHeight: variables.clientHeight,
                variablesScrollHeight: variables.scrollHeight,
            };
        }""")
        assert panel_metrics["editorScrollHeight"] > panel_metrics["editorHeight"], panel_metrics
        assert panel_metrics["variablesScrollHeight"] > panel_metrics["variablesHeight"], panel_metrics
        page.wait_for_timeout(900)

        bounds = page.evaluate("""() => {
            const root = document.querySelector('[data-agreement-builder]');
            const viewport = document.querySelector('[data-builder-preview-viewport]');
            const previewPage = document.querySelector('[data-agreement-builder-preview-page]');
            return {
                bodyFits: document.documentElement.scrollHeight <= window.innerHeight && document.body.scrollHeight <= window.innerHeight,
                rootFits: root.scrollHeight <= root.clientHeight,
                viewportScrolls: viewport.scrollHeight > viewport.clientHeight,
                viewportHeight: viewport.clientHeight,
                viewportContentHeight: viewport.scrollHeight,
                iframeHeight: document.querySelector('[data-agreement-builder-preview]').offsetHeight,
                a4Width: previewPage.offsetWidth,
                a4MinHeight: previewPage.offsetHeight,
            };
        }""")
        assert bounds["bodyFits"] is True
        assert bounds["rootFits"] is True
        assert bounds["viewportScrolls"] is True, bounds
        assert 790 <= bounds["a4Width"] <= 800
        assert bounds["a4MinHeight"] >= 1120

        def active_view_buttons():
            return page.locator('[data-builder-view-mode][aria-pressed="true"]').count()

        page.locator('[data-builder-view-mode="editor"]').click()
        assert active_view_buttons() == 1
        assert page.locator('[data-builder-view-mode="editor"]').get_attribute("aria-pressed") == "true"
        assert page.locator("[data-builder-preview-pane]").is_hidden()
        assert page.locator("[data-agreement-builder-variable-panel]").is_visible()

        page.locator("[data-builder-variables-toggle]").click()
        assert page.locator("[data-agreement-builder-variable-panel]").is_hidden()
        assert page.locator('[data-builder-view-mode="editor"]').get_attribute("aria-pressed") == "true"
        page.locator("[data-builder-variables-toggle]").click()

        page.locator('[data-builder-view-mode="preview"]').click()
        assert active_view_buttons() == 1
        assert page.locator("[data-builder-editor-pane]").is_hidden()
        assert page.locator("[data-builder-preview-pane]").is_visible()
        assert page.locator("[data-agreement-builder-variable-panel]").is_visible()

        page.locator('[data-builder-view-mode="split"]').click()
        page.locator('[data-builder-size-mode="editor-wide"]').click()
        assert page.locator("[data-agreement-builder]").get_attribute("data-layout") == "editor-wide"
        page.locator('[data-builder-size-mode="preview-wide"]').click()
        assert page.locator('[data-builder-size-mode="editor-wide"]').get_attribute("aria-pressed") == "false"
        assert page.locator('[data-builder-size-mode="preview-wide"]').get_attribute("aria-pressed") == "true"
        page.locator('[data-builder-size-mode="preview-wide"]').click()
        assert page.locator("[data-agreement-builder]").get_attribute("data-layout") == "balanced"

        page.locator('[data-builder-zoom="100"]').click()
        assert page.locator("[data-builder-zoom-value]").inner_text() == "100%"
        page.locator('[data-builder-zoom="out"]').click(click_count=8)
        assert page.locator("[data-builder-zoom-value]").inner_text() == "50%"
        page.locator('[data-builder-zoom="in"]').click()
        assert page.locator("[data-builder-zoom-value]").inner_text() == "60%"
        page.locator('[data-builder-zoom="fit"]').click()
        assert page.locator('[data-builder-zoom="fit"]').get_attribute("aria-pressed") == "true"
        fit_metrics = page.evaluate("""() => {
            const viewport = document.querySelector('[data-builder-preview-viewport]');
            const pageNode = document.querySelector('[data-agreement-builder-preview-page]');
            return {viewportWidth: viewport.clientWidth, pageWidth: pageNode.getBoundingClientRect().width};
        }""")
        assert fit_metrics["pageWidth"] <= fit_metrics["viewportWidth"] - 40
        page.locator('[data-builder-zoom="in"]').click()
        assert page.locator('[data-builder-zoom="fit"]').get_attribute("aria-pressed") == "false"

        scroll_before = page.evaluate("""() => {
            const editor = document.querySelector('[data-builder-editor-pane]');
            const variables = document.querySelector('[data-agreement-builder-variable-panel]');
            const preview = document.querySelector('[data-builder-preview-viewport]');
            editor.scrollTop = 420;
            variables.scrollTop = 310;
            preview.scrollTop = 480;
            return {editor: editor.scrollTop, variables: variables.scrollTop, preview: preview.scrollTop};
        }""")
        next_render = preview_render_count["value"] + 1
        page.evaluate("""() => {
            const editable = [...document.querySelectorAll('[data-builder-block] [contenteditable=true]')].at(-1);
            editable.innerHTML += ' zmiana';
            editable.dispatchEvent(new Event('input', {bubbles: true}));
        }""")
        page.wait_for_function(f"document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('render-{next_render}')")
        page.wait_for_timeout(100)
        scroll_after = page.evaluate("""() => ({
            editor: document.querySelector('[data-builder-editor-pane]').scrollTop,
            variables: document.querySelector('[data-agreement-builder-variable-panel]').scrollTop,
            preview: document.querySelector('[data-builder-preview-viewport]').scrollTop,
        })""")
        assert scroll_before["editor"] > 0 and scroll_after["editor"] == scroll_before["editor"]
        assert scroll_before["variables"] > 0 and scroll_after["variables"] == scroll_before["variables"]
        assert scroll_before["preview"] > 0 and abs(scroll_after["preview"] - scroll_before["preview"]) <= 2
        browser.close()


def test_agreement_builder_launcher_uses_reusable_popup_and_reports_blocker(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    form_id = create_form(admin_app, slug="builder_popup", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder"}})
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)
    script = (Path(__file__).resolve().parents[1] / "static" / "js" / "agreement_builder.js").read_text(encoding="utf-8")
    popup_stub = """<script>
        window.popupCalls = [];
        window.popupFocusCalls = 0;
        window.popupObject = {closed: false, location: {href: ''}, focus() { window.popupFocusCalls += 1; }};
        window.open = (...args) => { window.popupCalls.push(args); return window.popupObject; };
    </script>"""
    html_with_script = html.replace("<head>", '<head><base href="https://admin.test/">').replace("</body>", popup_stub + f"<script>{script}</script></body>")

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page()
        page.route("https://admin.test/**", lambda route: route.fulfill(status=200, content_type="text/plain", body=""))
        page.set_content(html_with_script, wait_until="domcontentloaded")
        button = page.locator("[data-agreement-builder-popup]")
        assert button.get_attribute("target") is None

        button.click()
        button.click()

        calls = page.evaluate("window.popupCalls")
        assert len(calls) == 1
        assert calls[0][0].endswith(f"/admin/forms/{form_id}/documents/agreement/builder")
        assert calls[0][1] == "agreementBuilder"
        assert "popup=yes" in calls[0][2]
        assert "resizable=yes" in calls[0][2]
        assert page.evaluate("window.popupFocusCalls") == 2

        page.evaluate("window.dispatchEvent(new MessageEvent('message', {origin: window.location.origin, data: {type: 'agreement-builder-saved', formId: %d, documentType: 'agreement', updatedAt: '2026-08-11T12:30:00+02:00'}}))" % form_id)
        assert "Ostatni zapis:" in page.locator("[data-agreement-builder-parent-status]").inner_text()

        blocked = browser.new_page()
        blocked_html = html.replace("<head>", '<head><base href="https://admin.test/">').replace("</body>", f"<script>window.open=()=>null;</script><script>{script}</script></body>")
        blocked.set_content(blocked_html, wait_until="domcontentloaded")
        blocked.locator("[data-agreement-builder-popup]").click()
        error = blocked.locator("[data-agreement-builder-popup-error]")
        assert error.is_visible()
        assert "zablokowała otwarcie kreatora" in error.inner_text()
        browser.close()


def test_agreement_docx_upload_also_prepares_editable_builder_draft(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="docx_builder_import", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder"}})
    document = Document()
    document.add_paragraph("§ 1.")
    document.add_paragraph("Przedmiot umowy")
    document.add_paragraph("Uczestnik {{ participant_name }}")
    buffer = io.BytesIO()
    document.save(buffer)
    login(admin_client)

    uploaded = admin_client.post(
        f"/admin/forms/{form_id}/agreement-template/docx",
        data={"csrf_token": admin_csrf(admin_client), "agreement_docx_template": (io.BytesIO(buffer.getvalue()), "umowa.docx")},
        content_type="multipart/form-data",
    )

    assert uploaded.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["contract_template_source"] == "docx"
        assert any(block["type"] == "agreement_section" for block in workflow["contract_builder_document"]["blocks"])
        assert any("participant_name" in str(block) for block in workflow["contract_builder_document"]["blocks"])

    imported = admin_client.post(
        f"/admin/forms/{form_id}/agreement-template/docx/import-builder",
        data={"csrf_token": admin_csrf(admin_client)},
    )

    assert imported.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["contract_template_source"] == "docx"
        assert workflow["contract_builder_status"] == "draft"


def test_form_edit_exposes_typed_qualification_field_definitions(admin_app, admin_client):
    create_user(admin_app)
    definition = {
        "title": "Qualification form",
        "fields": [
            {"name": "consent", "label": "Zgoda", "type": "checkbox"},
            {
                "name": "region",
                "label": "Region",
                "type": "select",
                "options": ["lubuskie", "wielkopolskie"],
            },
            {
                "name": "topics",
                "label": "Tematy",
                "type": "multi_select",
                "options": ["Excel", "Kadry"],
            },
            {"name": "age", "label": "Wiek", "type": "number"},
            {"name": "start_date", "label": "Data", "type": "date"},
        ],
    }
    form_id = create_form(admin_app, slug="qualification_ui", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    fields_json = html.split("data-qualification-fields>", 1)[1].split("</script>", 1)[0]
    fields = {field["name"]: field for field in json.loads(fields_json)}

    assert fields["consent"]["type"] == "checkbox"
    assert fields["region"]["options"] == ["lubuskie", "wielkopolskie"]
    assert fields["topics"]["options"] == ["Excel", "Kadry"]
    assert 'options = [{value: "TAK", label: "TAK"}, {value: "NIE", label: "NIE"}]' in html
    assert 'control.multiple = kind === "options-multiple"' in html
    assert 'control.type = kind;' in html
    assert "qualificationEmptyOperators" in html
    assert "Ten operator nie wymaga wartości oczekiwanej." in html


def test_form_edit_marks_requested_tab_as_active(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="active_tab_form", definition_json=readable_workflow_definition())
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=workflow").get_data(as_text=True)

    assert 'data-initial-tab="workflow"' in html
    assert 'name="active_tab" value="workflow"' in html
    assert 'data-form-tab-target="workflow" aria-controls="form-tab-workflow" aria-selected="true"' in html


def test_regular_admin_cannot_open_permissions_or_advanced_tabs(admin_app, admin_client):
    user_id = create_user(admin_app, role="admin")
    form_id = create_form(
        admin_app,
        slug="role_tabs_form",
        user_id=user_id,
        definition_json=readable_workflow_definition(),
    )
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=advanced").get_data(as_text=True)

    assert "Ustawienia zaawansowane" not in html
    assert "Pełna konfiguracja formularza JSON" not in html
    assert "Uprawnienia" not in html
    assert 'data-initial-tab="basic"' in html


def test_form_save_returns_to_active_tab(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="return_tab_form", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=appearance").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "appearance",
            "name": "Workflow Form",
            "slug": "return_tab_form",
            "title": "Workflow Form",
            "workflow_json": json.dumps(readable_workflow_definition()["workflow"]),
            "label_color": "#112233",
        },
    )

    assert response.status_code == 302
    assert response.location.endswith(f"/admin/forms/{form_id}/edit?tab=appearance")


def test_workflow_validation_opens_agreement_tab(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="error_tab_form", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "workflow",
            "name": "Workflow Form",
            "slug": "error_tab_form",
            "title": "Workflow Form",
            "workflow_name": "Workflow",
            "workflow_initial_step": "submission",
            "workflow_builder_json": json.dumps(readable_workflow_definition()["workflow"]),
            "requires_contract": "on",
            "contract_template_html": "<p>Umowa</p>",
        },
    )

    assert response.status_code == 400
    response_html = response.get_data(as_text=True)
    assert 'data-initial-tab="agreement"' in response_html
    assert "Proces wymaga umowy" in response_html


def test_invalid_full_json_opens_advanced_tab(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="json_error_tab", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=advanced").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "advanced",
            "name": "Workflow Form",
            "slug": "json_error_tab",
            "title": "Workflow Form",
            "use_form_definition_json": "on",
            "form_definition_json": "{niepoprawny",
        },
    )

    assert response.status_code == 400
    assert 'data-initial-tab="advanced"' in response.get_data(as_text=True)


def test_fields_editor_uses_same_tabs(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="fields_tabs_form")
    login(admin_client)

    html = admin_client.get(f"/admin/forms/{form_id}/fields").get_data(as_text=True)

    assert 'class="admin-form-tab is-active">Pola formularza</a>' in html
    assert "Regulaminy i dokumenty" in html
    assert "Wizualny układ formularza" in html
    assert "data-form-canvas" in html
    assert "new_sort_order" not in html


def test_form_upload_uses_tabs_and_polish_labels(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)

    html = admin_client.get("/admin/forms/upload").get_data(as_text=True)

    assert "Utwórz formularz" in html
    assert 'data-upload-tab="fields"' in html
    assert "Obsługiwane formaty" in html
    assert "Kolejność sortowania" in html
    assert "Powrót" in html


def test_form_upload_renders_four_separate_tab_panels(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)

    html = admin_client.get("/admin/forms/upload").get_data(as_text=True)

    assert html.count('role="tabpanel"') == 4
    assert html.count("data-upload-panel=") == 4
    fields_panel = html.split('id="upload-panel-fields"', 1)[1].split("</section>", 1)[0]
    basic_panel = html.split('id="upload-panel-basic"', 1)[1].split("</section>", 1)[0]
    emails_panel = html.split('id="upload-panel-emails"', 1)[1].split("</section>", 1)[0]
    appearance_panel = html.split('id="upload-panel-appearance"', 1)[1].split("</section>", 1)[0]

    assert 'name="form_file"' in fields_panel
    assert 'name="name"' not in fields_panel
    assert 'name="mail_mode"' not in fields_panel
    assert 'name="name"' in basic_panel
    assert 'name="form_file"' not in basic_panel
    assert 'name="mail_mode"' not in basic_panel
    assert 'name="mail_mode"' in emails_panel
    assert 'name="form_file"' not in emails_panel
    assert 'name="name"' not in emails_panel
    assert 'name="logo_id"' in appearance_panel
    assert 'name="form_file"' not in appearance_panel
    assert 'name="mail_mode"' not in appearance_panel
    assert 'id="upload-panel-basic"' in html and 'id="upload-panel-basic" class="admin-upload-panel" role=' in html
    assert 'panel.classList.toggle("is-active", isActive)' in html
    assert 'button.setAttribute("aria-selected", String(isActive))' in html


def test_form_upload_validation_returns_to_fields_tab_and_preserves_values(admin_app, admin_client):
    create_user(admin_app)
    login(admin_client)
    html = admin_client.get("/admin/forms/upload").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    response = admin_client.post(
        "/admin/forms/upload",
        data={
            "csrf_token": token,
            "active_tab": "emails",
            "name": "Zachowana nazwa",
            "form_smtp_mail_from": "nadawca@example.com",
        },
    )
    response_html = response.get_data(as_text=True)

    assert response.status_code == 400
    assert 'name="active_tab" value="fields"' in response_html
    assert 'id="upload-panel-fields" class="admin-upload-panel is-active"' in response_html
    assert 'id="upload-panel-emails" class="admin-upload-panel"' in response_html
    assert 'value="Zachowana nazwa"' in response_html
    assert 'value="nadawca@example.com"' in response_html


def test_super_admin_can_import_full_form_definition_json(admin_app, admin_client):
    create_user(admin_app)
    form_id = create_form(admin_app, slug="full_json_form", definition_json=readable_workflow_definition())
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=advanced").get_data(as_text=True)
    token = html.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
    imported = readable_workflow_definition()
    imported["fields"] = [{"name": "imported_field", "type": "text", "label": "Pole z importu"}]

    response = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "advanced",
            "name": "Workflow Form",
            "slug": "full_json_form",
            "title": "Workflow Form",
            "use_form_definition_json": "on",
            "form_definition_json": json.dumps(imported, ensure_ascii=False),
        },
    )

    assert response.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        assert form.definition_json["fields"][0]["name"] == "imported_field"
        assert db.query(FormField).filter_by(form_id=form_id, name="imported_field", active=True).one()


def test_form_tabs_have_mobile_css():
    css = (Path(__file__).parents[1] / "static" / "css" / "admin.css").read_text(encoding="utf-8")

    assert ".admin-form-tab-mobile" in css
    assert "@media (max-width: 720px)" in css
    assert ".admin-form-tabs" in css


def test_declaration_builder_ui_uses_shared_component_without_training_blocks(admin_app, admin_client):
    create_user(admin_app)
    definition = {
        "title": "Deklaracja",
        "fields": [{"name": "zgoda", "label": "Zgoda", "type": "checkbox"}],
        "workflow": {"requires_declaration": True, "declaration_template_source": "builder"},
    }
    form_id = create_form(admin_app, slug="declaration_builder_ui", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    login(admin_client)

    edit = admin_client.get(f"/admin/forms/{form_id}/edit?tab=declaration").get_data(as_text=True)
    builder = admin_client.get(f"/admin/forms/{form_id}/documents/declaration/builder")
    html = builder.get_data(as_text=True)

    assert builder.status_code == 200
    assert 'name="declaration_template_source" value="builder"' in edit
    assert "Microsoft Word (.docx)" in edit
    assert "HTML — tryb zaawansowany" in edit
    assert 'data-document-builder-popup' in edit
    assert f"/admin/forms/{form_id}/documents/declaration/builder" in edit
    assert 'data-document-type="declaration"' in html
    assert 'data-builder-add="form_field"' in html
    assert 'data-builder-add="participant_address"' in html
    assert 'data-builder-add="participant_contact"' in html
    assert 'data-builder-add="statement"' in html
    assert 'data-builder-add="participant_signature"' in html
    assert 'data-builder-add="training_table"' not in html
    assert 'data-builder-add="agreement_section"' not in html
    assert "training_name" not in html
    assert "agreement_builder.js" in html


@pytest.mark.parametrize("source", ["builder", "docx", "html"])
def test_document_source_selection_persists_after_form_save(admin_app, admin_client, source):
    create_user(admin_app)
    definition = {
        "title": "Dokumenty",
        "fields": [],
        "workflow": {
            "requires_contract": True,
            "requires_declaration": True,
            "contract_template_source": "builder",
            "declaration_template_source": "builder",
            "contract_template_html": "<p>Umowa</p>",
            "declaration_template_html": "<p>Deklaracja</p>",
            "contract_docx_template": {
                "original_filename": "umowa.docx",
                "storage_path": "templates/agreement/umowa.docx",
                "html": "<p>Umowa Word</p>",
                "valid": True,
            },
            "declaration_docx_template": {
                "original_filename": "deklaracja.docx",
                "storage_path": "templates/declaration/deklaracja.docx",
                "html": "<p>Deklaracja Word</p>",
                "valid": True,
            },
        },
    }
    form_id = create_form(admin_app, slug=f"document-source-{source}", name="Dokumenty", definition_json=definition)
    login(admin_client)
    edit = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)
    token = edit.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]

    saved = admin_client.post(
        f"/admin/forms/{form_id}/edit",
        data={
            "csrf_token": token,
            "active_tab": "agreement",
            "name": "Dokumenty",
            "slug": f"document-source-{source}",
            "title": "Dokumenty",
            "workflow_controls_present": "1",
            "workflow_name": "Dokumenty",
            "workflow_initial_step": "",
            "requires_contract": "on",
            "requires_declaration": "on",
            "contract_template_source": source,
            "declaration_template_source": source,
            "contract_template_html": "<p>Umowa</p>",
            "declaration_template_html": "<p>Deklaracja</p>",
        },
    )

    assert saved.status_code == 302
    reloaded = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)
    assert re.search(rf'name="contract_template_source" value="{source}" checked', reloaded)
    assert re.search(rf'name="declaration_template_source" value="{source}" checked', reloaded)


def test_document_source_cards_switch_panels_and_variable_catalog_without_reload(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    definition = {
        "title": "Dokumenty",
        "fields": [{"name": "zgoda_projektowa", "label": "Zgoda projektowa", "type": "checkbox"}],
        "workflow": {
            "requires_contract": True,
            "requires_declaration": True,
            "contract_template_source": "builder",
            "declaration_template_source": "builder",
        },
    }
    form_id = create_form(admin_app, slug="document_source_cards", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/edit?tab=agreement").get_data(as_text=True)
    assert re.search(r'data-document-variable-catalog="agreement"\s+hidden', html)
    assert re.search(r'data-document-variable-catalog="declaration"\s+hidden', html)
    assert "Aktualny plik: None" not in html
    assert html.count(">Wgraj DOCX</button>") == 2
    assert "Pobierz oryginał" not in html
    script = (Path(__file__).resolve().parents[1] / "static" / "js" / "agreement_builder.js").read_text(encoding="utf-8")
    html = html.replace("<head>", '<head><base href="https://admin.test/">').replace("</body>", f"<script>{script}</script></body>")

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page()
        page.route("https://admin.test/**", lambda route: route.fulfill(status=200, content_type="text/plain", body=""))
        page.set_content(html, wait_until="domcontentloaded")
        page.evaluate("Object.defineProperty(navigator, 'clipboard', {configurable: true, value: {writeText: (value) => { window.copiedVariable = value; return Promise.resolve(); }}})")

        for document_type in ("agreement", "declaration"):
            page.get_by_role("tab", name="Umowa" if document_type == "agreement" else "Deklaracja", exact=True).click()
            scope = page.locator(f'[data-document-source-scope][data-document-type="{document_type}"]')
            catalog = scope.locator(f'[data-document-variable-catalog="{document_type}"]')
            assert scope.locator('input[value="builder"]').is_checked()
            assert scope.locator('[data-document-source-option].is-active').count() == 1
            assert scope.locator('[data-source-panel="builder"]').evaluate("node => !node.hidden")
            assert catalog.evaluate("node => node.hidden")

            for source in ("docx", "html", "builder"):
                scope.locator(f'[data-document-source-option][data-source="{source}"]').click()
                assert scope.locator(f'input[value="{source}"]').is_checked()
                assert scope.locator('[data-document-source-option].is-active').count() == 1
                for panel_source in ("builder", "docx", "html"):
                    assert scope.locator(f'[data-source-panel="{panel_source}"]').evaluate("node => node.hidden") is (panel_source != source)
                assert catalog.evaluate("node => node.hidden") is (source == "builder")

        page.get_by_role("tab", name="Umowa", exact=True).click()
        agreement_scope = page.locator('[data-document-source-scope][data-document-type="agreement"]')
        agreement_scope.locator('[data-document-source-option][data-source="docx"]').click()
        catalog = agreement_scope.locator('[data-document-variable-catalog="agreement"]')
        cards = catalog.locator('[data-document-variable-card]')
        total_cards = cards.count()
        catalog.locator('[data-document-variable-search]').fill("szkolenie")
        visible_cards = catalog.locator('[data-document-variable-card]:not([hidden])')
        assert 0 < visible_cards.count() < total_cards
        visible_cards.first.click()
        assert page.evaluate("window.copiedVariable").startswith("{{")
        assert "Skopiowano" in catalog.locator('[data-document-variable-copy-status]').inner_text()
        browser.close()


def test_document_configuration_uses_shared_source_and_variable_partials():
    template = (Path(__file__).resolve().parents[1] / "templates" / "admin" / "forms" / "edit.html").read_text(encoding="utf-8")
    source_partial = (Path(__file__).resolve().parents[1] / "templates" / "admin" / "documents" / "_template_source_selector.html").read_text(encoding="utf-8")
    variable_partial = (Path(__file__).resolve().parents[1] / "templates" / "admin" / "documents" / "_variable_catalog.html").read_text(encoding="utf-8")
    css = (Path(__file__).resolve().parents[1] / "static" / "css" / "admin.css").read_text(encoding="utf-8")

    assert template.count('{% include "admin/documents/_template_source_selector.html" %}') == 2
    assert template.count('{% include "admin/documents/_variable_catalog.html" %}') == 2
    assert 'role="radiogroup"' in source_partial
    assert 'type="radio"' in source_partial
    assert 'class="document-source-input visually-hidden"' in source_partial
    assert "has_docx_template" in source_partial
    assert 'data-document-variable-search' in variable_partial
    assert "repeat(3, minmax(0, 1fr))" in css
    assert "repeat(2, minmax(0, 1fr))" in css


def test_declaration_builder_draft_activation_preview_and_pdf_are_stateless(admin_app, admin_client, monkeypatch):
    create_user(admin_app)
    definition = {
        "title": "Deklaracja",
        "fields": [{"name": "zgoda", "label": "Zgoda", "document_label": "Zgoda kwalifikacyjna", "type": "checkbox"}],
        "qualification_conditions": {
            "enabled": True,
            "conditions": [{"field_name": "zgoda", "field_label": "Zgoda", "operator": "equals", "expected_value": True, "is_active": True}],
        },
        "workflow": {"requires_declaration": True, "declaration_template_source": "html", "declaration_template_html": "<p>Stary HTML</p>"},
    }
    form_id = create_form(admin_app, slug="declaration_builder_save", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    login(admin_client)
    document = {
        "version": 1,
        "document_type": "declaration",
        "blocks": [
            {"type": "heading", "level": 1, "runs": [{"text": "Deklaracja {{ participant_name }}", "bold": True}]},
            {"type": "paragraph", "runs": [{"text": "Regulaminu projektu", "bold": True, "italic": True, "underline": True}]},
            {
                "type": "ordered_list",
                "items": [
                    {"content": "Punkt główny", "level": 0},
                    {"content": "Podpunkt", "level": 1},
                ],
                "list_styles": [
                    {"level": 0, "marker": "decimal-paren", "indent_mm": 0},
                    {"level": 1, "marker": "decimal-compound", "indent_mm": 9},
                    {"level": 2, "marker": "alpha-dot", "indent_mm": 18},
                    {"level": 3, "marker": "upper-roman-dot", "indent_mm": 27},
                ],
            },
            {"type": "form_field", "field": "zgoda", "label": "Zgoda", "display": "yes_no"},
            {"type": "criteria_table", "criteria": [{"field_key": "zgoda"}], "show_number": True, "show_header": True},
            {"type": "participant_signature", "label": "Podpis uczestnika"},
        ],
    }
    payload = {"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(document), "action": "draft"}

    draft = admin_client.post(f"/admin/forms/{form_id}/declaration-template/builder", data=payload)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["declaration_template_source"] == "html"
        assert workflow["declaration_builder_status"] == "draft"

    payload["action"] = "activate"
    activated = admin_client.post(f"/admin/forms/{form_id}/declaration-template/builder", data=payload)
    preview = admin_client.post(
        f"/admin/forms/{form_id}/documents/declaration/preview",
        data={"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(document), "preview_mode": "example"},
    )
    pdf_calls = []
    monkeypatch.setattr(
        admin_app.extensions["services"].document_service.pdf_render_service,
        "render_document_pdf_bytes",
        lambda **kwargs: pdf_calls.append(kwargs) or b"%PDF-1.4\ndeclaration",
    )
    pdf = admin_client.post(
        f"/admin/forms/{form_id}/documents/declaration/example.pdf",
        data={"csrf_token": admin_csrf(admin_client), "builder_json": json.dumps(document), "preview_mode": "example"},
    )

    assert draft.status_code == 200
    assert activated.status_code == 200
    assert preview.status_code == 200
    assert "Jan Kowalski" in preview.get_json()["html"]
    assert "Tak" in preview.get_json()["html"]
    assert "Zgoda kwalifikacyjna" in preview.get_json()["html"]
    assert ">1)</span>" in preview.get_json()["html"]
    assert ">1.1</span>" in preview.get_json()["html"]
    marked_phrase = '<strong><em><span class="document-text-underline">Regulaminu projektu</span></em></strong>'
    assert marked_phrase in preview.get_json()["html"]
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    assert marked_phrase in pdf_calls[0]["template_html"]
    assert "document-criteria-table" in pdf_calls[0]["template_html"]
    assert "--document-list-indent: 9mm" in pdf_calls[0]["template_html"]
    reopened = admin_client.get(f"/admin/forms/{form_id}/documents/declaration/builder").get_data(as_text=True)
    assert 'data-builder-add="criteria_table"' in reopened
    assert '"marker": "decimal-compound"' in reopened
    assert '"field_key": "zgoda"' in reopened
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["declaration_template_source"] == "builder"
        assert workflow["declaration_builder_status"] == "active"
        assert workflow["declaration_builder_active_document"]["document_type"] == "declaration"
        assert workflow["declaration_builder_active_document"]["blocks"][2]["items"][1]["level"] == 1
        assert workflow["declaration_builder_active_document"]["blocks"][4]["criteria"] == [{"field_key": "zgoda"}]
        assert db.query(FormSubmission).filter(FormSubmission.form_slug == "declaration_builder_save").count() == 0
        assert db.query(SubmissionFile).count() == 0
        assert db.query(SubmissionWorkflowEvent).count() == 0
        assert db.query(SubmissionTraining).count() == 0


def test_declaration_docx_upload_prepares_builder_draft_and_keeps_source_docx(admin_app, admin_client):
    create_user(admin_app)
    definition = {
        "title": "Deklaracja",
        "fields": [{"name": "zgoda", "label": "Zgoda", "type": "checkbox"}],
        "qualification_conditions": {
            "enabled": True,
            "conditions": [{"field_name": "zgoda", "operator": "equals", "expected_value": True, "is_active": True}],
        },
        "workflow": {"requires_declaration": True, "declaration_template_source": "builder"},
    }
    form_id = create_form(admin_app, slug="declaration_docx", definition_json=definition)
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        form = db.get(Form, form_id)
        sync_form_fields(db, form, definition)
        db.commit()
    document = Document()
    paragraph = document.add_paragraph("Uczestnik ")
    run = paragraph.add_run("{{ participant_name }}")
    run.bold = True
    document.add_paragraph("TAK: {{ criterion_zgoda_yes_checked }}")
    buffer = io.BytesIO()
    document.save(buffer)
    login(admin_client)

    uploaded = admin_client.post(
        f"/admin/forms/{form_id}/declaration-template/docx",
        data={"csrf_token": admin_csrf(admin_client), "declaration_docx_template": (io.BytesIO(buffer.getvalue()), "deklaracja.docx")},
        content_type="multipart/form-data",
    )
    imported = admin_client.post(
        f"/admin/forms/{form_id}/declaration-template/docx/import-builder",
        data={"csrf_token": admin_csrf(admin_client)},
    )

    assert uploaded.status_code == 302
    assert imported.status_code == 302
    with create_session_factory(admin_app.config["DATABASE_URL"])() as db:
        workflow = db.get(Form, form_id).definition_json["workflow"]
        assert workflow["declaration_template_source"] == "docx"
        assert workflow["declaration_builder_status"] == "draft"
        assert workflow["declaration_docx_template"]["document_type"] == "declaration"
        assert workflow["declaration_docx_template"]["unknown_variables"] == []
        runs = workflow["declaration_builder_document"]["blocks"][0]["runs"]
        assert any(run["bold"] and "participant_name" in run["text"] for run in runs)


def test_document_builder_browser_preserves_selection_and_toggles_inline_runs(admin_app, admin_client):
    playwright_api = pytest.importorskip("playwright.sync_api")
    create_user(admin_app)
    form_id = create_form(admin_app, slug="builder_inline_runs", definition_json={"title": "Umowa", "fields": [], "workflow": {"requires_contract": True, "contract_template_source": "builder"}})
    login(admin_client)
    html = admin_client.get(f"/admin/forms/{form_id}/documents/agreement/builder").get_data(as_text=True)
    script = (Path(__file__).resolve().parents[1] / "static" / "js" / "agreement_builder.js").read_text(encoding="utf-8")
    document_css = (Path(__file__).resolve().parents[1] / "static" / "css" / "document_template.css").read_text(encoding="utf-8")
    html = html.replace("<head>", '<head><base href="https://inline.test/">').replace("</body>", f"<script>{script}</script></body>")

    with playwright_api.sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=True)
        except playwright_api.Error as exception:
            pytest.skip(f"Brak przeglądarki Playwright: {exception}")
        page = browser.new_page()
        preview_requests = []

        def route_request(route):
            body = route.request.post_data or ""
            if route.request.url.endswith("/documents/agreement/preview"):
                preview_requests.append(route.request)
                formatted = '"text":"Regulaminu projektu","bold":true,"italic":true,"underline":true' in body
                preview_html = (
                    f'<style>{document_css}</style><main class="document"><strong><em><span class="document-text-underline">Regulaminu projektu</span></em></strong></main>'
                    if formatted else "<main>Podgląd</main>"
                )
                route.fulfill(status=200, content_type="application/json", body=json.dumps({"ok": True, "html": preview_html}))
            else:
                route.fulfill(status=200, content_type="application/json", body=json.dumps({"ok": True}))

        page.route("https://inline.test/**", route_request)
        page.goto("https://inline.test/workspace")
        page.set_content(html, wait_until="domcontentloaded")
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('Podgląd')")
        preview_count_before_formatting = len(preview_requests)
        page.locator('[data-builder-add="paragraph"]').first.click()
        editable = page.locator("[data-builder-block]").last.locator("[contenteditable=true]")
        text = "Beneficjent zobowiązuje się przestrzegać postanowień Regulaminu projektu."
        phrase = "Regulaminu projektu"
        editable.fill(text)
        page.evaluate("""() => {
            const editable = [...document.querySelectorAll('[data-builder-block] [contenteditable=true]')].at(-1);
            const node = editable.firstChild;
            const phrase = 'Regulaminu projektu';
            const start = node.nodeValue.indexOf(phrase);
            const range = document.createRange();
            range.setStart(node, start);
            range.setEnd(node, start + phrase.length);
            const selection = window.getSelection();
            selection.removeAllRanges();
            selection.addRange(range);
        }""")

        page.locator('[data-builder-format="bold"]').click()
        page.locator('[data-builder-format="italic"]').click()
        page.locator('[data-builder-format="underline"]').click()

        runs = page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).runs")
        phrase_start = text.index(phrase)
        assert runs == [
            {"text": text[:phrase_start], "bold": False, "italic": False, "underline": False},
            {"text": phrase, "bold": True, "italic": True, "underline": True},
            {"text": text[phrase_start + len(phrase):], "bold": False, "italic": False, "underline": False},
        ]
        for mark in ("bold", "italic", "underline"):
            button = page.locator(f'[data-builder-format="{mark}"]')
            assert button.get_attribute("aria-pressed") == "true"
            button.click()
            assert page.evaluate(f"JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).runs[1].{mark}") is False
            button.click()
            assert page.evaluate(f"JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).runs[1].{mark}") is True
        assert page.evaluate("window.getSelection().toString()") == phrase
        assert page.locator("[data-agreement-builder-status]").inner_text() == "Niezapisane zmiany"

        page.locator('[data-builder-action="undo"]').click()
        assert page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).runs[1].underline") is False
        page.locator('[data-builder-action="redo"]').click()
        assert page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).runs[1].underline") is True
        page.wait_for_function("document.querySelector('[data-agreement-builder-preview]').srcdoc.includes('document-text-underline')")
        page.wait_for_function("getComputedStyle(document.querySelector('[data-agreement-builder-preview]').contentDocument.querySelector('strong')).fontWeight === '700'")
        assert page.evaluate("getComputedStyle(document.querySelector('[data-agreement-builder-preview]').contentDocument.querySelector('strong')).fontWeight") == "700"
        assert len(preview_requests) > preview_count_before_formatting
        before_collapsed_click = page.locator("[data-agreement-builder-json]").input_value()
        page.evaluate("""() => {
            const editable = [...document.querySelectorAll('[data-builder-block] [contenteditable=true]')].at(-1);
            const node = editable.querySelector('[data-builder-run]')?.firstChild || editable.firstChild;
            const range = document.createRange();
            range.setStart(node, 1);
            range.collapse(true);
            const selection = window.getSelection();
            selection.removeAllRanges();
            selection.addRange(range);
        }""")
        page.locator('[data-builder-format="bold"]').click()
        assert page.locator("[data-agreement-builder-json]").input_value() == before_collapsed_click

        page.locator('[data-builder-add="ordered_list"]').first.click()
        list_block = page.locator("[data-builder-block]").last
        list_block.locator("[data-builder-add-item]").click()
        second_item = list_block.locator("[data-builder-list-item]").nth(1)
        second_item.focus()
        page.keyboard.press("Tab")
        assert page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).items[1].level") == 1
        page.keyboard.press("Shift+Tab")
        assert page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1).items[1].level") == 0
        page.locator('[data-builder-list-level="increase"]').click()
        list_block.locator('[data-builder-list-marker][data-level="1"]').select_option("alpha-dot")
        list_block.locator('[data-builder-list-indent][data-level="1"]').fill("12")
        list_model = page.evaluate("JSON.parse(document.querySelector('[data-agreement-builder-json]').value).blocks.at(-1)")
        assert list_model["items"][1]["level"] == 1
        assert list_model["list_styles"][1]["marker"] == "alpha-dot"
        assert list_model["list_styles"][1]["indent_mm"] == 12
        browser.close()
