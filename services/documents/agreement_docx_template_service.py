from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from services.documents.agreement_template_context_service import AgreementVariableCatalog
from services.documents.docx_template_parser import PARSER_VERSION, parse_docx_template


DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def parse_stored_agreement_docx(storage, metadata: Mapping[str, Any]):
    document_type = str(metadata.get("document_type") or "agreement").strip().casefold()
    path = str(metadata.get("storage_path") or "").strip()
    if path:
        try:
            return parse_docx_template(storage.read_bytes(path), document_type=document_type)
        except Exception as exc:
            raise ValueError("Nie udało się odczytać zapisanego szablonu DOCX.") from exc
    html = str(metadata.get("html") or "").strip()
    if html:
        from services.documents.docx_template_parser import ParsedDocxTemplate

        return ParsedDocxTemplate(
            html=html,
            variables=tuple(str(name) for name in metadata.get("variables") or []),
            warnings=tuple(str(item) for item in metadata.get("warnings") or []),
        )
    raise ValueError("Nie wgrano szablonu umowy Word.")


class AgreementDocxTemplateService:
    def __init__(self, storage):
        self.storage = storage

    def upload(
        self,
        *,
        form_slug: str,
        filename: str,
        content: bytes,
        fields: Iterable[Any],
        uploaded_by_user_id: int | None,
        document_type: str = "agreement",
    ) -> dict:
        document_type = str(document_type or "agreement").strip().casefold()
        if document_type not in {"agreement", "declaration"}:
            raise ValueError("Nieobsługiwany typ dokumentu.")
        label = "umowy" if document_type == "agreement" else "deklaracji"
        if Path(filename or "").suffix.casefold() != ".docx":
            raise ValueError(f"Szablon {label} musi być plikiem DOCX.")
        parsed = parse_docx_template(content, document_type=document_type)
        if document_type == "declaration":
            from services.documents.declaration_template_context_service import DeclarationVariableCatalog

            known = DeclarationVariableCatalog.context_names(fields)
        else:
            known = AgreementVariableCatalog.context_names(fields)
        unknown = sorted(set(parsed.variables) - known)
        storage_path = self.storage_path(form_slug, document_type=document_type)
        self._ensure_directory(form_slug, document_type=document_type)
        self.storage.write_bytes(storage_path, content, DOCX_MIME_TYPE)
        return {
            "storage_path": storage_path,
            "document_type": document_type,
            "original_filename": Path(filename).name,
            "mime_type": DOCX_MIME_TYPE,
            "size_bytes": len(content),
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "uploaded_by_user_id": uploaded_by_user_id,
            "parser_version": PARSER_VERSION,
            "variables": list(parsed.variables),
            "warnings": list(parsed.warnings),
            "unknown_variables": unknown,
            "valid": not unknown,
            "html": parsed.html,
            "builder_document": parsed.builder_document,
        }

    def download(self, metadata: Mapping[str, Any]) -> bytes:
        path = str(metadata.get("storage_path") or "").strip()
        if not path:
            raise FileNotFoundError("Brak zapisanego szablonu DOCX.")
        return self.storage.read_bytes(path)

    def parse_stored_template(self, metadata: Mapping[str, Any]):
        return parse_stored_agreement_docx(self.storage, metadata)

    def delete(self, metadata: Mapping[str, Any]) -> None:
        path = str(metadata.get("storage_path") or "").strip()
        if path:
            self.storage.delete(path, missing_ok=True)

    def sample_docx(self, fields: Iterable[Any], *, document_type: str = "agreement") -> bytes:
        document = Document()
        if document_type == "declaration":
            title = document.add_heading("DEKLARACJA UCZESTNICTWA", level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph("sporządzona w dniu {{ generated_date }}")
            document.add_heading("Dane uczestnika", level=2)
            document.add_paragraph("{{ imiona }} {{ nazwisko }}\nPESEL: {{ pesel }}")
            document.add_paragraph("Adres: {{ participant_address_inline }}")
            document.add_paragraph("E-mail: {{ email }}\nTelefon: {{ telefon }}")
            document.add_heading("Oświadczenie", level=2)
            paragraph = document.add_paragraph("Oświadczam, że dane podane w formularzu są prawdziwe.")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            document.add_paragraph("\n\n___________________\nCzytelny podpis uczestnika")
            buffer = BytesIO()
            document.save(buffer)
            return buffer.getvalue()
        title = document.add_heading("UMOWA UCZESTNICTWA W PROJEKCIE", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        number = document.add_paragraph("nr {{ agreement_number }}")
        number.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("zawarta w dniu {{ generated_date }} pomiędzy:")
        document.add_paragraph("[NAZWA INSTYTUCJI]")
        document.add_paragraph("a")
        document.add_paragraph("{{ imiona }} {{ nazwisko }}\nE-mail: {{ email }}\nTelefon: {{ telefon }}")
        document.add_heading("§ 1", level=2)
        document.add_heading("Przedmiot umowy", level=3)
        paragraph = document.add_paragraph("Przedmiotem umowy jest udział w szkoleniu {{ training_name }}. Cena szkolenia wynosi {{ training_price_formatted }}, a łączna wartość wsparcia {{ all_selected_trainings_total_formatted }}.")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Szkolenie"
        table.cell(0, 1).text = "Cena"
        table.cell(1, 0).text = "{{ training_name }}"
        table.cell(1, 1).text = "{{ training_price_formatted }}"
        document.add_heading("§ 2", level=2)
        document.add_heading("Postanowienia", level=3)
        document.add_paragraph("[Tutaj wpisz treść umowy.]")
        signatures = document.add_table(rows=2, cols=2)
        signatures.cell(0, 0).text = "Uczestnik"
        signatures.cell(0, 1).text = "Urząd"
        signatures.cell(1, 0).text = "___________________"
        signatures.cell(1, 1).text = "___________________"
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def storage_path(self, form_slug: str, *, document_type: str = "agreement") -> str:
        filename = "agreement-template.docx" if document_type == "agreement" else "declaration-template.docx"
        return f"{self.storage.output_dir}/{form_slug}/templates/{document_type}/{filename}"

    def _ensure_directory(self, form_slug: str, *, document_type: str = "agreement") -> None:
        self.storage.mkdir(f"{self.storage.output_dir}/{form_slug}")
        self.storage.mkdir(f"{self.storage.output_dir}/{form_slug}/templates")
        self.storage.mkdir(f"{self.storage.output_dir}/{form_slug}/templates/{document_type}")
