from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from services.documents.agreement_docx_template_service import AgreementDocxTemplateService
from services.documents.docx_template_parser import DocxTemplateParseError, parse_docx_template
from services.form_config_service import FormConfigService


def _docx_bytes(build):
    document = Document()
    build(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_parser_preserves_split_jinja_variable_heading_alignment_table_and_empty_paragraph():
    def build(document):
        heading = document.add_heading("Umowa ", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run("{{")
        heading.add_run(" agreement_")
        heading.add_run("number ")
        heading.add_run("}}")
        document.add_paragraph("")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Szkolenie"
        table.cell(0, 1).text = "{{ training_name }}"

    parsed = parse_docx_template(_docx_bytes(build))

    assert parsed.variables == ("agreement_number", "training_name")
    assert "{{ agreement_number }}" in parsed.html
    assert '<h1 class="document-paragraph document-align-center">' in parsed.html
    assert "document-paragraph--empty" in parsed.html
    assert '<table class="document-table document-table--docx">' in parsed.html


def test_parser_escapes_word_text_but_keeps_jinja_placeholder():
    parsed = parse_docx_template(_docx_bytes(lambda document: document.add_paragraph("A < B & {{ pesel }} {% unsafe %}")))
    assert "A &lt; B &amp; {{ pesel }}" in parsed.html
    assert "&#123;% unsafe %&#125;" in parsed.html
    assert parsed.warnings


def test_parser_renders_legal_paragraph_and_following_title_outside_lists():
    def build(document):
        document.add_paragraph("§ 2.", style="List Number")
        document.add_paragraph("Przedmiot umowy")

    parsed = parse_docx_template(_docx_bytes(build))

    assert '<div class="document-section-number">§ 2.</div>' in parsed.html
    assert '<div class="document-section-title">Przedmiot umowy</div>' in parsed.html
    assert "<li" not in parsed.html


def test_parser_closes_numbered_list_before_legal_paragraph():
    def build(document):
        document.add_paragraph("Punkt pierwszy", style="List Number")
        document.add_paragraph("Punkt drugi", style="List Number")
        document.add_paragraph("§ 3.", style="List Number")
        document.add_paragraph("Prawa i obowiązki")

    parsed = parse_docx_template(_docx_bytes(build))

    assert parsed.html.count("<li ") == 2
    assert parsed.html.count('<ol class="document-list">') == 1
    assert parsed.html.index("</ol>") < parsed.html.index("document-section-number")
    assert '<div class="document-section-title">Prawa i obowiązki</div>' in parsed.html


def test_parser_skips_empty_numbered_paragraph_and_groups_same_num_id():
    def build(document):
        document.add_paragraph("Pierwszy", style="List Number")
        document.add_paragraph("", style="List Number")
        document.add_paragraph("Drugi", style="List Number")

    parsed = parse_docx_template(_docx_bytes(build))

    assert parsed.html.count('<ol class="document-list">') == 1
    assert parsed.html.count("<li ") == 2
    assert "<li></li>" not in parsed.html
    assert "<li>&nbsp;</li>" not in parsed.html


def test_parser_rejects_empty_and_invalid_documents():
    with pytest.raises(DocxTemplateParseError):
        parse_docx_template(b"")
    with pytest.raises(DocxTemplateParseError):
        parse_docx_template(b"not a docx")


class _MemoryStorage:
    output_dir = "output"

    def __init__(self):
        self.files = {}
        self.directories = []

    def mkdir(self, path):
        self.directories.append(path)

    def write_bytes(self, path, content, content_type):
        self.files[path] = content

    def read_bytes(self, path):
        return self.files[path]

    def delete(self, path, missing_ok=False):
        self.files.pop(path, None)


def test_upload_marks_unknown_variables_without_crashing_and_supports_replace_download_delete():
    storage = _MemoryStorage()
    service = AgreementDocxTemplateService(storage)
    content = _docx_bytes(lambda document: document.add_paragraph("{{ agreement_number }} {{ missing_value }}"))

    metadata = service.upload(
        form_slug="sample",
        filename="umowa.docx",
        content=content,
        fields=[],
        uploaded_by_user_id=7,
    )

    assert metadata["valid"] is False
    assert metadata["unknown_variables"] == ["missing_value"]
    assert service.download(metadata) == content
    service.delete(metadata)
    assert storage.files == {}


def test_preview_reloads_current_docx_bytes_from_storage():
    storage = _MemoryStorage()
    service = AgreementDocxTemplateService(storage)
    path = service.storage_path("sample")
    storage.files[path] = _docx_bytes(lambda document: document.add_paragraph("Aktualny {{ agreement_number }}"))

    parsed = service.parse_stored_template({"storage_path": path, "html": "<p>stary</p>"})

    assert "Aktualny {{ agreement_number }}" in parsed.html
    assert "stary" not in parsed.html


def test_form_config_selects_docx_html_without_changing_legacy_html_source():
    service = FormConfigService()
    docx_config = service.normalize_form_config({
        "workflow": {
            "requires_contract": True,
            "managed_documents": True,
            "contract_template_source": "docx",
            "contract_template_html": "<p>legacy</p>",
            "contract_docx_template": {"html": "<p>{{ agreement_number }}</p>", "valid": True},
        }
    })
    html_config = service.normalize_form_config({
        "workflow": {"requires_contract": True, "managed_documents": True, "contract_template_html": "<p>legacy</p>"}
    })

    assert next(item for item in docx_config["documents"] if item["id"] == "agreement")["template_html"] == "<p>{{ agreement_number }}</p>"
    assert next(item for item in html_config["documents"] if item["id"] == "agreement")["template_html"] == "<p>legacy</p>"
