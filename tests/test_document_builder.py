import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from jinja2.sandbox import SandboxedEnvironment

from services.documents.declaration_template_context_service import (
    DeclarationVariableCatalog,
    build_declaration_render_context,
    declaration_builder_variable_catalog,
    declaration_criteria_catalog,
    declaration_variable_catalog,
)
from services.documents.document_builder_service import (
    document_builder_block_types,
    normalize_inline_runs,
    normalize_document_builder_document,
    render_document_builder_template,
    validate_document_builder_document,
)
from services.documents.docx_template_parser import parse_docx_template
from services.documents.agreement_docx_template_service import AgreementDocxTemplateService
from testing_support import InMemoryStorage
from services.form_config_service import FormConfigService


@pytest.mark.parametrize("document_type", ["declaration", "agreement"])
def test_dotted_fields_are_shared_layout_blocks_with_default_labels(document_type):
    document = {"version": 1, "blocks": [{"type": "dotted_fields"}]}

    normalized = normalize_document_builder_document(document, document_type)
    template = render_document_builder_template(normalized, document_type)

    assert "dotted_fields" in document_builder_block_types(document_type)
    assert normalized["blocks"] == [{
        "type": "dotted_fields",
        "left_label": "miejscowość i data",
        "right_label": "podpis",
    }]
    assert validate_document_builder_document(document, document_type=document_type) == []
    assert template.count('class="document-dotted-fields__line"') == 2
    assert template.count('class="document-dotted-fields__label"') == 2
    assert "miejscowość i data" in template and "podpis" in template
    assert normalize_document_builder_document(normalized, document_type) == normalized


def test_dotted_fields_keep_custom_labels_and_block_order():
    document = {"blocks": [
        {"type": "paragraph", "content": "Przed"},
        {"type": "dotted_fields", "left_label": "Miejsce podpisania", "right_label": "Podpis osoby"},
        {"type": "paragraph", "content": "Po"},
    ]}

    normalized = normalize_document_builder_document(document, "declaration")
    template = render_document_builder_template(normalized, "declaration")

    assert [block["type"] for block in normalized["blocks"]] == ["paragraph", "dotted_fields", "paragraph"]
    assert template.index("Przed") < template.index("Miejsce podpisania") < template.index("Po")
    assert "Podpis osoby" in template
    assert validate_document_builder_document(normalized, document_type="declaration") == []


def test_dotted_fields_escape_html_and_jinja_in_labels():
    document = {"blocks": [{
        "type": "dotted_fields",
        "left_label": '\"><img src=x onerror=alert(1)>',
        "right_label": "{{ missing_variable }} {% if secret %}",
    }]}

    template = render_document_builder_template(document, "agreement")
    rendered = SandboxedEnvironment(autoescape=True).from_string(template).render()

    assert "<img" not in template and "<img" not in rendered
    assert "{{ missing_variable }}" not in template
    assert "{% if secret %}" not in template
    assert "&#123;&#123;" in template
    assert validate_document_builder_document(document, document_type="agreement") == []


def test_dotted_fields_reject_empty_label_and_leave_older_documents_unchanged():
    old = {"version": 1, "blocks": [{"type": "paragraph", "content": "Dotychczasowy dokument"}]}
    original = json.loads(json.dumps(old))
    normalized = normalize_document_builder_document(old, "agreement")
    assert old == original
    assert [block["type"] for block in normalized["blocks"]] == ["paragraph"]
    assert "document-dotted-fields" not in render_document_builder_template(old, "agreement")
    assert validate_document_builder_document(old, document_type="agreement") == []

    invalid = {"blocks": [{"type": "dotted_fields", "left_label": "", "right_label": "Podpis"}]}
    assert any("etykiet" in error.message for error in validate_document_builder_document(invalid, document_type="agreement"))


def test_dotted_fields_are_available_in_shared_editor_without_submission_fields():
    root = Path(__file__).resolve().parents[1]
    editor = (root / "templates/admin/forms/_document_builder.html").read_text(encoding="utf-8")
    script = (root / "static/js/agreement_builder.js").read_text(encoding="utf-8")

    assert 'data-builder-add="dotted_fields"' in editor
    assert 'data-builder-field="left_label"' in script
    assert 'data-builder-field="right_label"' in script
    assert 'data-builder-move="up"' in script and 'data-builder-remove' in script


def test_inline_runs_are_merged_persisted_and_rendered_without_breaking_jinja():
    document = {
        "version": 1,
        "document_type": "agreement",
        "blocks": [{
            "type": "paragraph",
            "format": {"alignment": "justify"},
            "runs": [
                {"text": "Uczestnik: ", "bold": False},
                {"text": "{{ participant_name }}", "bold": True},
                {"text": " podpisuje", "bold": True},
                {"text": " umowę", "italic": True, "underline": True},
            ],
        }],
    }

    normalized = normalize_document_builder_document(document, "agreement")
    template = render_document_builder_template(normalized, "agreement")

    assert len(normalized["blocks"][0]["runs"]) == 3
    assert normalized["blocks"][0]["runs"][1]["text"] == "{{ participant_name }} podpisuje"
    assert "<strong>{{ participant_name }} podpisuje</strong>" in template
    assert 'class="document-text-underline"' in template
    assert "document-align-justify" in template
    assert validate_document_builder_document(normalized, document_type="agreement") == []
    rendered = SandboxedEnvironment(autoescape=True).from_string(template).render(participant_name="Jan Kowalski")
    assert "Jan Kowalski" in rendered


def test_inline_run_normalization_removes_only_editor_markers_across_repeated_round_trips():
    significant_text = "  Początek\n\u00a0środek końca  "
    runs = [
        {
            "text": "\u200b\ufeff" + significant_text + "\u200c\u200d\u2060\x00",
            "bold": True,
        }
    ]

    for _ in range(5):
        runs = json.loads(json.dumps(normalize_inline_runs(runs), ensure_ascii=False))

    assert runs == [
        {
            "text": significant_text,
            "bold": True,
            "italic": False,
            "underline": False,
        }
    ]
    template = render_document_builder_template(
        {"version": 1, "blocks": [{"type": "paragraph", "runs": runs}]},
        "agreement",
    )
    assert "  Początek<br>\u00a0środek końca  " in template


def test_inline_renderer_formats_only_requested_legal_phrase_in_stable_wrapper_order():
    document = {
        "version": 1,
        "document_type": "agreement",
        "blocks": [{
            "type": "paragraph",
            "format": {"alignment": "justify"},
            "runs": [
                {"text": "Beneficjent zobowiązuje się przestrzegać postanowień "},
                {
                    "text": "Regulaminu projektu",
                    "bold": True,
                    "italic": True,
                    "underline": True,
                },
                {"text": "."},
            ],
        }],
    }

    template = render_document_builder_template(document, "agreement")

    expected = (
        "Beneficjent zobowiązuje się przestrzegać postanowień "
        '<strong><em><span class="document-text-underline">'
        "Regulaminu projektu"
        "</span></em></strong>."
    )
    assert expected in template
    assert template.count("<strong>") == 1


def test_inline_renderer_keeps_whole_jinja_variable_inside_marked_run():
    document = {
        "version": 1,
        "document_type": "agreement",
        "blocks": [{
            "type": "paragraph",
            "runs": [
                {"text": "Uczestnik: "},
                {"text": "{{ participant_name }}", "bold": True},
            ],
        }],
    }

    template = render_document_builder_template(document, "agreement")
    rendered = SandboxedEnvironment(autoescape=True).from_string(template).render(participant_name="Jan Kowalski")

    assert "<strong>{{ participant_name }}</strong>" in template
    assert "<strong>Jan Kowalski</strong>" in rendered


def test_shared_document_css_does_not_reset_semantic_bold_to_normal():
    css = (Path(__file__).resolve().parents[1] / "static" / "css" / "document_template.css").read_text(encoding="utf-8")

    assert ".document strong," in css
    assert ".document b," in css
    assert "font-weight: 700;" in css.split(".document strong,", 2)[-1].split("}", 1)[0]


def test_ordered_list_hierarchy_styles_round_trip_and_render_markers():
    document = {
        "version": 1,
        "blocks": [{
            "type": "ordered_list",
            "items": [
                {"content": "Pierwszy", "level": 0},
                {"content": "Pierwszy podpunkt", "level": 1},
                {"content": "Głębiej", "level": 2},
                {"content": "Drugi podpunkt", "level": 1},
                {"content": "Drugi", "level": 0},
            ],
            "list_styles": [
                {"level": 0, "marker": "decimal-dot", "indent_mm": 0},
                {"level": 1, "marker": "decimal-compound", "indent_mm": 8},
                {"level": 2, "marker": "alpha-paren", "indent_mm": 16},
                {"level": 3, "marker": "upper-roman-dot", "indent_mm": 24},
            ],
        }],
    }

    normalized = normalize_document_builder_document(document, "agreement")
    template = render_document_builder_template(normalized, "agreement")

    assert normalize_document_builder_document(normalized, "agreement") == normalized
    assert [item["level"] for item in normalized["blocks"][0]["items"]] == [0, 1, 2, 1, 0]
    assert [style["marker"] for style in normalized["blocks"][0]["list_styles"]] == [
        "decimal-dot", "decimal-compound", "alpha-paren", "upper-roman-dot"
    ]
    assert [marker in template for marker in (">1.</span>", ">1.1</span>", ">a)</span>", ">1.2</span>", ">2.</span>")] == [True] * 5
    assert "--document-list-indent: 16mm" in template
    assert validate_document_builder_document(normalized, document_type="agreement") == []


def test_ordered_list_rejects_level_jumps_and_invalid_configuration():
    invalid = {
        "version": 1,
        "blocks": [{
            "type": "ordered_list",
            "items": [{"content": "A", "level": 0}, {"content": "B", "level": 3}],
            "list_styles": [
                {"level": 0, "marker": "unknown", "indent_mm": -1},
                {"level": 1, "marker": "decimal-dot", "indent_mm": 7},
            ],
        }],
    }

    errors = validate_document_builder_document(invalid, document_type="agreement")

    assert any("pominąć poziomu" in error.message for error in errors)
    assert any("poziomów 1–4" in error.message or "marker" in error.message.lower() for error in errors)


def test_declaration_criteria_table_uses_qualification_source_and_tri_state_context():
    fields = [
        {"name": "zgoda", "label": "Zgoda formularza", "document_label": "Zgoda w deklaracji", "type": "checkbox"},
        {
            "name": "status",
            "label": "Status uczestnika",
            "type": "radio",
            "options": [
                {"value": "tak", "label": "Tak"},
                {"value": "nie", "label": "Nie"},
                {"value": "odmowa", "label": "Odmowa podania informacji"},
            ],
        },
    ]
    definition = {
        "fields": fields,
        "qualification_conditions": {
            "enabled": True,
            "conditions": [
                {"field_name": "zgoda", "is_active": True},
                {"field_name": "status", "is_active": True},
                {"field_name": "pomijane", "is_active": False},
            ],
        },
    }
    criteria = declaration_criteria_catalog(definition, fields)
    builder_catalog = declaration_builder_variable_catalog(fields, form_definition=definition)
    document = {
        "version": 1,
        "blocks": [{
            "type": "criteria_table",
            "criteria": [{"field_key": "zgoda", "label": "stara kopia"}, {"field_key": "status"}],
            "show_number": True,
            "show_header": True,
        }],
    }

    normalized = normalize_document_builder_document(document, "declaration")
    assert [item["field_key"] for item in criteria] == ["zgoda", "status"]
    assert criteria[0]["label"] == "Zgoda w deklaracji"
    assert criteria[1]["answer_type"] == "choice"
    assert any(item["type"] == "criterion" and item["field_key"] == "status" for item in builder_catalog)
    assert normalized["blocks"][0]["criteria"] == [{"field_key": "zgoda"}, {"field_key": "status"}]
    assert validate_document_builder_document(
        normalized,
        fields,
        "declaration",
        form_definition=definition,
    ) == []

    template = render_document_builder_template(normalized, "declaration")
    context = build_declaration_render_context(
        {"zgoda": None, "status": "odmowa"},
        form_definition=definition,
        fields=fields,
    )
    rendered = SandboxedEnvironment(autoescape=True).from_string(template).render(**context)

    assert context["criterion_zgoda_yes_checked"] == ""
    assert context["criterion_zgoda_no_checked"] == ""
    assert context["criterion_status_other_enabled"] == "X"
    assert context["criterion_status_other_checked"] == "X"
    assert context["criterion_status_other_label"] == "Odmowa podania informacji"
    assert "Zgoda w deklaracji" in rendered
    assert "Odmowa podania informacji" in rendered


def test_declaration_builder_uses_form_fields_and_rejects_training_blocks():
    fields = [{"name": "zgoda", "label": "Zgoda na udział", "type": "checkbox"}]
    document = {
        "version": 1,
        "blocks": [
            {"type": "heading", "level": 1, "runs": [{"text": "Deklaracja", "bold": True}]},
            {"type": "form_field", "field": "zgoda", "label": "Zgoda na udział", "display": "yes_no"},
            {"type": "statement", "runs": [{"text": "Potwierdzam prawdziwość danych."}]},
            {"type": "participant_signature", "label": "Podpis uczestnika"},
        ],
    }

    assert validate_document_builder_document(document, fields, "declaration") == []
    template = render_document_builder_template(document, "declaration")
    assert "{{ zgoda_yes_no }}" in template
    assert "document--declaration" in template

    invalid = {"version": 1, "blocks": [{"type": "training_table", "columns": ["name"]}]}
    errors = validate_document_builder_document(invalid, fields, "declaration")
    assert errors and "Nieobsługiwany typ bloku" in errors[0].message


def test_declaration_context_has_boolean_representations_and_no_training_catalog():
    fields = [
        {"name": "zgoda", "label": "Zgoda", "type": "checkbox"},
        {"name": "tematy", "label": "Tematy", "type": "multi_select"},
    ]
    context = build_declaration_render_context(
        {"zgoda": "on", "tematy": ["Kadry", "Płace"], "imiona": "Anna", "nazwisko": "Nowak"},
        form_definition={"title": "Projekt", "fields": fields},
        fields=fields,
    )
    names = DeclarationVariableCatalog.context_names(fields)
    categories = {item["category"] for item in declaration_variable_catalog(fields)}

    assert context["zgoda_display"] == "Tak"
    assert context["zgoda_yes_no"] == "Tak"
    assert context["tematy_display"] == "Kadry, Płace"
    assert context["selected_trainings"] == []
    assert "training_name" not in names
    assert "selected_trainings" not in names
    assert "Deklaracja" not in categories
    assert "Zgłoszenie" in categories


def test_declaration_catalog_uses_assigned_fields_and_legacy_keeps_all_fields():
    legacy_fields = [
        {"name": "imie", "label": "Imię", "type": "text"},
        {"name": "numer_legitymacji", "label": "Numer legitymacji", "type": "text"},
    ]
    assert {"imie", "numer_legitymacji"}.issubset({
        item["name"] for item in declaration_variable_catalog(
            legacy_fields, form_definition={"fields": legacy_fields}
        )
    })

    assigned_fields = [
        {**legacy_fields[0], "document_usage": {"declaration": False}},
        {**legacy_fields[1], "document_usage": {"declaration": True}},
    ]
    names = {
        item["name"] for item in declaration_variable_catalog(
            assigned_fields, form_definition={"fields": assigned_fields}
        )
    }
    assert "numer_legitymacji" in names
    assert "imie" not in names


def test_declaration_field_filter_keeps_nonempty_section_and_assigned_text():
    from services.documents.declaration_flow_service import DeclarationFlowService

    fields = [
        {"type": "section", "label": "Pusta sekcja"},
        {"name": "excluded", "label": "Poza deklaracją", "type": "text",
         "document_usage": {"declaration": False}},
        {"type": "section", "label": "Sekcja deklaracji"},
        {"name": "intro", "type": "static_text", "content": "Treść deklaracji",
         "document_usage": {"declaration": True}},
        {"name": "included", "label": "W deklaracji", "type": "text",
         "document_usage": {"declaration": True}},
        {"type": "section", "label": "Informacja"},
        {"name": "note", "type": "static_text", "content": "Ważna informacja",
         "document_usage": {"declaration": True}},
    ]
    document = {"id": "declaration", "label": "Deklaracja", "fields": fields}
    definition = DeclarationFlowService.document_fields_definition(document)
    assert [field.get("name") or field.get("label") for field in definition["fields"]] == [
        "Sekcja deklaracji", "intro", "included", "Informacja", "note"]
    from services.documents.declaration_template_context_service import declaration_variable_catalog
    variable_names = {item["name"] for item in declaration_variable_catalog(fields)}
    assert "included" in variable_names
    assert "note" not in variable_names and "note_display" not in variable_names

    legacy = [{key: value for key, value in field.items() if key != "document_usage"}
              for field in fields]
    assert "excluded" in {field.get("name") for field in
                          DeclarationFlowService.document_fields_definition({**document, "fields": legacy})["fields"]}


def test_declaration_address_uses_submission_data_behind_empty_columns():
    from services.documents.declaration_template_context_service import (
        build_declaration_render_context, declaration_preview_context,
    )

    # Given: legacy DB columns are blank while the submission JSON has the address.
    address = {"ulica": "Przykładowa", "nr_budynku": "12", "nr_lokalu": "3",
               "kod_pocztowy": "65-001", "miejscowosc": "Zielona Góra"}
    row = {**{name: "" for name in address}, "data_json": address}
    context = build_declaration_render_context(row)
    assert context["participant_address"] == "ul. Przykładowa 12/3, 65-001 Zielona Góra"
    assert context["participant_address_inline"] == context["participant_address"]

    # Expected: no empty slash or punctuation when address parts are missing.
    no_apartment = build_declaration_render_context({"data_json": {**address, "nr_lokalu": ""}})
    assert no_apartment["participant_address"] == "ul. Przykładowa 12, 65-001 Zielona Góra"
    partial = build_declaration_render_context({"data_json": {
        "ulica": "null", "nr_budynku": None, "nr_lokalu": "None",
        "kod_pocztowy": "65-001", "miejscowosc": "Zielona Góra"}})
    assert partial["participant_address"] == "65-001 Zielona Góra"
    aliases = build_declaration_render_context({"data_json": {
        "street": "  Polna  ", "building_number": "4", "postal_code": "65-001", "city": "Zielona   Góra"}})
    assert aliases["participant_address"] == "ul. Polna 4, 65-001 Zielona Góra"

    # When: preview uses a real, partial submission, its example address is cleared.
    preview = declaration_preview_context(submission={"ulica": "", "data_json": {
        "kod_pocztowy": "65-001", "miejscowosc": "Zielona Góra"}})
    assert preview["participant_address"] == "65-001 Zielona Góra"
    assert declaration_preview_context(submission={"data_json": {}})["participant_address"] == ""


def test_docx_import_preserves_inline_styles_and_whole_split_jinja_token():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Dane ")
    bold = paragraph.add_run("uczestnika ")
    bold.bold = True
    token_open = paragraph.add_run("{{")
    token_open.bold = True
    token_middle = paragraph.add_run(" participant_")
    token_middle.italic = True
    token_end = paragraph.add_run("name }}")
    token_end.underline = True
    suffix = paragraph.add_run(" ważne")
    suffix.italic = True
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_docx_template(buffer.getvalue(), document_type="declaration")
    block = parsed.builder_document["blocks"][0]

    assert parsed.variables == ("participant_name",)
    assert "{{ participant_name }}" in parsed.html
    assert "&#123;&#123;" not in parsed.html
    assert any(run["bold"] and "uczestnika" in run["text"] for run in block["runs"])
    assert any(run["italic"] and "ważne" in run["text"] for run in block["runs"])
    SandboxedEnvironment(autoescape=True).parse(parsed.html)


def test_form_config_supports_shared_builder_for_declaration_and_legacy_html():
    service = FormConfigService()
    fresh = service.normalize_form_config({"workflow": {"requires_declaration": True}})
    legacy = service.normalize_form_config({"workflow": {"requires_declaration": True, "declaration_template_html": "<p>Legacy</p>"}})

    assert fresh["workflow"]["declaration_template_source"] == "builder"
    declaration = next(item for item in fresh["documents"] if item["id"] == "declaration")
    assert declaration["template_source"] == "builder"
    assert declaration["builder_document"]["document_type"] == "declaration"
    assert legacy["workflow"]["declaration_template_source"] == "html"
    assert next(item for item in legacy["documents"] if item["id"] == "declaration")["template_html"] == "<p>Legacy</p>"


def test_form_config_preserves_selected_template_sources_for_both_document_types():
    service = FormConfigService()
    normalized = service.normalize_form_config({
        "workflow": {
            "requires_contract": True,
            "requires_declaration": True,
            "contract_template_source": "docx",
            "declaration_template_source": "html",
            "contract_docx_template": {"storage_path": "templates/agreement.docx"},
            "declaration_template_html": "<p>{{ participant_name }}</p>",
        }
    })

    assert normalized["workflow"]["contract_template_source"] == "docx"
    assert normalized["workflow"]["declaration_template_source"] == "html"
    agreement = next(item for item in normalized["documents"] if item["id"] == "agreement")
    declaration = next(item for item in normalized["documents"] if item["id"] == "declaration")
    assert agreement["template_source"] == "docx"
    assert declaration["template_source"] == "html"


def test_custom_docx_uses_existing_parser_and_isolated_template_path():
    # Given: a DOCX template for an additional generated document.
    source = Document()
    source.add_paragraph("Zaświadczenie dla {{ participant_name }}")
    buffer = BytesIO()
    source.save(buffer)
    storage = InMemoryStorage({})
    service = AgreementDocxTemplateService(storage)
    # When: the shared DOCX service imports it with a generic document ID.
    metadata = service.upload(form_slug="form", filename="certificate.docx", content=buffer.getvalue(),
        fields=[], uploaded_by_user_id=1, document_type="document", document_id="employment_certificate",
        form_definition={"fields": []}, version_key="1-1-2")
    # Expected: the template is version scoped and the renderer sees its text.
    assert "/employment_certificate/versions/1-1-2/" in metadata["storage_path"]
    assert "{{ participant_name }}" in service.parse_stored_template(metadata).html
