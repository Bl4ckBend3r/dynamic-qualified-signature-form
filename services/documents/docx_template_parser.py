from __future__ import annotations

import html
import re
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.ns import qn
from docx.opc.exceptions import PackageNotFoundError
from jinja2 import TemplateSyntaxError, meta
from jinja2.sandbox import SandboxedEnvironment


PARSER_VERSION = "1.3"
_VARIABLE_RE = re.compile(r"{{\s*([A-Za-z_][A-Za-z0-9_]*)\s*}}", re.DOTALL)
_JINJA_TOKEN_RE = re.compile(r"({{.*?}}|{%.*?%}|{#.*?#})", re.DOTALL)
_JINJA_FRAGMENT_RE = re.compile(r"({{|}}|{%|%}|{#|#})")
_SUPPORTED_JINJA_STATEMENTS = {"if", "elif", "else", "endif", "for", "endfor", "set"}
_LEGAL_SECTION_RE = re.compile(r"^\s*§\s*\d+[A-Za-z]?\s*\.?\s*$")
_TEXT_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+")


class DocxTemplateParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocxTemplate:
    html: str
    variables: tuple[str, ...]
    warnings: tuple[str, ...]
    builder_document: dict | None = None


@dataclass(frozen=True)
class _ListInfo:
    tag: str
    identity: str
    level: int


def parse_docx_template(content: bytes, *, document_type: str = "agreement") -> ParsedDocxTemplate:
    if not content:
        raise DocxTemplateParseError("Plik DOCX jest pusty.")
    try:
        document = Document(BytesIO(content))
    except (PackageNotFoundError, BadZipFile, ValueError, KeyError) as exc:
        raise DocxTemplateParseError("Plik nie jest poprawnym dokumentem DOCX.") from exc

    document_type = str(document_type or "agreement").strip().casefold()
    if document_type not in {"agreement", "declaration"}:
        raise DocxTemplateParseError("Nieobsługiwany typ dokumentu DOCX.")
    variables: set[str] = set()
    warnings: list[str] = []
    blocks: list[str] = []
    pending_list: list[str] = []
    pending_list_tag: str | None = None
    pending_list_identity: str | None = None
    expects_section_title = False

    def flush_list() -> None:
        nonlocal pending_list_tag, pending_list_identity
        if pending_list_tag and pending_list:
            blocks.append(f"<{pending_list_tag} class=\"document-list\">{''.join(pending_list)}</{pending_list_tag}>")
        pending_list.clear()
        pending_list_tag = None
        pending_list_identity = None

    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            raw_text = block.text or ""
            text = _paragraph_text(block, variables, warnings)
            if _LEGAL_SECTION_RE.fullmatch(raw_text):
                flush_list()
                blocks.append(f'<div class="document-section-number">{text}</div>')
                expects_section_title = True
                continue

            list_info = _list_info(block)
            if not _visible_text(text):
                if list_info:
                    # Word can retain numPr on a visually empty paragraph. It must
                    # not consume a visible list number or produce an empty <li>.
                    continue
                flush_list()
                blocks.append(_paragraph_html(block, text))
                continue

            if expects_section_title:
                expects_section_title = False
                if _looks_like_section_title(block, raw_text, list_info):
                    flush_list()
                    blocks.append(f'<div class="document-section-title">{text}</div>')
                    continue

            if list_info:
                if pending_list_identity and pending_list_identity != list_info.identity:
                    flush_list()
                if pending_list_identity is None:
                    pending_list_tag = list_info.tag
                    pending_list_identity = list_info.identity
                level = max(0, min(list_info.level, 8))
                pending_list.append(f'<li class="document-list__item document-list__item--level-{level}">{text}</li>')
                continue
            flush_list()
            heading_lines = _visual_lines(text) if _heading_level(block) else []
            if len(heading_lines) > 1:
                blocks.append(_paragraph_html(block, heading_lines[0]))
                alignment = _paragraph_alignment(block, default="center")
                for line in heading_lines[1:]:
                    blocks.append(f'<p class="document-paragraph document-align-{alignment} document-heading-context">{line}</p>')
                continue
            blocks.append(_paragraph_html(block, text))
        else:
            flush_list()
            blocks.append(_table_html(block, variables, warnings))
    flush_list()

    if not blocks or not any(_visible_text(block) for block in blocks):
        raise DocxTemplateParseError("Dokument DOCX nie zawiera treści szablonu.")
    if _contains_unsupported_xml(document):
        warnings.append("Dokument zawiera obrazy, pola tekstowe lub obiekty, których parser nie przenosi do HTML.")

    template_html = f'<main class="document document--{document_type} document--docx">' + "".join(blocks) + "</main>"
    try:
        parsed_jinja = SandboxedEnvironment(autoescape=True).parse(template_html)
        variables.update(meta.find_undeclared_variables(parsed_jinja))
    except TemplateSyntaxError as exc:
        warnings.append(f"Niepoprawna składnia Jinja: {exc.message}.")
    builder_document = _builder_document(document, variables, warnings, document_type=document_type)
    return ParsedDocxTemplate(
        html=template_html,
        variables=tuple(sorted(variables)),
        warnings=tuple(dict.fromkeys(warnings)),
        builder_document=builder_document,
    )


def _iter_blocks(parent: DocxDocument):
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _paragraph_text(paragraph: Paragraph, variables: set[str], warnings: list[str]) -> str:
    return _runs_html(_paragraph_runs(paragraph, variables, warnings))


def _paragraph_runs(paragraph: Paragraph, variables: set[str], warnings: list[str]) -> list[dict]:
    raw = paragraph.text or ""
    matches = list(_JINJA_TOKEN_RE.finditer(raw))
    variables.update(match.group(1) for match in _VARIABLE_RE.finditer(raw))
    stripped = _JINJA_TOKEN_RE.sub("", raw)
    if _JINJA_FRAGMENT_RE.search(stripped):
        warnings.append("Wykryto niepełny znacznik Jinja. Sprawdź pary nawiasów {{ }}, {% %} lub {# #}.")

    styled_ranges: list[tuple[int, int, dict[str, bool]]] = []
    range_cursor = 0
    for run in paragraph.runs:
        run_text = run.text or ""
        if not run_text:
            continue
        style = {"bold": bool(run.bold), "italic": bool(run.italic), "underline": bool(run.underline)}
        styled_ranges.append((range_cursor, range_cursor + len(run_text), style))
        range_cursor += len(run_text)
    if range_cursor != len(raw):
        styled_ranges = [(0, len(raw), {"bold": False, "italic": False, "underline": False})]

    pieces: list[dict] = []
    cursor = 0
    for match in matches:
        _append_literal_runs(pieces, raw, cursor, match.start(), styled_ranges)
        token = match.group(0)
        _append_model_run(pieces, token, _style_at(styled_ranges, match.start()))
        if not _is_supported_jinja_token(token):
            warnings.append("Nieobsługiwany znacznik Jinja został zachowany jako zwykły tekst.")
        cursor = match.end()
    _append_literal_runs(pieces, raw, cursor, len(raw), styled_ranges)
    return pieces


def _append_literal_runs(pieces: list[dict], raw: str, start: int, end: int, ranges: list[tuple[int, int, dict[str, bool]]]) -> None:
    cursor = start
    while cursor < end:
        matching = next((item for item in ranges if item[0] <= cursor < item[1]), None)
        next_end = min(end, matching[1]) if matching else end
        style = matching[2] if matching else {"bold": False, "italic": False, "underline": False}
        _append_model_run(pieces, raw[cursor:next_end], style)
        cursor = next_end


def _append_model_run(pieces: list[dict], text: str, style: dict[str, bool]) -> None:
    if not text:
        return
    run = {"text": text, **style}
    token_boundary = bool(_JINJA_TOKEN_RE.fullmatch(text)) or bool(
        pieces and _JINJA_TOKEN_RE.fullmatch(str(pieces[-1].get("text") or ""))
    )
    if pieces and not token_boundary and all(pieces[-1].get(name) == run[name] for name in ("bold", "italic", "underline")):
        pieces[-1]["text"] += text
    else:
        pieces.append(run)


def _style_at(ranges: list[tuple[int, int, dict[str, bool]]], position: int) -> dict[str, bool]:
    match = next((item for item in ranges if item[0] <= position < item[1]), None)
    return dict(match[2]) if match else {"bold": False, "italic": False, "underline": False}


def _runs_html(runs: list[dict]) -> str:
    pieces = []
    for run in runs:
        raw = str(run.get("text") or "")
        rendered = raw if _JINJA_TOKEN_RE.fullmatch(raw) and _is_supported_jinja_token(raw) else _escape_literal(raw)
        rendered = rendered.replace("\n", "<br>")
        if run.get("underline"):
            rendered = f'<span class="document-text-underline">{rendered}</span>'
        if run.get("italic"):
            rendered = f"<em>{rendered}</em>"
        if run.get("bold"):
            rendered = f"<strong>{rendered}</strong>"
        pieces.append(rendered)
    return "".join(pieces)


def _paragraph_html(paragraph: Paragraph, text: str) -> str:
    heading_level = _heading_level(paragraph)
    tag = f"h{heading_level}" if heading_level else "p"
    classes = ["document-paragraph"]
    alignment = _paragraph_alignment(paragraph)
    if alignment:
        classes.append(f"document-align-{alignment}")
    if not text:
        classes.append("document-paragraph--empty")
    return f'<{tag} class="{" ".join(classes)}">{text or "&nbsp;"}</{tag}>'


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = str(getattr(paragraph.style, "name", "") or "").casefold()
    match = re.search(r"(?:heading|nagłówek|naglowek)\s*([1-3])", style_name)
    return int(match.group(1)) if match else None


def _paragraph_alignment(paragraph: Paragraph, *, default: str | None = None) -> str | None:
    return {
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.LEFT: "left",
    }.get(paragraph.alignment, default)


def _visual_lines(content: str) -> list[str]:
    return [line.strip() for line in re.split(r"<br\s*/?>", content, flags=re.IGNORECASE) if _visible_text(line)]


def _looks_like_section_title(paragraph: Paragraph, raw_text: str, list_info: _ListInfo | None) -> bool:
    if list_info or _TEXT_NUMBER_RE.match(raw_text):
        return False
    style_name = str(getattr(paragraph.style, "name", "") or "").casefold()
    if re.search(r"(?:heading|nagłówek|naglowek)\s*[1-3]", style_name):
        return True
    candidate = raw_text.strip()
    return bool(candidate) and len(candidate) <= 160 and not candidate.endswith((".", ";", ":", "?", "!"))


def _list_info(paragraph: Paragraph) -> _ListInfo | None:
    style = str(getattr(paragraph.style, "name", "") or "").casefold()
    num_pr = _paragraph_num_pr(paragraph)
    if num_pr is None and not any(token in style for token in ("list", "lista", "bullet", "number", "punkt")):
        return None
    num_id = _xml_value(getattr(num_pr, "numId", None)) if num_pr is not None else None
    level = int(_xml_value(getattr(num_pr, "ilvl", None)) or 0) if num_pr is not None else 0
    number_format = _numbering_format(paragraph, num_id, level)
    tag = "ul" if number_format == "bullet" or any(token in style for token in ("bullet", "punkt")) else "ol"
    identity = f"num:{num_id}" if num_id is not None else f"style:{style}:{tag}"
    return _ListInfo(tag=tag, identity=identity, level=level)


def _paragraph_num_pr(paragraph: Paragraph):
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None and paragraph_properties.numPr is not None:
        return paragraph_properties.numPr
    style_element = getattr(getattr(paragraph, "style", None), "element", None)
    style_properties = getattr(style_element, "pPr", None)
    return getattr(style_properties, "numPr", None)


def _xml_value(element) -> str | None:
    value = getattr(element, "val", None)
    return str(value) if value is not None else None


def _numbering_format(paragraph: Paragraph, num_id: str | None, level: int) -> str | None:
    if num_id is None:
        return None
    try:
        numbering = paragraph.part.numbering_part.element
        abstract_id = None
        for number in numbering.findall(qn("w:num")):
            if number.get(qn("w:numId")) == num_id:
                reference = number.find(qn("w:abstractNumId"))
                abstract_id = reference.get(qn("w:val")) if reference is not None else None
                break
        if abstract_id is None:
            return None
        for abstract in numbering.findall(qn("w:abstractNum")):
            if abstract.get(qn("w:abstractNumId")) != abstract_id:
                continue
            for level_element in abstract.findall(qn("w:lvl")):
                if int(level_element.get(qn("w:ilvl")) or 0) != level:
                    continue
                format_element = level_element.find(qn("w:numFmt"))
                return format_element.get(qn("w:val")) if format_element is not None else None
    except (AttributeError, KeyError, TypeError, ValueError):
        return None
    return None


def _table_html(table: Table, variables: set[str], warnings: list[str]) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            contents = "".join(_paragraph_html(p, _paragraph_text(p, variables, warnings)) for p in cell.paragraphs)
            cells.append(f"<td>{contents or '&nbsp;'}</td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<table class="document-table document-table--docx"><tbody>' + "".join(rows) + "</tbody></table>"


def _contains_unsupported_xml(document: DocxDocument) -> bool:
    xml = document.element.xml
    return any(marker in xml for marker in ("<w:drawing", "<w:pict", "<w:txbxContent", "<w:object"))


def _visible_text(markup: str) -> str:
    return re.sub(r"<[^>]+>|&nbsp;", "", markup).strip()


def _escape_literal(value: str) -> str:
    return html.escape(value).replace("{", "&#123;").replace("}", "&#125;")


def _is_supported_jinja_token(token: str) -> bool:
    if token.startswith("{{") or token.startswith("{#"):
        return True
    statement = token[2:-2].strip().split(maxsplit=1)
    return bool(statement and statement[0] in _SUPPORTED_JINJA_STATEMENTS)


def _builder_document(document: DocxDocument, variables: set[str], warnings: list[str], *, document_type: str) -> dict:
    """Convert the same Word structure to stable editable document blocks."""
    blocks: list[dict] = []
    pending_items: list[dict] = []
    pending_list_type = "ordered_list"
    pending_identity: str | None = None
    pending_section_number: str | None = None

    def flush_list() -> None:
        nonlocal pending_identity
        if pending_items:
            blocks.append({"type": pending_list_type, "items": list(pending_items)})
        pending_items.clear()
        pending_identity = None

    def flush_section() -> None:
        nonlocal pending_section_number
        if pending_section_number is not None:
            if document_type == "agreement":
                blocks.append({"type": "agreement_section", "number": pending_section_number, "title": "Paragraf umowy"})
            else:
                blocks.append(_plain_text_block("heading", pending_section_number, alignment="left", bold=True, level=2))
            pending_section_number = None

    for block in _iter_blocks(document):
        if isinstance(block, Table):
            flush_list()
            flush_section()
            rows = []
            for row in block.rows:
                rows.append([_cell_builder_value(cell.paragraphs, variables, warnings) for cell in row.cells])
            blocks.append({"type": "table", "header": bool(rows), "rows": rows})
            continue

        raw_text = block.text or ""
        runs = _paragraph_runs(block, variables, warnings)
        content = _runs_html(runs)
        plain_content = "".join(run["text"] for run in runs)
        list_info = _list_info(block)
        if _LEGAL_SECTION_RE.fullmatch(raw_text):
            flush_list()
            flush_section()
            pending_section_number = plain_content
            continue
        if pending_section_number is not None:
            if _visible_text(content):
                if document_type == "agreement":
                    blocks.append({"type": "agreement_section", "number": pending_section_number, "title": plain_content})
                else:
                    blocks.append(_plain_text_block("heading", f"{pending_section_number} {plain_content}".strip(), alignment="left", bold=True, level=2))
                pending_section_number = None
                continue
            flush_section()
        if not _visible_text(content):
            if list_info:
                continue
            flush_list()
            continue
        if list_info:
            if pending_identity and pending_identity != list_info.identity:
                flush_list()
            if pending_identity is None:
                pending_identity = list_info.identity
                pending_list_type = "bullet_list" if list_info.tag == "ul" else "ordered_list"
            pending_items.append({"content": plain_content, "runs": runs, "level": max(0, min(list_info.level, 8))})
            continue

        flush_list()
        heading_level = _heading_level(block)
        alignment = _paragraph_alignment(block, default="center" if heading_level else "left")
        if heading_level:
            lines = _split_model_runs_lines(runs)
            first_runs = lines[0] if lines else runs
            blocks.append(_runs_text_block("heading", first_runs, alignment=alignment, level=heading_level))
            for line_runs in lines[1:]:
                blocks.append(_runs_text_block("paragraph", line_runs, alignment=alignment))
        else:
            blocks.append(_runs_text_block("paragraph", runs, alignment=alignment))

    flush_list()
    flush_section()
    return {"version": 1, "document_type": document_type, "blocks": blocks}


def _cell_builder_value(paragraphs, variables: set[str], warnings: list[str]):
    combined: list[dict] = []
    for index, paragraph in enumerate(paragraphs):
        if index:
            _append_model_run(combined, "\n", {"bold": False, "italic": False, "underline": False})
        for run in _paragraph_runs(paragraph, variables, warnings):
            _append_model_run(combined, run["text"], {name: bool(run.get(name)) for name in ("bold", "italic", "underline")})
    return {"content": "".join(run["text"] for run in combined), "runs": combined}


def _split_model_runs_lines(runs: list[dict]) -> list[list[dict]]:
    lines: list[list[dict]] = [[]]
    for run in runs:
        parts = str(run.get("text") or "").split("\n")
        style = {name: bool(run.get(name)) for name in ("bold", "italic", "underline")}
        for index, part in enumerate(parts):
            if part:
                _append_model_run(lines[-1], part, style)
            if index < len(parts) - 1:
                lines.append([])
    return [line for line in lines if any(str(run.get("text") or "").strip() for run in line)]


def _runs_text_block(block_type: str, runs: list[dict], *, alignment: str, level: int | None = None) -> dict:
    result = {
        "type": block_type,
        "format": {"alignment": alignment},
        "content": "".join(run["text"] for run in runs),
        "runs": runs,
    }
    if level is not None:
        result["level"] = level
    return result


def _plain_text_block(block_type: str, content: str, *, alignment: str, bold: bool, level: int | None = None) -> dict:
    return _runs_text_block(
        block_type,
        [{"text": content, "bold": bold, "italic": False, "underline": False}],
        alignment=alignment,
        level=level,
    )
